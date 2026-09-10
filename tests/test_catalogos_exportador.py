"""O Excel gerado a partir da base.

O ficheiro é para consultar, e por isso o que estes testes guardam é o que se
lê nele: uma linha por referência, uma coluna por espessura, e o aviso escrito
onde a tabela do fornecedor se contradiz. Um exportador que perca uma coluna de
preço não rebenta — dá um ficheiro bonito e errado.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db.catalogos import SCHEMA_CATALOGOS, catalogos_metadata
from app.services.catalogos import exportador
from app.services.catalogos.base import ArtigoCatalogo, TabelaCatalogo
from app.services.catalogos.importador import importar


@pytest.fixture()
def session():
    engine = create_engine(
        "sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False}
    )

    @event.listens_for(engine, "connect")
    def _anexar(dbapi_connection, _record):  # noqa: ANN001
        dbapi_connection.execute(f"ATTACH DATABASE ':memory:' AS {SCHEMA_CATALOGOS}")

    catalogos_metadata.create_all(engine)
    try:
        with Session(engine) as db:
            yield db
    finally:
        engine.dispose()


def _placa(espessura: str, preco: str, **extra) -> ArtigoCatalogo:
    dados = dict(
        chave_natural=f"F037|ST76|PB STD|{espessura}",
        referencia="F037",
        descricao=f"EGGER F037 ST76 · Travertino Taormina · {espessura}",
        unidade="M2",
        nome_design="Travertino Taormina",
        familia="Eurodekor",
        seccao="Eurodekor Aglomerado",
        grupo="Grupo 8",
        fabricante="EGGER",
        substrato="PB STD",
        espessura_mm=Decimal(espessura.replace("mm", "")),
        acabamento="ST76",
        preco=Decimal(preco),
        atributos={
            "folha": "Stock_B&F_Egger",
            "espessura": espessura,
            "tipo_produto": "Eurodekor Aglomerado revestido",
            "substrato_origem": "Aglomerado",
        },
    )
    dados.update(extra)
    return ArtigoCatalogo(**dados)


def _ferragem(**extra) -> ArtigoCatalogo:
    dados = dict(
        chave_natural="22.8000|Kit capas Cinza",
        referencia="22.8000",
        descricao="BLUM · Kit capas Cinza",
        unidade="UN",
        familia="AVENTOS",
        seccao="BL.1.1.",
        grupo="Preço un.",
        fabricante="BLUM",
        preco=Decimal("9.02"),
        atributos={"folha": "Somapil_BLUM"},
    )
    dados.update(extra)
    return ArtigoCatalogo(**dados)


def _tabela(*artigos, **extra) -> TabelaCatalogo:
    dados = dict(
        fornecedor="Balbino & Faustino",
        fabricante="EGGER",
        nome="EGGER Balbino & Faustino",
        referencia_tabela="BF-82",
        data_tabela=date(2026, 4, 20),
        ficheiro_origem="12_Placas_Referencias_COMPLETO.xlsx#Stock_B&F_Egger",
        ficheiro_hash="a" * 64,
        unidade_preco="M2",
        artigos=artigos,
    )
    dados.update(extra)
    return TabelaCatalogo(**dados)


def _abrir(caminho):
    return load_workbook(caminho, data_only=True)


def test_a_base_vazia_recusa_escrever(session: Session, tmp_path) -> None:
    with pytest.raises(RuntimeError, match="importar_catalogos"):
        exportador.exportar(session, tmp_path / "x.xlsx")


def test_uma_linha_por_referencia_e_uma_coluna_por_espessura(
    session: Session, tmp_path
) -> None:
    importar(
        session,
        _tabela(
            _placa("8mm", "17.98"), _placa("19mm", "21.02"), _placa("10mm", "18.90")
        ),
    )
    session.commit()

    destino = tmp_path / "gerado.xlsx"
    resultado = exportador.exportar(session, destino)

    assert resultado.separadores == 1
    assert resultado.referencias == 1
    assert resultado.precos == 3

    livro = _abrir(destino)
    assert livro.sheetnames == ["Índice", "Stock_B&F_Egger"]
    folha = livro["Stock_B&F_Egger"]
    cabecalho = [celula.value for celula in folha[4]]
    # As espessuras saem por ordem de grandeza, não por ordem alfabética: por
    # texto o 10mm vinha antes do 8mm.
    assert cabecalho[-4:] == [
        "Preço Tabela 8mm",
        "Preço Tabela 10mm",
        "Preço Tabela 19mm",
        "Aviso",
    ]
    linha = {nome: celula.value for nome, celula in zip(cabecalho, folha[5])}
    assert linha["Referência"] == "F037"
    assert linha["ST/Acab"] == "ST76"
    assert linha["Substrato"] == "PB STD"
    assert linha["Substrato do fornecedor"] == "Aglomerado"
    assert linha["Tipo Produto"] == "Eurodekor Aglomerado revestido"
    assert linha["Preço Tabela 8mm"] == pytest.approx(17.98)
    assert linha["Preço Tabela 19mm"] == pytest.approx(21.02)
    # Uma célula vazia volta do openpyxl como None, não como "".
    assert not linha["Aviso"]


def test_o_indice_diz_de_onde_vem_cada_separador(session: Session, tmp_path) -> None:
    importar(session, _tabela(_placa("19mm", "21.02")))
    session.commit()

    destino = tmp_path / "gerado.xlsx"
    exportador.exportar(session, destino)

    folha = _abrir(destino)["Índice"]
    cabecalho = [celula.value for celula in folha[4]]
    linha = {nome: celula.value for nome, celula in zip(cabecalho, folha[5])}
    assert linha["Separador"] == "Stock_B&F_Egger"
    assert linha["Fornecedor"] == "Balbino & Faustino"
    assert linha["Tabela"] == "EGGER Balbino & Faustino"
    assert linha["Código"] == "BF-82"
    assert linha["Referências"] == 1
    assert linha["Preços"] == 1


def test_as_ferragens_saem_com_a_coluna_do_preco_unitario(
    session: Session, tmp_path
) -> None:
    importar(
        session,
        _tabela(
            _ferragem(),
            fornecedor="Somapil",
            fabricante="BLUM",
            nome="BLUM · Somapil",
            ficheiro_origem="12_Placas_Referencias_COMPLETO.xlsx#Somapil_BLUM",
            ficheiro_hash="b" * 64,
        ),
    )
    session.commit()

    destino = tmp_path / "gerado.xlsx"
    exportador.exportar(session, destino)

    folha = _abrir(destino)["Somapil_BLUM"]
    cabecalho = [celula.value for celula in folha[4]]
    assert "Preço Tabela 19mm" not in cabecalho
    assert cabecalho[-2:] == [exportador.CHAVE_PRECO_UNITARIO, "Aviso"]
    linha = {nome: celula.value for nome, celula in zip(cabecalho, folha[5])}
    assert linha["Descrição"] == "BLUM · Kit capas Cinza"
    assert linha["Família"] == "AVENTOS"
    assert linha[exportador.CHAVE_PRECO_UNITARIO] == pytest.approx(9.02)


def test_onde_a_tabela_se_contradiz_o_aviso_diz_qual_vale(
    session: Session, tmp_path
) -> None:
    def finsa(design: str, preco: str) -> ArtigoCatalogo:
        return _placa(
            "19mm",
            preco,
            chave_natural=f"688B|{design}|YOKU|AGL HID|19mm",
            referencia="688B",
            nome_design=design,
            acabamento="YOKU",
            atributos={
                "folha": "Stock_B&F_Egger",
                "espessura": "19mm",
                "substrato_origem": "AGL HID",
            },
        )

    importar(session, _tabela(finsa("CARYA WOOD", "17.36"), finsa("TIVOLI ASH", "16.63")))
    session.commit()

    destino = tmp_path / "gerado.xlsx"
    exportador.exportar(session, destino)

    folha = _abrir(destino)["Stock_B&F_Egger"]
    cabecalho = [celula.value for celula in folha[4]]
    avisos = {}
    for linha in folha.iter_rows(min_row=5, values_only=True):
        valores = dict(zip(cabecalho, linha))
        avisos[valores["Nome Design"]] = valores["Aviso"]

    assert len(avisos) == 2
    for aviso in avisos.values():
        assert "vale o mais alto, 17.36" in aviso


def test_um_separador_por_folha(session: Session, tmp_path) -> None:
    importar(session, _tabela(_placa("19mm", "21.02")))
    importar(
        session,
        _tabela(
            _ferragem(),
            fornecedor="Somapil",
            fabricante="BLUM",
            nome="BLUM · Somapil",
            ficheiro_origem="12_Placas_Referencias_COMPLETO.xlsx#Somapil_BLUM",
            ficheiro_hash="b" * 64,
        ),
    )
    session.commit()

    destino = tmp_path / "gerado.xlsx"
    resultado = exportador.exportar(session, destino)

    assert resultado.separadores == 2
    assert _abrir(destino).sheetnames == ["Índice", "Stock_B&F_Egger", "Somapil_BLUM"]


def test_o_nome_do_separador_cabe_no_excel_e_nao_se_repete() -> None:
    """Trinta e um caracteres, e nenhum nome perdido por colisão."""
    usados: set[str] = set()
    comprido = "Stock_de_um_fornecedor_com_nome_muito_grande_2026"
    primeiro = exportador._nome_separador(comprido, usados)
    segundo = exportador._nome_separador(comprido, usados)

    assert len(primeiro) <= 31
    assert len(segundo) <= 31
    assert primeiro != segundo
    assert exportador._nome_separador("Stock/Placas", usados) == "Stock-Placas"


def test_o_caminho_por_omissao_fica_na_pasta_dos_gerados(tmp_path) -> None:
    caminho = exportador.caminho_por_omissao(tmp_path)
    assert caminho.parent.name == exportador.PASTA_GERADOS
    assert caminho.name == exportador.FICHEIRO_GERADO


def test_exportar_duas_vezes_substitui_o_ficheiro(session: Session, tmp_path) -> None:
    """O ficheiro gerado é sempre o mesmo, atualizado — não se acumulam cópias."""
    importar(session, _tabela(_placa("19mm", "21.02")))
    session.commit()

    destino = tmp_path / "sub" / "gerado.xlsx"
    exportador.exportar(session, destino)
    exportador.exportar(session, destino)

    assert list(destino.parent.iterdir()) == [destino]
