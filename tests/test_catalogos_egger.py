"""O adaptador de EGGER: o unpivot das espessuras e o que ele não deixa passar.

O que estes testes guardam é sobretudo aquilo que **falha em silêncio** se não
for guardado: um cabeçalho escrito de outra maneira, uma coluna de espessura
que muda de palavra, um tipo de produto novo. Foi assim que 3 869 artigos
estiveram invisíveis na Pesquisa IA sem ninguém saber — o adaptador tem de
rebentar ou avisar, nunca devolver menos e calar-se.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.catalogos import egger
from app.services.catalogos.base import FormatoInesperado

NOTAS_BF = (
    "Tabela de Produtos EGGER Balbino & Faustino - 2026",
    "Base atualizada a partir do PDF BF-82 2026/04/20 (substitui tabela anterior de 2026/03/12).",
    "Mantido o separador para consulta.",
)
TIPO_BF = "Eurodekor Tableros de partículas revestidos E1E05 TSCA P2"
#: A WoodSide escreve o mesmo produto com o «El E05» que o OCR do PDF deixou.
TIPO_WOODSIDE = "Eurodekor Tableros de partículas revestidos El E05 TSCA P2"


def _escrever(
    caminho: Path,
    nome_folha: str,
    *,
    notas: tuple[str, ...] = NOTAS_BF,
    espessuras: tuple[str, ...] = ("8mm", "19mm"),
    linhas: list[list[object]],
) -> Path:
    """Escreve um separador com a mesma forma que os do ficheiro real.

    As ``linhas`` vêm como ``[ref, st, design, grupo, *(esp, preço) por
    espessura, tipo, fornecedor, obs]`` — sem o ``id``, que o adaptador ignora.
    """
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = nome_folha
    for nota in notas:
        worksheet.append([nota])

    cabecalho: list[str] = ["id", "Referência", "ST", "Nome Design", "Grupo"]
    for espessura in espessuras:
        cabecalho += [f"Esp {espessura}", f"Preço Tabela {espessura}"]
    cabecalho += ["Tipo Produto", "Fornecedor", "Observações"]
    worksheet.append(cabecalho)

    for indice, linha in enumerate(linhas, start=1):
        completa = list(linha)
        if len(completa) < len(cabecalho) - 1:
            completa += [None] * (len(cabecalho) - 1 - len(completa))
        worksheet.append([indice, *completa])

    workbook.save(caminho)
    return caminho


def _folha(nome: str = "Stock_B&F_Egger") -> egger.FolhaEgger:
    return egger.FolhaEgger(
        folha=nome, fornecedor="Balbino & Faustino", nome="EGGER Balbino & Faustino"
    )


LINHA_F037 = ["F037", "ST76", "Travertino Taormina", 8, "SIM", 17.98, "SIM", 21.02, TIPO_BF, "Balbino & Faustino", None]


# ---------------------------------------------------------------------------
# O unpivot
# ---------------------------------------------------------------------------


def test_cada_espessura_da_um_artigo(tmp_path: Path) -> None:
    """Duas colunas de espessura numa linha são dois artigos, não um."""
    caminho = _escrever(
        tmp_path / "t.xlsx", "Stock_B&F_Egger", linhas=[list(LINHA_F037)]
    )
    tabela = egger.ler_folha(caminho, _folha())

    assert len(tabela.artigos) == 2
    espessuras = sorted(a.espessura_mm for a in tabela.artigos)
    assert espessuras == [Decimal("8"), Decimal("19")]
    precos = {a.espessura_mm: a.preco for a in tabela.artigos}
    assert precos[Decimal("8")] == Decimal("17.98")
    assert precos[Decimal("19")] == Decimal("21.02")
    assert not tabela.avisos


def test_dez_espessuras_dao_dez_artigos(tmp_path: Path) -> None:
    """A WoodSide tem dez espessuras por decorativo — 209 linhas, 2 090 artigos."""
    espessuras = ("8mm", "10mm", "12mm", "16mm", "19mm", "22mm", "25mm", "28mm", "30mm", "38mm")
    pares: list[object] = []
    for i, _ in enumerate(espessuras):
        pares += ["SIM", 6.19 + i]
    caminho = _escrever(
        tmp_path / "w.xlsx",
        "Stock_WoodSide_Egger",
        espessuras=espessuras,
        linhas=[["W908", "SM", "Blanco tiza", 0, *pares, TIPO_WOODSIDE, "WoodSide", None]],
    )
    tabela = egger.ler_folha(
        caminho,
        egger.FolhaEgger(
            folha="Stock_WoodSide_Egger", fornecedor="WoodSide", nome="EGGER WoodSide"
        ),
    )

    assert len(tabela.artigos) == 10
    assert tabela.com_preco == 10
    assert sorted(a.espessura_mm for a in tabela.artigos)[-1] == Decimal("38")


def test_espessura_que_nao_existe_fica_de_fora(tmp_path: Path) -> None:
    """Um «NAO» na coluna Esp não gera artigo — e não gera aviso nenhum."""
    caminho = _escrever(
        tmp_path / "t.xlsx",
        "Stock_B&F_Egger",
        linhas=[["F037", "ST76", "Travertino", 8, "NAO", None, "SIM", 21.02, TIPO_BF, "Balbino & Faustino", None]],
    )
    tabela = egger.ler_folha(caminho, _folha())

    assert [a.espessura_mm for a in tabela.artigos] == [Decimal("19")]
    assert not tabela.avisos


def test_preco_sem_a_marca_sim_nao_se_perde(tmp_path: Path) -> None:
    """Nunca se descarta um preço por causa da coluna do lado — mas avisa-se."""
    caminho = _escrever(
        tmp_path / "t.xlsx",
        "Stock_B&F_Egger",
        linhas=[["F037", "ST76", "Travertino", 8, "NAO", 17.98, "SIM", 21.02, TIPO_BF, "Balbino & Faustino", None]],
    )
    tabela = egger.ler_folha(caminho, _folha())

    assert len(tabela.artigos) == 2
    assert any("tem preço mas a coluna Esp diz" in aviso for aviso in tabela.avisos)


def test_marcado_como_disponivel_e_sem_preco_entra_sem_preco(tmp_path: Path) -> None:
    """A tabela lista o artigo mas não lhe dá preço. Acontece, e diz-se."""
    caminho = _escrever(
        tmp_path / "t.xlsx",
        "Stock_B&F_Egger",
        linhas=[["F037", "ST76", "Travertino", 8, "SIM", None, "SIM", 21.02, TIPO_BF, "Balbino & Faustino", None]],
    )
    tabela = egger.ler_folha(caminho, _folha())

    assert len(tabela.artigos) == 2
    assert tabela.com_preco == 1
    assert any("sem preço" in aviso for aviso in tabela.avisos)


# ---------------------------------------------------------------------------
# A chave natural
# ---------------------------------------------------------------------------


def test_o_st_entra_na_chave_natural(tmp_path: Path) -> None:
    """O W908 existe em SM e em ST7: sem o ST, um apagava o outro."""
    caminho = _escrever(
        tmp_path / "t.xlsx",
        "Stock_B&F_Egger",
        linhas=[
            ["W908", "SM", "Blanco tiza", 0, "SIM", 6.19, "SIM", 8.74, TIPO_BF, "Balbino & Faustino", None],
            ["W908", "ST7", "Blanco tiza", 0, "SIM", 6.19, "SIM", 8.74, TIPO_BF, "Balbino & Faustino", None],
        ],
    )
    tabela = egger.ler_folha(caminho, _folha())

    chaves = {a.chave_natural for a in tabela.artigos}
    assert len(chaves) == 4
    assert "W908|SM|PB STD|8mm" in chaves
    assert "W908|ST7|PB STD|8mm" in chaves


# ---------------------------------------------------------------------------
# O substrato
# ---------------------------------------------------------------------------


def test_as_duas_grafias_do_eurodekor_dao_o_mesmo_substrato(tmp_path: Path) -> None:
    """«E1E05» e «El E05» são a mesma placa — o OCR é que as separou."""
    caminho = _escrever(
        tmp_path / "t.xlsx",
        "Stock_B&F_Egger",
        linhas=[
            ["F037", "ST76", "A", 8, "SIM", 17.98, "SIM", 21.02, TIPO_BF, "Balbino & Faustino", None],
            ["F038", "ST76", "B", 8, "SIM", 17.98, "SIM", 21.02, TIPO_WOODSIDE, "Balbino & Faustino", None],
        ],
    )
    tabela = egger.ler_folha(caminho, _folha())

    assert {a.substrato for a in tabela.artigos} == {"PB STD"}
    assert not tabela.avisos


def test_tipo_de_produto_desconhecido_avisa_e_nao_adivinha(tmp_path: Path) -> None:
    """Um substrato errado é pior do que nenhum: é dele que depende comparar."""
    caminho = _escrever(
        tmp_path / "t.xlsx",
        "Stock_B&F_Egger",
        linhas=[["F037", "ST76", "A", 8, "SIM", 17.98, "SIM", 21.02, "Coisa nova do catálogo", "Balbino & Faustino", None]],
    )
    tabela = egger.ler_folha(caminho, _folha())

    assert {a.substrato for a in tabela.artigos} == {None}
    assert any("sem substrato conhecido" in aviso for aviso in tabela.avisos)
    assert all(a.chave_natural.startswith("F037|ST76|") for a in tabela.artigos)


# ---------------------------------------------------------------------------
# A cabeça da tabela
# ---------------------------------------------------------------------------


def test_data_e_codigo_saem_das_notas_do_separador(tmp_path: Path) -> None:
    """A primeira data é a desta tabela; a segunda é a que ela substituiu."""
    caminho = _escrever(
        tmp_path / "t.xlsx", "Stock_B&F_Egger", linhas=[list(LINHA_F037)]
    )
    tabela = egger.ler_folha(caminho, _folha())

    assert tabela.data_tabela == date(2026, 4, 20)
    assert tabela.referencia_tabela == "BF-82"
    # O nome não leva o ano de propósito: é ele que identifica a tabela de uma
    # versão para a outra, e o ano está na data_tabela.
    assert tabela.nome == "EGGER Balbino & Faustino"
    assert tabela.fabricante == "EGGER"
    assert tabela.fornecedor == "Balbino & Faustino"
    assert tabela.unidade_preco == "M2"
    assert tabela.observacoes is not None and "BF-82" in tabela.observacoes


def test_uma_tabela_sem_codigo_nao_inventa_nenhum(tmp_path: Path) -> None:
    caminho = _escrever(
        tmp_path / "t.xlsx",
        "Stock_B&F_Egger",
        notas=("Tabela EGGER WoodSide", "A partir do PDF Eurodekor WoodSide 2026/03/11."),
        linhas=[list(LINHA_F037)],
    )
    tabela = egger.ler_folha(caminho, _folha())

    assert tabela.referencia_tabela is None
    assert tabela.data_tabela == date(2026, 3, 11)


def test_o_grupo_dez_continua_a_ser_dez(tmp_path: Path) -> None:
    """Um ``rstrip('.0')`` transformava o Grupo 10 em Grupo 1."""
    caminho = _escrever(
        tmp_path / "t.xlsx",
        "Stock_B&F_Egger",
        linhas=[["F037", "ST76", "A", 10, "SIM", 17.98, "SIM", 21.02, TIPO_BF, "Balbino & Faustino", None]],
    )
    tabela = egger.ler_folha(caminho, _folha())

    assert {a.grupo for a in tabela.artigos} == {"Grupo 10"}


def test_a_coluna_fornecedor_que_nao_bate_avisa(tmp_path: Path) -> None:
    caminho = _escrever(
        tmp_path / "t.xlsx",
        "Stock_B&F_Egger",
        linhas=[["F037", "ST76", "A", 8, "SIM", 17.98, "SIM", 21.02, TIPO_BF, "Outro Fornecedor", None]],
    )
    tabela = egger.ler_folha(caminho, _folha())

    assert any("mas este separador é do" in aviso for aviso in tabela.avisos)


# ---------------------------------------------------------------------------
# O hash
# ---------------------------------------------------------------------------


def test_o_hash_e_do_separador_e_nao_do_ficheiro(tmp_path: Path) -> None:
    """O mesmo xlsx traz treze separadores: um hash por ficheiro colidia."""
    caminho = tmp_path / "dois.xlsx"
    _escrever(caminho, "Stock_B&F_Egger", linhas=[list(LINHA_F037)])
    # acrescenta o segundo separador ao mesmo ficheiro
    from openpyxl import load_workbook

    workbook = load_workbook(caminho)
    folha2 = workbook.create_sheet("Stock_WoodSide_Egger")
    for nota in NOTAS_BF:
        folha2.append([nota])
    folha2.append(
        ["id", "Referência", "ST", "Nome Design", "Grupo", "Esp 8mm", "Preço Tabela 8mm", "Tipo Produto", "Fornecedor", "Observações"]
    )
    folha2.append([1, "W908", "SM", "Blanco tiza", 0, "SIM", 6.19, TIPO_WOODSIDE, "WoodSide", None])
    workbook.save(caminho)

    tabelas = egger.ler_tabelas(caminho)
    assert len(tabelas) == 2
    assert tabelas[0].ficheiro_hash != tabelas[1].ficheiro_hash
    assert tabelas[0].ficheiro_origem.endswith("#Stock_B&F_Egger")


def test_o_mesmo_conteudo_da_sempre_o_mesmo_hash(tmp_path: Path) -> None:
    """Sem isto, reimportar a mesma tabela criava-a outra vez."""
    a = _escrever(tmp_path / "a.xlsx", "Stock_B&F_Egger", linhas=[list(LINHA_F037)])
    b = _escrever(tmp_path / "b.xlsx", "Stock_B&F_Egger", linhas=[list(LINHA_F037)])

    assert egger.ler_folha(a, _folha()).ficheiro_hash == egger.ler_folha(
        b, _folha()
    ).ficheiro_hash


def test_mexer_num_preco_muda_o_hash(tmp_path: Path) -> None:
    linha_nova = list(LINHA_F037)
    linha_nova[5] = 18.50
    a = _escrever(tmp_path / "a.xlsx", "Stock_B&F_Egger", linhas=[list(LINHA_F037)])
    b = _escrever(tmp_path / "b.xlsx", "Stock_B&F_Egger", linhas=[linha_nova])

    assert egger.ler_folha(a, _folha()).ficheiro_hash != egger.ler_folha(
        b, _folha()
    ).ficheiro_hash


# ---------------------------------------------------------------------------
# O que rebenta, em vez de devolver zero em silêncio
# ---------------------------------------------------------------------------


def test_sem_cabecalho_reconhecivel_rebenta(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Stock_B&F_Egger"
    worksheet.append(["Código", "Acab", "Design", "Grupo", "Esp 8mm", "Preço Tabela 8mm"])
    worksheet.append(["F037", "ST76", "A", 8, "SIM", 17.98])
    caminho = tmp_path / "sem.xlsx"
    workbook.save(caminho)

    with pytest.raises(FormatoInesperado, match="cabeçalho"):
        egger.ler_folha(caminho, _folha())


def test_sem_colunas_de_espessura_rebenta(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Stock_B&F_Egger"
    worksheet.append(["id", "Referência", "ST", "Nome Design", "Grupo", "Preço Tabela"])
    worksheet.append([1, "F037", "ST76", "A", 8, 17.98])
    caminho = tmp_path / "sem_esp.xlsx"
    workbook.save(caminho)

    with pytest.raises(FormatoInesperado, match="espessura"):
        egger.ler_folha(caminho, _folha())


def test_separador_que_nao_existe_diz_quais_existem(tmp_path: Path) -> None:
    caminho = _escrever(tmp_path / "t.xlsx", "Outra_Coisa", linhas=[list(LINHA_F037)])

    with pytest.raises(FormatoInesperado, match="Outra_Coisa"):
        egger.ler_folha(caminho, _folha())


def test_cabecalho_sem_uma_unica_linha_rebenta(tmp_path: Path) -> None:
    caminho = _escrever(tmp_path / "t.xlsx", "Stock_B&F_Egger", linhas=[])

    with pytest.raises(FormatoInesperado, match="zero linhas"):
        egger.ler_folha(caminho, _folha())


def test_ficheiro_que_nao_existe_rebenta(tmp_path: Path) -> None:
    with pytest.raises(FormatoInesperado, match="não encontrado"):
        egger.ler_folha(tmp_path / "nao_existe.xlsx", _folha())
