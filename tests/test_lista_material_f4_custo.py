"""F4 da Lista Material: preço das ferragens V3 → PHC → IMOS e custo com rigor só no fim.

Preços reais da obra 1610 (PHC, 21-09-2026): dobradiça FF00060 2,59 €/UN;
parafuso FF00036 1,267 € por 100 («%»); corrediça FF00017 5,268 €/PAR.
"""
from __future__ import annotations

from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from PySide6.QtWidgets import QApplication

from app.services import analise_lista_material_service as svc
from app.services import custo_ferragens_service as custo
from app.ui.dialogs import analise_lista_material_dialog as ui
from app.ui.dialogs.mapear_ferragens_dialog import MapearFerragensDialog

PHC = {
    "FF00060": {"Ref": "FF00060", "Descricao": "DOBRADICA BLUM 75B1550", "Preco_Custo": 2.59,
                "Preco_Ultimo": 2.35, "Unidade": "UN", "Data_Preco": "14.09.2026"},
    "FF00036": {"Ref": "FF00036", "Descricao": "PARAFUSO 3,5X35", "Preco_Custo": 1.267,
                "Preco_Ultimo": 1.7, "Unidade": "%", "Data_Preco": "14.09.2026"},
    "FF00017": {"Ref": "FF00017", "Descricao": "CORREDIÇA", "Preco_Custo": 5.268,
                "Preco_Ultimo": 6.9, "Unidade": "PAR", "Data_Preco": "02.09.2026"},
    "FF01499": {"Ref": "FF01499", "Descricao": "PE BONE", "Preco_Custo": 0,
                "Preco_Ultimo": 0.154, "Unidade": "UN", "Data_Preco": ""},
}


def _line(ref, qt=10, imos="", kind="Ferragens"):
    return svc.cost_line(kind, f"ferragem:{ref}", ref, qt, "un", ref_phc=ref, imos_price=imos)


def test_phc_converte_a_centena_e_pede_confirmacao_no_par():
    dobradica = custo.preco_phc(_line("FF00060"), PHC["FF00060"])
    assert dobradica["net"] == "2.59" and dobradica["fonte"] == "PHC" and not dobradica["confirmar"]
    parafuso = custo.preco_phc(_line("FF00036"), PHC["FF00036"])
    assert Decimal(parafuso["net"]) == Decimal("0.01267")
    corredica = custo.preco_phc(_line("FF00017"), PHC["FF00017"])
    assert "PAR" in corredica["confirmar"]
    cavilha = {"Ref": "FC00304", "Descricao": "CAVILHA (10KG)", "Preco_Custo": 7.34, "Preco_Ultimo": 0, "Unidade": "KG"}
    assert "KG" in custo.preco_phc(_line("FC00304"), cavilha)["confirmar"]
    # À unidade, mas 20× o IMOS: também não entra sozinho.
    longe = custo.preco_phc(_line("FF00060", imos="0.10"), PHC["FF00060"])
    assert "longe do IMOS" in longe["confirmar"]
    # Sem preço de custo usa o último preço de compra.
    assert custo.preco_phc(_line("FF01499"), PHC["FF01499"])["net"] == "0.154"
    assert custo.preco_phc(_line("X"), None) is None


def test_ordem_v3_phc_imos():
    v3 = {"id": 7, "net": "2.10", "unit": "un"}
    assert custo.resolver_ferragem(_line("FF00060", imos="2.47"), v3, PHC) is v3
    assert custo.resolver_ferragem(_line("FF00060", imos="2.47"), None, PHC)["fonte"] == "PHC"
    imos = custo.resolver_ferragem(_line("FF09999", imos="2.47"), None, PHC)
    assert imos["fonte"] == "IMOS" and "PROVISÓRIO" in imos["mapping_source"]
    assert custo.resolver_ferragem(_line("FF09999"), None, PHC) is None
    # A cavilha ao quilo nunca entra pelo PHC: fica o IMOS provisório.
    cavilha = {"FC00304": {"Ref": "FC00304", "Preco_Custo": 7.34, "Unidade": "KG"}}
    assert custo.resolver_ferragem(_line("FC00304", imos="0.01"), None, cavilha)["fonte"] == "IMOS"
    assert custo.resolver_ferragem(_line("FF00017"), None, PHC) is None
    assert custo.resolver_ferragem(_line("FF00060", kind="Orlas"), None, PHC) is None
    cost, _ = svc.calculate_cost(_line("FF00036", qt=400), custo.preco_phc(_line("FF00036"), PHC["FF00036"]))
    assert cost == Decimal("5.068")


def test_consulta_phc_so_aceita_referencias_no_formato():
    assert custo.refs_validas(["ff00060", "FF00060", "?? (?)", "FF1'; DROP", "", None]) == ["FF00060"]


def _plan():
    return [{"name": "p", "version": "v", "total": "10.5", "materials": []}]


def test_custo_so_e_final_com_a_obra_fechada_e_tudo_apurado():
    lines = [svc.cost_line("Placas", "placa:A", "A", 10, "m2"), _line("FF00060")]
    precos = {"placa:A": {"id": 1, "net": "5"}, "ferragem:FF00060": {"id": 2, "net": "2"}}
    tempos = {"sectors": [{"state": "Concluído"}, {"state": "Não aplicável"}]}
    final = custo.estado_do_custo("Finalizado", lines, precos, _plan(), tempos)
    assert final.final and "FINAL" in final.titulo
    em_producao = custo.estado_do_custo("Producao", lines, precos, _plan(), tempos)
    assert not em_producao.final and "PROVISÓRIO" in em_producao.titulo
    imos = dict(precos, **{"ferragem:FF00060": {"net": "2", "fonte": "IMOS"}})
    assert not custo.estado_do_custo("Arquivado", lines, imos, _plan(), tempos).final
    sem_plano = custo.estado_do_custo("Finalizado", lines, precos, [], tempos)
    assert any(not ok and "sem plano de corte" in texto for ok, texto in sem_plano.pontos)
    sem_tempos = custo.estado_do_custo("Finalizado", lines, precos, _plan(), {"sectors": [{"state": "Em curso / estado por confirmar"}]})
    assert not sem_tempos.final


@pytest.fixture
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "Lista_Material_1610_01_26_JF_VIVA.xlsx"
    w = Workbook()
    w.active.title = "LISTAGEM_CUT_RITE"
    w.active.append(["Title"])
    w.active.append(["Material", "Descricao", "Qt", "Esp"])
    w.active.append(["AGL_MLM_BRANCO_19MM", "TETO", 2, 19])
    s = w.create_sheet("5_Custo_Obra_Ferragens")
    s.append(["Nome iMos (Nome Uniao)", "Jogo de Unioes (iMos)", "Descricao", "Ref PHC", "Ref Fornecedor",
              "Fornecedor", "Na lista", "Comp", "Larg", "Esp", "Qt", "Un", "€ / un", "€"])
    s.append(["FERRAGENS"] + [""] * 13)
    s.append(["BL_DOB", "", "Dobradiça", "FF00060", "75B1550 BLUM", "BLUM", "", "", "", "", 95, "un", "2,47", 234.65])
    s.append(["PRF", "", "Parafuso", "FF00036", "", "HAFELE", "", "", "", "", 400, "un", "0,03", 12])
    s.append(["CANTO", "", "Canto", "??  (?)", "", "ARTIMOL", "", "", "", "", 1, "un", "1,10", 1.1])
    # A origem das ferragens é o 1_FERRAGENS (editado à mão), não o 5_Custo.
    f = w.create_sheet("1_FERRAGENS")
    f.append(["Listagem Ferragem"])
    f.append(["Imagem\n", "Ref PHC", None, "Ref Fornecedor", "Descrição 1", None, None, "Qt.", "UN", "Artg.\n"])
    f.append([None, "FF00060", None, "75B1550 BLUM", "Dobradiça\nCOR -> nickle", None, None, 95, "\n", "RP_A_03"])
    f.append([None, "FF00036", None, "016.10.699", "Parafuso", None, None, 400, "\n", "RP_A_04"])
    f.append([None, "??", None, "??", "Canto", None, None, 1, "\n", "LV_A_01"])
    w.save(path)
    return path


def _dialog(session, workbook, monkeypatch, estado="Producao"):
    monkeypatch.setattr(ui, "query_woodstore", lambda _: [])
    monkeypatch.setattr(ui.custo_ferragens, "ler_precos_phc", lambda _s, refs: {r: PHC[r] for r in refs if r in PHC})
    return ui.AnaliseListaMaterialDialog(
        session, workbook_path=workbook, plan_name="1610_01_01_26_JF_VIVA", cutrite_folder=workbook.parent,
        user=SimpleNamespace(role="admin", username="paulo"), obra_info={"estado": estado})


def test_analise_preenche_phc_e_imos_e_diz_que_e_provisorio(app, session, workbook, monkeypatch):
    dialog = _dialog(session, workbook, monkeypatch)
    try:
        fontes = {l["name"]: custo.fonte(dialog.prices.get(l["key"])) for l in dialog.lines if l["kind"] == "Ferragens"}
        assert fontes == {"Dobradiça": "PHC", "Parafuso": "PHC", "Canto": "IMOS"}
        assert "PROVISÓRIO" in dialog.rigor_label.text()
        assert "Ferragens: V3 0 · PHC 2 · IMOS provisório 1" in dialog.rigor_label.text()
        dialog._update_prices()   # não pode perder os preços PHC/IMOS (não têm id)
        assert {custo.fonte(dialog.prices.get(l["key"])) for l in dialog.lines if l["kind"] == "Ferragens"} == {"PHC", "IMOS"}
    finally:
        dialog.close()


def test_passo_a_passo_usa_phc_imos_e_salta(app, session, workbook, monkeypatch):
    dialog = _dialog(session, workbook, monkeypatch)
    try:
        pending = [l for l in dialog.lines if l["kind"] == "Ferragens"]
        prices = {}
        wizard = MapearFerragensDialog(pending, prices, [], dialog.phc, on_v3=lambda *_: None)
        assert wizard.progress.text().startswith("1 de 3")
        assert "2,59 €" in wizard.phc_button.text()
        wizard._use_phc()
        assert wizard.index == 1
        wizard._go(1)                     # saltar o parafuso
        assert not wizard.phc_button.isEnabled()   # «??  (?)» não é ref PHC
        wizard._use_imos()
        assert wizard.changed == 2
        assert {custo.fonte(p) for p in prices.values()} == {"PHC", "IMOS"}
        assert all(p["escolha_manual"] for p in prices.values())
    finally:
        dialog.close()


# ---- Origem nos separadores 1_FERRAGENS / 2_PURCH / 3_SPP (pedido de 21-09-2026) ----

def _livro_separadores(com_purch=True, com_custo=True):
    w = Workbook()
    f = w.active
    f.title = "1_FERRAGENS"
    f.append(["Listagem Ferragem"])
    f.append(["Imagem\n", "Ref PHC", None, "Ref Fornecedor", "Descrição 1", None, None, "Qt.", "UN", "Artg.\n"])
    f.append([None, "FF00043", None, "80395.07", "Suporte Prateleira", None, None, 128, "\n", "LV_A_02"])
    f.append([None, "FF00043 ", None, "80395.07", "Suporte Prateleira", None, None, 20, "\n", "RP_A_01"])
    f.append([None, "80593813", None, "Calçeiro", "Calçeiro Extraivel 11 | 805.93.813", None, None, 1, "\n", "RP_A_05(b)"])
    s = w.create_sheet("3_SPP")
    s.append(["Listagem Ferragem"])
    s.append(["Imagem", "Ref PHC", None, "Ref Fornecedor", "Descrição 1", None, None, "Qt.\n", None, None, "Comp.", None, "Artg."])
    s.append([None, "FF00412     ", None, "", "Varao Oval 30X15\nCOR -> Aluminio\n[990 mm]x[15 mm]x[30 mm]", None, None, 2, None, None, "989\n", None, "RP_A_04"])
    s.append([None, "FF00412", None, "", "Varao Oval 30X15\n[468 mm]x[15 mm]x[30 mm]", None, None, 1, None, None, None, None, "RP_A_01"])
    if com_purch:
        p = w.create_sheet("2_PURCH")
        p.append(["Listagem Ferragem"])
        p.append(["Imagem", None, "Acessorio #", "Comp X Larg X Esp", None, None, None, None, "Qt.", None, "Artg."])
        p.append([None, None, "", "600 X 610 X 860", None, None, None, None, 1, None, "13"])
    if com_custo:
        c = w.create_sheet("5_Custo_Obra_Ferragens")
        c.append(["Nome iMos (Nome Uniao)", "Jogo de Unioes (iMos)", "Descricao", "Ref PHC", "Ref Fornecedor",
                  "Fornecedor", "Na lista", "Comp", "Larg", "Esp", "Qt", "Un", "€ / un", "€"])
        c.append(["FERRAGENS"])
        c.append(["SUP", "", "Suporte Prateleira", "FF00043", "", "EMUCA", "", "", "", "", 999, "un", "0,15", 0])
        c.append(["CAVILHA_08p1X30_COLA", "", "Cavilha 8 X 30", "FC00304", "", "ARTIMOL", "fora", "", "", "", 606, "un", "0,01", 6.06])
        c.append(["MAQ", "", "Maquina lavar", "  (?)", "", "", "", "", "", "", 2, "un", "0", 0])
    return w


def test_ferragens_saem_dos_separadores_editados():
    linhas, avisos = custo.linhas_dos_separadores(_livro_separadores())
    por_nome = {l["name"]: l for l in linhas}
    suporte = por_nome["Suporte Prateleira"]
    # A mesma Ref PHC em duas linhas soma; a quantidade é a do separador (148), não a do IMOS (999).
    assert suporte["quantity"] == "148" and suporte["source_sheet"] == "1_FERRAGENS"
    assert suporte["imos_price"] == "0.15" and suporte["articles"] == "LV_A_02, RP_A_01"
    # SPP em metros: 2 × 989 mm + 1 × 468 mm (do texto, quando falta o Comp.).
    assert por_nome["Varao Oval 30X15"]["quantity"] == "2.446" and por_nome["Varao Oval 30X15"]["unit"] == "ml"
    assert por_nome["Calçeiro Extraivel 11 | 805.93.813"]["ref_phc"] == ""   # 80593813 não é Ref PHC
    assert por_nome["Objeto comprado 600 X 610 X 860"]["kind"] == "Comprados"
    # As cavilhas «fora da lista» contam; a máquina do 5_Custo nunca entra por lá.
    assert por_nome["CAVILHA_08p1X30_COLA"]["quantity"] == "606"
    assert "MAQ" not in por_nome and "Maquina lavar" not in por_nome
    assert any("1_FERRAGENS (3), 2_PURCH (1), 3_SPP (2)" in a for a in avisos)


def test_2_purch_apagado_deixa_de_contar_e_sem_separadores_nao_ha_ferragens():
    linhas, _ = custo.linhas_dos_separadores(_livro_separadores(com_purch=False))
    assert not [l for l in linhas if l["kind"] == "Comprados"]
    vazio = Workbook()
    linhas, avisos = custo.linhas_dos_separadores(vazio)
    assert linhas == [] and any("ferragens sem custo" in a for a in avisos)
    _, avisos = custo.linhas_dos_separadores(_livro_separadores(com_custo=False))
    assert any("Sem 5_Custo_Obra_Ferragens" in a for a in avisos)


def test_mapeamentos_antigos_reaproveitados_pela_ref_phc():
    from app.services import analise_custo_mapeamento_service as maps
    linha = svc.cost_line("Ferragens", "ferragem:Ferragens:FF00132:un", "MAELGA TIC-TAC", 22, "un",
                          ref_phc="FF00132", source_sheet="1_FERRAGENS")
    assert maps.mapping_identity(linha) == ["Ferragens", "REF", "FF00132", "un"]
    antigo = {"k": {"identity": ["Ferragens", "FECHO_TIC_TAC", "un", "X", "FF00132"], "mp_id": 5}}
    mp = SimpleNamespace(id=5, ref_le="FER1", ref_phc="", nome_imos="", descricao="Tic-tac",
                         unidade="un", preco_liquido=Decimal("2"))
    price = maps.resolve_price(linha, [mp], antigo)
    assert price["id"] == 5 and price["net"] == "2"


def test_nao_contabilizar_conta_como_resolvido():
    linha = _line("FF09999")
    precos = {linha["key"]: custo.preco_excluido(linha)}
    assert svc.calculate_cost(linha, precos[linha["key"]])[0] == 0
    estado = custo.estado_do_custo("Finalizado", [linha], precos, _plan(), {"sectors": [{"state": "Concluído"}]})
    assert estado.final
    assert any("não consideradas nesta obra 1" in texto for _, texto in estado.pontos)


def test_conciliacao_com_o_imos_diz_o_que_foi_retirado_acrescentado_e_fora():
    linhas, avisos = custo.linhas_dos_separadores(_livro_separadores(com_purch=False))
    texto = " | ".join(avisos)
    assert "retirado à mão, não conta): MAQ" in texto            # máquina de lavar
    assert "Fora da lista mas contam" in texto and "CAVILHA_08p1X30_COLA" in texto
    # O calceiro não está no 5_Custo deste livro: foi acrescentado à mão.
    assert "acrescentado ou mudado à mão, conta): " in texto and "Calçeiro" in texto


def test_nao_considerar_nesta_obra_nao_mexe_no_preco_nem_no_mapeamento(app, session, workbook, monkeypatch):
    dialog = _dialog(session, workbook, monkeypatch)
    try:
        linha = next(l for l in dialog.lines if l["name"] == "Dobradiça")
        preco = dict(dialog.prices[linha["key"]])
        antes = dialog.cost_summary.text()
        row = dialog.lines.index(linha)
        dialog.cost_table.selectRow(row)
        dialog._toggle_excluded()
        assert linha["key"] in dialog.excluded
        assert dialog.prices[linha["key"]] == preco                 # o preço fica
        assert dialog._effective_prices()[linha["key"]]["net"] == "0"
        assert "NÃO CONSIDERADA NESTA OBRA" in dialog.cost_table.item(row, 10).text()
        assert "não consideradas nesta obra 1" in dialog.rigor_label.text()
        assert antes != dialog.cost_summary.text()
        dialog._save()
        dialog._reload()                                            # fica gravado na análise desta obra
        assert linha["key"] in dialog.excluded
        dialog.cost_table.selectRow(row)
        dialog._toggle_excluded()                                   # e volta a contar
        assert linha["key"] not in dialog.excluded
    finally:
        dialog.close()


def test_passo_a_passo_nao_considerar_nesta_obra(app, session, workbook, monkeypatch):
    dialog = _dialog(session, workbook, monkeypatch)
    try:
        pending = [l for l in dialog.lines if l["kind"] == "Ferragens"]
        prices, excluded = {}, set()
        wizard = MapearFerragensDialog(pending, prices, [], dialog.phc, on_v3=lambda *_: None, excluded=excluded)
        wizard._exclude()
        assert excluded == {pending[0]["key"]} and prices == {}
    finally:
        dialog.close()
