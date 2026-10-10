"""Menu IMOS IX › Traduções do iX: página, janelas e aviso ao entrar."""

from __future__ import annotations

import inspect
from datetime import date, datetime
from pathlib import Path
from time import monotonic, sleep

import pytest
from openpyxl import Workbook
from PySide6.QtWidgets import QApplication, QDialog, QWidget
from sqlalchemy import BigInteger, create_engine
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app.db.base import Base
from app.services import imos_traducoes_service as servico
from app.services.permission_service import (
    DEFAULT_USER_PERMISSIONS,
    DESCRICOES_ACESSOS,
    MENU_PERMISSIONS,
)
from app.ui.dialogs import imos_traducoes_dialogs as janelas
from app.ui.helpers import traducoes_imos as config
from app.ui.pages import imos_ix_page as pagina_mod

BOM = b"\xef\xbb\xbf"
MSG = "PTG;7134\t;Comissão\r\nPTG;1220\t;Descrição\r\nPTG;32250 ;Largura da Tag\r\n"


@compiles(BigInteger, "sqlite")
def _bigint(type_, compiler, **kw):  # noqa: ANN001
    return "INTEGER"


@pytest.fixture(scope="module")
def qapp():
    return QApplication.instance() or QApplication([])


@pytest.fixture()
def ambiente(qapp, tmp_path, monkeypatch):
    """Um iX de mentira: imos.msg + Excel numa pasta temporária, programas fechados."""
    pasta = tmp_path / "iX CAD 2025" / "BIN" / "MSG"
    pasta.mkdir(parents=True)
    msg = pasta / "imos.msg"
    msg.write_bytes(BOM + MSG.encode("utf-8"))
    livro = Workbook()
    for linha, (ref, texto) in enumerate(
        [("PTG;7134", "Enc PHC:"), ("PTG;1220", "Materiais Usados"), ("PTG;32250", "Largura da Tag"),
         ("PTG;99", "Não existe")],
        start=6,
    ):
        livro.active.cell(linha, 2, ref)
        livro.active.cell(linha, 3, texto)
    excel = tmp_path / "imos_msg.xlsx"
    livro.save(excel)

    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    fabrica = sessionmaker(bind=engine)
    monkeypatch.setattr(pagina_mod, "SessionLocal", fabrica)
    monkeypatch.setattr(config, "caminho_escolhido_neste_pc", lambda: "")
    monkeypatch.setattr(config, "caminho_excel", lambda _s: str(excel))
    monkeypatch.setattr(servico, "localizar_imos_msg", lambda **_k: msg)
    monkeypatch.setattr(config, "caminho_imos_msg", lambda: (msg, False))
    programas = {"iX CAD": False, "iX Organizer": False}
    monkeypatch.setattr(servico, "estado_programas", lambda *_a: dict(programas))
    monkeypatch.setattr(servico, "_executaveis_a_correr",
                        lambda: {exe for exe, nome in servico.PROGRAMAS_IX.items() if programas[nome]})
    for modulo in (pagina_mod, janelas):
        monkeypatch.setattr(modulo.diario_bordo, "registar_acao", lambda *_a, **_k: None)
    yield {"msg": msg, "excel": excel, "programas": programas, "sessao": fabrica}
    engine.dispose()


def _aba(ambiente) -> pagina_mod.TraducoesIxAba:
    pagina = pagina_mod.ImosIxPage(user_id=2)
    pagina.carregar()
    aba = pagina.traducoes
    aba._pagina_do_teste = pagina  # a página (dona da aba) vive enquanto o teste a usar
    return aba


def test_aba_mostra_o_que_vai_mudar(ambiente) -> None:
    aba = _aba(ambiente)
    assert aba.msg_edit.text() == str(ambiente["msg"])
    assert aba.msg_edit.isReadOnly()
    assert "iX CAD 2025" in aba.msg_nota.text()
    assert aba.table.rowCount() == 4
    estados = {aba.table.item(l, 0).text(): aba.table.item(l, 3).text() for l in range(4)}
    assert estados == {
        "PTG;7134": "Vai mudar",
        "PTG;1220": "Vai mudar",
        "PTG;32250": "Já está",
        "PTG;99": "Não existe no iX",
    }
    assert aba.table.item(0, 1).text() == "Comissão"
    # O que vai mudar fica em cima, depois o que não existe, por fim o que já está.
    assert [aba.table.item(l, 3).text() for l in range(4)] == [
        "Vai mudar", "Vai mudar", "Não existe no iX", "Já está",
    ]
    assert aba.ficha_por_aplicar.text() == "2 por aplicar"
    assert "2 tradução(ões) por aplicar" in aba.status_label.text()
    assert aba.aplicar_button.isEnabled()
    assert not aba.excel_procurar_button.isEnabled()  # só o administrador muda o Excel
    for botao in (aba.verificar_button, aba.aplicar_button, aba.repor_button, aba.msg_procurar_button):
        assert botao.toolTip()


def test_com_o_ix_aberto_o_supervisor_pede_para_fechar(ambiente) -> None:
    ambiente["programas"]["iX CAD"] = True
    aba = _aba(ambiente)
    assert "iX CAD: aberto" in aba.ficha_cad.text()
    assert "Feche o iX CAD antes de aplicar" in aba.status_label.text()


def test_sem_ix_no_pc(ambiente, monkeypatch) -> None:
    monkeypatch.setattr(config, "caminho_imos_msg", lambda: (None, False))
    aba = _aba(ambiente)
    assert "Não encontrei o iX CAD" in aba.msg_nota.text()
    assert not aba.aplicar_button.isEnabled() and not aba.repor_button.isEnabled()
    assert aba.table.rowCount() == 4  # o Excel continua à vista


def test_aplicar_pela_aba(ambiente, monkeypatch) -> None:
    aba = _aba(ambiente)
    monkeypatch.setattr(pagina_mod, "garantir_programas_fechados", lambda _p: True)

    class _Janela:
        def __init__(self, msg, lista, _parent):
            self.resultado = servico.aplicar(msg, lista, listar_processos=set)
            self.erro = ""

        def exec(self):
            return 1

    monkeypatch.setattr(pagina_mod, "AplicarTraducoesDialog", _Janela)
    aba.aplicar()
    assert "PTG;7134\t;Enc PHC:" in ambiente["msg"].read_text(encoding="utf-8-sig")
    assert aba.ficha_por_aplicar.text() == "0 por aplicar"
    assert "2 tradução(ões) aplicada(s)" in aba.status_label.text()
    assert len(servico.listar_copias(ambiente["msg"])) == 1
    # Outra vez: nada para aplicar, nem cópia nova.
    aba.aplicar()
    assert "Nada para aplicar" in aba.status_label.text()
    assert len(servico.listar_copias(ambiente["msg"])) == 1


def test_cancelar_no_fechar_programas_nao_mexe(ambiente, monkeypatch) -> None:
    aba = _aba(ambiente)
    original = ambiente["msg"].read_bytes()
    monkeypatch.setattr(pagina_mod, "garantir_programas_fechados", lambda _p: False)
    aba.aplicar()
    assert ambiente["msg"].read_bytes() == original
    assert "Cancelado" in aba.status_label.text()


def test_janela_de_aplicar_mostra_linha_a_linha(ambiente) -> None:
    lista = servico.ler_excel(ambiente["excel"])
    janela = janelas.AplicarTraducoesDialog(ambiente["msg"], lista)
    limite = monotonic() + 5
    while janela.resultado is None and not janela.erro and monotonic() < limite:
        QApplication.processEvents()
        sleep(0.01)
    QApplication.processEvents()
    registo = janela.registo.toPlainText()
    assert janela.resultado is not None, janela.erro
    assert "Cópia de segurança feita: imos.msg.copia_" in registo
    assert "«Comissão»  →  «Enc PHC:»" in registo
    assert janela.fechar_button.isEnabled() and janela.pasta_button.isVisibleTo(janela)
    assert "2 alterada(s)" in janela.resumo_label.text()
    assert "1 não existem no imos.msg (PTG;99)" in janela.resumo_label.text()
    janela.accept()


def test_janela_fechar_programas_so_deixa_continuar_com_tudo_fechado(qapp) -> None:
    estado = {"iX CAD": True, "iX Organizer": False}
    janela = janelas.FecharProgramasDialog(estado=lambda: dict(estado))
    assert not janela.continuar_button.isEnabled()
    assert "À espera que feche: iX CAD" in janela.status_label.text()
    estado["iX CAD"] = False
    assert janela.verificar()
    assert janela.continuar_button.isEnabled()
    janela.reject()
    assert janelas.garantir_programas_fechados(None, estado=lambda: {"iX CAD": False})


def test_repor_copia_pela_janela(ambiente, monkeypatch) -> None:
    original = ambiente["msg"].read_bytes()
    servico.aplicar(ambiente["msg"], servico.ler_excel(ambiente["excel"]), listar_processos=set)
    janela = janelas.ReporCopiaDialog(ambiente["msg"])
    assert janela.tabela.rowCount() == 1
    monkeypatch.setattr(janelas.QMessageBox, "question",
                        lambda *_a, **_k: janelas.QMessageBox.StandardButton.Yes)
    monkeypatch.setattr(janelas.QMessageBox, "information", lambda *_a, **_k: None)
    janela._repor()
    assert ambiente["msg"].read_bytes() == original
    assert janela.result() == QDialog.DialogCode.Accepted
    assert len(servico.listar_copias(ambiente["msg"])) == 2  # a de antes + a do «antes de repor»


def test_preferencia_do_aviso_fica_na_conta(ambiente) -> None:
    aba = _aba(ambiente)
    assert aba.aviso_check.isChecked()
    aba.aviso_check.setChecked(False)
    with ambiente["sessao"]() as s:
        assert not config.aviso_ligado(s, 2)
        assert config.aviso_ligado(s, 3)  # os outros continuam com o padrão
    assert "Aviso desligado" in aba.status_label.text()


# ---- aviso ao entrar ----------------------------------------------------------------------
def test_aviso_so_pergunta_quando_ha_traducoes_por_aplicar(qapp) -> None:
    abertos, perguntas = [], []

    def perguntar(_janela, estado):
        perguntas.append(estado.por_aplicar)
        return True

    aviso = config.AvisoTraducoesImos(
        QWidget(), ativo=False, user_id=2, abrir_menu=lambda: abertos.append(1),
        perguntar=perguntar, guardar_ultimo=lambda _d: None,
    )
    aviso._ao_ler(config.EstadoAviso(por_aplicar=0))
    aviso._ao_ler(config.EstadoAviso(ligado=False))
    aviso._ao_ler(config.EstadoAviso(problema="Sem I:"))
    assert perguntas == [] and abertos == []
    aviso._ao_ler(config.EstadoAviso(por_aplicar=3, versao="iX CAD 2025"))
    assert perguntas == [3] and abertos == [1]


def test_aviso_uma_vez_por_dia_nos_dias_uteis(qapp) -> None:
    lidos, guardados = [], []
    aviso = config.AvisoTraducoesImos(
        QWidget(), ativo=True, user_id=2, abrir_menu=lambda: None,
        ler_estado=lambda: lidos.append(1) or config.EstadoAviso(por_aplicar=0),
        ler_ultimo=lambda: None, guardar_ultimo=guardados.append,
        agora=lambda: datetime(2026, 10, 12, 9, 0),  # segunda-feira
    )
    try:
        aviso.verificar_se_e_hora()
        limite = monotonic() + 3
        while not guardados and monotonic() < limite:
            QApplication.processEvents()
            sleep(0.01)
        assert lidos == [1] and guardados == [date(2026, 10, 12)]
    finally:
        aviso.parar()
    sabado = config.AvisoTraducoesImos(
        QWidget(), ativo=False, user_id=2, abrir_menu=lambda: None,
        agora=lambda: datetime(2026, 10, 10, 9, 0),
    )
    sabado.ativo = True
    sabado.verificar_se_e_hora()
    assert not sabado._a_ler


def test_caminho_escolhido_que_ja_nao_existe_e_ignorado(tmp_path) -> None:
    existe = tmp_path / "imos.msg"
    existe.write_bytes(b"x")
    assert config.caminho_imos_msg(escolhido=lambda: str(existe), localizar=lambda: None) == (existe, True)
    assert config.caminho_imos_msg(
        escolhido=lambda: str(tmp_path / "foi_apagado.msg"), localizar=lambda: Path("auto")
    ) == (Path("auto"), False)


# ---- menu e acessos -------------------------------------------------------------------------
def test_menu_nasce_desligado_e_tem_descricao() -> None:
    assert "menu.imos_ix" in MENU_PERMISSIONS
    assert DEFAULT_USER_PERMISSIONS["menu.imos_ix"] is False
    assert "imos.msg" in DESCRICOES_ACESSOS["menu.imos_ix"].o_que_faz


def test_janela_principal_liga_o_menu_com_icone_e_o_aviso() -> None:
    from app.ui.main_window import MainWindow

    fonte = inspect.getsource(MainWindow)
    assert '_criar_item("IMOS IX", "imos_ix")' in fonte
    assert 'icone_imagem("imos_ix.png")' in fonte
    assert '"imos_ix": "menu.imos_ix"' in fonte
    assert "AvisoTraducoesImos(" in fonte
    icone = Path(pagina_mod.__file__).parents[1] / "assets" / "icons" / "imos_ix.png"
    assert icone.is_file()


def test_ficheiro_escolhido_a_mao_pode_voltar_ao_automatico(ambiente, monkeypatch) -> None:
    guardados = []
    monkeypatch.setattr(config, "caminho_imos_msg", lambda: (ambiente["msg"], True))
    monkeypatch.setattr(config, "guardar_caminho_neste_pc", guardados.append)
    aba = _aba(ambiente)
    assert "Escolhido à mão" in aba.msg_nota.text()
    assert aba.msg_auto_button.isVisibleTo(aba) and aba.msg_auto_button.toolTip()
    monkeypatch.setattr(config, "caminho_imos_msg", lambda: (ambiente["msg"], False))
    aba.msg_auto_button.click()
    assert guardados == [""]
    assert not aba.msg_auto_button.isVisibleTo(aba)
