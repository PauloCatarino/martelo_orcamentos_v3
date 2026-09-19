from __future__ import annotations

import inspect
from datetime import date
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.services.lista_material_acabamentos_service import (
    PlanoLacagem,
    _area_impressao,
    _configurar_pagina,
    analisar_acabamentos,
    palavra_de_acabamento,
    texto_rodape,
    tipo_acabamento,
)
from app.services.lista_material_pdf_service import (
    ACABAMENTOS_DOCUMENT_ID,
    DEFAULT_DOCUMENTS,
    document_filename,
    inspect_pdf_documents,
)


@pytest.mark.parametrize(
    "texto",
    [
        "RASGO ALHETA 3MM CF + 1LARG D_INTERROMPIDO_LACAR 9001",
        "ESTICADOR_PUX FRESADO_LACAR 9001 2F",
        "Lacagem mate",
        "LAKAGEM RAL 9010",
        "c/ envernisamento",
        "VERNIS incolor",
        "verniz",
        "ACABMENTO especial",
        "pintrua cliente",
        "Pintado a branco",
        "velatura carvalho",
        "LACADO",
        "cor RAL 7016",
        "Envernizar 2 faces",
    ],
)
def test_notas_com_acabamento_sao_detetadas_mesmo_mal_escritas(texto) -> None:
    assert palavra_de_acabamento(texto)


@pytest.mark.parametrize(
    "texto",
    [
        "S/ LACAGEM",
        "sem verniz",
        "NÃO LACAR",
        "ESTÁ EM OBRA",
        "RASGO CALHA LED CLIENTE_17X6",
        "RASGO ALHETA ORLA 1 COMP",
        "PLACA INTEIRA",
        "LATERAL",
        "PINO 8MM",
        "PVC_1.0_COR",
        "",
        None,
    ],
)
def test_notas_sem_acabamento_nao_dao_aviso(texto) -> None:
    assert palavra_de_acabamento(texto) == ""


def test_tipo_de_lacagem_aceita_maiusculas_e_espacos() -> None:
    assert tipo_acabamento("A-1 Face + Topos") == "A-1 Face + Topos"
    assert tipo_acabamento(" C-1 face+topos+50mm  contraface ") == (
        "c-1 Face + Topos + 50mm Contraface"
    )
    assert tipo_acabamento("E- Outro") == "E- Outro"
    assert tipo_acabamento("0") == ""
    assert tipo_acabamento("LACAR") == ""
    assert tipo_acabamento(None) == ""


def _lista_material(path, *, pecas, linhas_lacagem=6, cabecalho=True, observacoes=None):
    """Excel mínimo com a estrutura do modelo Lista_Material_IMOS_MARTELO."""
    workbook = Workbook()
    definicoes = workbook.active
    definicoes.title = "DEFENICOES"
    definicoes["E3"] = "0621_06_26_WERNAGEN"

    listagem = workbook.create_sheet("LISTAGEM_CUT_RITE")
    colunas = {"A": "Descricao", "B": "Material", "C": "Comp", "D": "Larg",
               "E": "Qt", "K": "Artigo", "L": "Notas", "S": "ID",
               "X": "Esp.Mat", "Z": "Tipo_Lacagem"}
    for letra, nome in colunas.items():
        listagem[f"{letra}2"] = nome
    for indice, (descricao, qt, notas, tipo) in enumerate(pecas, start=3):
        listagem[f"A{indice}"] = descricao
        listagem[f"E{indice}"] = qt
        listagem[f"L{indice}"] = notas
        listagem[f"S{indice}"] = indice - 2
        listagem[f"X{indice}"] = 19
        listagem[f"Z{indice}"] = tipo

    lacagem = workbook.create_sheet("Lacagem")
    lacagem["B2"] = "ACABAMENTO\n(LACAGEM / VERNIZ / PINTURA)"
    lacagem["B4"] = "ACABAMENTO"
    lacagem["D6"] = "TIPO\n(Lacado, verniz, etc)"
    lacagem["D8"] = "ACABAMENTO\n(Brilho, matte, etc)"
    lacagem["D10"] = "COR\n(ou RAL/Ref.)"
    if cabecalho:
        lacagem["G6"] = "LACADO"
        lacagem["G8"] = "MATE"
        lacagem["G10"] = "RAL 9001"
    for letra, nome in zip("BCDEFGHIJ", ("Descrição", "Material", "Comp", "Larg",
                                         "Quant", "Legenda", "Observações", "RP", "ID")):
        lacagem[f"{letra}15"] = nome
    for linha in range(17, 17 + linhas_lacagem):
        origem = linha - 14
        lacagem[f"G{linha}"] = (
            f'=IF(NOT(LISTAGEM_CUT_RITE!X{origem}=""),LISTAGEM_CUT_RITE!Z{origem},"")'
        )
        lacagem[f"J{linha}"] = (
            f'=IF(NOT(LISTAGEM_CUT_RITE!X{origem}=""),LISTAGEM_CUT_RITE!S{origem},"")'
        )
    for linha, texto in (observacoes or {}).items():
        lacagem[f"H{linha}"] = texto
    workbook.save(path)
    return path


def test_analise_da_obra_soma_as_pecas_a_lacar_e_ignora_s_lacagem(tmp_path) -> None:
    path = _lista_material(
        tmp_path / "lista.xlsx",
        pecas=[
            ("Costa", 1, None, None),
            ("Painel CIMA PORTA_1", 1, "INTERROMPIDO_LACAR 9001", "A-1 Face + Topos"),
            ("Porta esquerda", 2, "FRESADO_LACAR 9001", "B-2 Face + Topos"),
            ("Regua_Vertical", 2, "S/ LACAGEM", None),
            ("Lateral", 0, "ESTÁ EM OBRA", None),
        ],
    )

    analise = analisar_acabamentos(path)

    assert [peca.descricao for peca in analise.pecas] == [
        "Painel CIMA PORTA_1",
        "Porta esquerda",
    ]
    assert analise.total_pecas == 3
    assert analise.sem_tipo == []
    assert analise.avisos() == []
    assert analise.disponivel is True
    assert analise.plano == PlanoLacagem(
        linha_cabecalho=15,
        primeira_linha=17,
        ultima_linha=22,
        primeira_coluna=2,
        ultima_coluna=10,
        coluna_legenda=7,
        coluna_quant=6,
        linhas_a_acrescentar=0,
    )
    assert analise.plano.linha_total == 23


def test_acabamento_nas_notas_sem_tipo_lacagem_da_aviso(tmp_path) -> None:
    path = _lista_material(
        tmp_path / "lista.xlsx",
        pecas=[
            ("Porta", 1, "LACAR 9001", "A-1 Face + Topos"),
            ("Tampo", 1, "LAKAGEM RAL 9010", None),
            ("Frente", 2, "Pintura", "LACAR"),
            ("Costa", 1, None, None),
        ],
        observacoes={20: "envernizar"},
    )

    analise = analisar_acabamentos(path)

    assert [(a.separador, a.linha, a.palavra) for a in analise.sem_tipo] == [
        ("LISTAGEM_CUT_RITE", 4, "LAKAGEM"),
        ("Lacagem", 20, "ENVERNIZAR"),
    ]
    assert [(a.linha, a.texto) for a in analise.tipo_invalido] == [(5, "LACAR")]
    texto = "\n".join(analise.avisos())
    assert "sem Tipo_Lacagem" in texto
    assert "ID 2 Tampo" in texto
    assert "nenhum dos 5" in texto


def test_cabecalho_por_preencher_e_avisado(tmp_path) -> None:
    path = _lista_material(
        tmp_path / "lista.xlsx",
        pecas=[("Porta", 1, "LACAR", "D-1 Face")],
        cabecalho=False,
    )

    assert analisar_acabamentos(path).cabecalho_em_falta == ["Tipo", "Acabamento", "Cor"]


def test_obra_com_mais_pecas_que_linhas_da_lacagem_acrescenta_as_que_faltam(tmp_path) -> None:
    # O modelo só tem um número fixo de linhas de fórmulas no separador
    # Lacagem: a peça da linha 10 da LISTAGEM ficava fora do PDF.
    pecas = [("Costa", 1, None, None)] * 7 + [("Porta nova", 3, "LACAR", "E- Outro")]
    path = _lista_material(tmp_path / "lista.xlsx", pecas=pecas, linhas_lacagem=5)

    analise = analisar_acabamentos(path)

    assert analise.plano.ultima_linha == 21  # cobre a LISTAGEM até à linha 7
    assert analise.plano.linhas_a_acrescentar == 3  # até à linha 10
    assert analise.plano.linha_total == 25
    assert analise.disponivel is True


def test_sem_nenhum_tipo_lacagem_o_documento_fica_indisponivel(tmp_path) -> None:
    path = _lista_material(tmp_path / "lista.xlsx", pecas=[("Costa", 1, None, None)])

    states = {s.document.identifier: s for s in inspect_pdf_documents(path)}
    estado = states[ACABAMENTOS_DOCUMENT_ID]

    assert estado.available is False
    assert "Tipo_Lacagem" in estado.reason


def test_documento_listagem_acabamentos_no_centro_de_exportacao(tmp_path) -> None:
    path = _lista_material(
        tmp_path / "lista.xlsx", pecas=[("Porta", 2, "LACAR", "A-1 Face + Topos")]
    )
    by_id = {document.identifier: document for document in DEFAULT_DOCUMENTS}
    documento = by_id[ACABAMENTOS_DOCUMENT_ID]

    assert documento.name == "Listagem Acabamentos"
    assert documento.category == "Acabamentos"
    assert documento.sheets == ("Lacagem",)
    # Último grupo: fica na caixa do canto inferior direito do Centro.
    assert documento.order == max(document.order for document in DEFAULT_DOCUMENTS)
    assert document_filename(documento, "0621_06_26_WERNAGEN") == (
        "3_Lacagem_0621_06_26_WERNAGEN.pdf"
    )
    states = {s.document.identifier: s for s in inspect_pdf_documents(path)}
    assert states[ACABAMENTOS_DOCUMENT_ID].available is True
    assert states[ACABAMENTOS_DOCUMENT_ID].export_sheets == ("Lacagem",)
    assert "2 peça(s)" in states[ACABAMENTOS_DOCUMENT_ID].reason


def test_pagina_a4_horizontal_com_rodape_data_obra_e_paginas() -> None:
    pagina = SimpleNamespace(PrintArea="$B$2:$J$48")
    folha = SimpleNamespace(PageSetup=pagina)
    excel = SimpleNamespace()
    plano = PlanoLacagem(15, 17, 47, 2, 10, 7, 6, linhas_a_acrescentar=3)

    _configurar_pagina(excel, folha, plano, "26.0621_06_01_WERNAGEN", date(2026, 9, 19))

    assert pagina.Orientation == 2  # horizontal
    assert pagina.PaperSize == 9  # A4
    assert pagina.FitToPagesWide == 1
    assert pagina.PrintArea == "$B$2:$J$51"
    assert pagina.PrintTitleRows == "$15:$16"
    assert pagina.LeftFooter == "19-09-2026"
    assert pagina.CenterFooter == "26.0621_06_01_WERNAGEN"
    assert pagina.RightFooter == "&P/&N"
    assert excel.PrintCommunication is True


def test_rodape_protege_o_e_comercial() -> None:
    assert texto_rodape("OBRA A & B") == "OBRA A && B"
    folha = SimpleNamespace(PageSetup=SimpleNamespace(PrintArea=""))
    assert _area_impressao(folha, PlanoLacagem(15, 17, 20, 2, 10, 7, 6)) == "$B$1:$J$21"


def test_exportacao_abre_o_excel_em_calculo_manual() -> None:
    # Com as macros desligadas, recalcular dava #NOME? na Esp.Mat (função VBA)
    # e todas as linhas da Lacagem saíam em erro.
    from app.services import lista_material_pdf_service as servico

    fonte = inspect.getsource(servico.export_pdf_documents)
    assert fonte.index("abrir_em_calculo_manual(excel)") < fonte.index("Workbooks.Open")
    assert "preparar_separador_lacagem(" in fonte


def test_dialogo_avisa_e_passa_o_nome_da_obra() -> None:
    from app.ui.dialogs.lista_material_pdf_dialog import ListaMaterialPdfDialog

    export_source = inspect.getsource(ListaMaterialPdfDialog._export)
    confirm_source = inspect.getsource(ListaMaterialPdfDialog._confirmar_acabamentos)

    assert "self._confirmar_acabamentos()" in export_source
    assert "obra_nome=self.obra_nome" in export_source
    assert "Voltar e corrigir" in confirm_source
    assert "Exportar assim mesmo" in confirm_source
