"""Traduções do iX (imos.msg): ler o Excel, comparar, aplicar com cópia, repor."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.services import imos_traducoes_service as s

BOM = b"\xef\xbb\xbf"

#: Um pedaço com o formato real: cabeçalho com «!», tab ou espaço antes do 2.º
#: «;», CRLF, várias línguas e um separador de linha Unicode dentro de um texto.
MSG = (
    "! Version 10.2\r\n"
    "!*** IMPORTANT: UTF-8 ***!\r\n"
    "GER;7134\t;Kommission\r\n"
    "ENG;7134\t;Commission\r\n"
    "PTG;7134\t;Comissão\r\n"
    "PTG;1220\t;Descrição\r\n"
    "PTG;32250 ;Activar/Desactivar Tags\r\n"
    "CHN;9000\t;文本\u2028第二行\r\n"
    "PTG;10279\t;Número de artigo\r\n"
)


def _msg(tmp_path: Path, texto: str = MSG, bom: bool = True) -> Path:
    pasta = tmp_path / "iX CAD 2025" / "BIN" / "MSG"
    pasta.mkdir(parents=True)
    caminho = pasta / "imos.msg"
    caminho.write_bytes((BOM if bom else b"") + texto.encode("utf-8"))
    return caminho


def _excel(tmp_path: Path, linhas, nome: str = "imos_msg.xlsx") -> Path:
    livro = Workbook()
    folha = livro.active
    for numero, (referencia, texto) in enumerate(linhas, start=6):
        folha.cell(numero, 2, referencia)
        folha.cell(numero, 3, texto)
    caminho = tmp_path / nome
    livro.save(caminho)
    return caminho


def _lista(tmp_path: Path) -> s.ListaTraducoes:
    return s.ler_excel(
        _excel(
            tmp_path,
            [
                ("PTG;7134", "Enc PHC:"),
                ("PTG;1220", "Materiais Usados"),
                ("PTG;32250", "Activar/Desactivar Tags"),
                ("PTG;99999", "Não existe"),
                (None, None),
                ("C:\\Program Files\\imos AG\\iX CAD 2025\\BIN\\MSG", None),
            ],
        )
    )


# ---- Excel ------------------------------------------------------------------------------
def test_referencias_normalizadas() -> None:
    assert s.normalizar_referencia("PTG;10280") == "PTG;10280"
    assert s.normalizar_referencia(" ptg ; 10280 ") == "PTG;10280"
    assert s.normalizar_referencia(10280.0) == "PTG;10280"
    assert s.normalizar_referencia("PTG;0123") == "PTG;123"
    assert s.normalizar_referencia(r"C:\Program Files\imos AG") is None
    assert s.normalizar_referencia("") is None
    assert s.normalizar_texto("Linha 1\nLinha 2 ") == r"Linha 1\nLinha 2"


def test_ler_excel_ignora_linhas_vazias_e_o_caminho_e_avisa_repetidas(tmp_path) -> None:
    excel = _excel(
        tmp_path,
        [
            ("PTG;7158", "Ref Cliente:"),
            ("PTG;7129", "Ref Cliente:"),
            ("PTG;7158", "Ref. do Cliente"),
            ("C:\\Program Files\\imos AG\\iX CAD 2025\\BIN\\MSG", None),
            ("PTG;5994", "."),
        ],
    )
    lista = s.ler_excel(excel)
    assert [t.chave for t in lista.traducoes] == ["PTG;7129", "PTG;7158", "PTG;5994"]
    assert lista.por_chave["PTG;7158"] == "Ref. do Cliente"  # vale a última
    assert lista.repetidas == ("PTG;7158",)


def test_excel_sem_traducoes_ou_inexistente(tmp_path) -> None:
    with pytest.raises(s.ErroTraducoes, match="não tem traduções"):
        s.ler_excel(_excel(tmp_path, [(None, "texto sem referência")]))
    with pytest.raises(s.ErroTraducoes, match="Não encontrei o Excel"):
        s.ler_excel(tmp_path / "nao_existe.xlsx")
    falso = tmp_path / "falso.xlsx"
    falso.write_text("isto não é um Excel", encoding="utf-8")
    with pytest.raises(s.ErroTraducoes, match="não é um Excel válido"):
        s.ler_excel(falso)


# ---- comparar -----------------------------------------------------------------------------
def test_comparar_diz_o_que_falta(tmp_path) -> None:
    estados = {e.chave: e for e in s.comparar(_msg(tmp_path), _lista(tmp_path))}
    assert estados["PTG;7134"].texto_atual == "Comissão"
    assert estados["PTG;7134"].estado == s.ESTADO_POR_APLICAR
    assert estados["PTG;32250"].estado == s.ESTADO_CERTA
    assert estados["PTG;99999"].estado == s.ESTADO_NAO_EXISTE
    assert s.contar_por_aplicar(estados.values()) == 2


def test_ficheiro_que_nao_e_utf8_nao_se_mexe(tmp_path) -> None:
    caminho = _msg(tmp_path)
    caminho.write_bytes("PTG;7134\t;Comissão\r\n".encode("cp1252"))
    with pytest.raises(s.ErroTraducoes, match="UTF-8"):
        s.comparar(caminho, _lista(tmp_path))


# ---- aplicar ------------------------------------------------------------------------------
def test_aplicar_faz_copia_e_muda_so_as_linhas_do_excel(tmp_path) -> None:
    caminho = _msg(tmp_path)
    original = caminho.read_bytes()
    passos = []
    resultado = s.aplicar(
        caminho,
        _lista(tmp_path),
        agora=datetime(2026, 10, 10, 10, 15, 0),
        listar_processos=lambda: {"explorer.exe"},
        ao_passo=lambda texto, tipo: passos.append((tipo, texto)),
    )
    # 1) a cópia é o ficheiro de antes, byte a byte
    assert resultado.copia == caminho.with_name("imos.msg.copia_2026-10-10_101500")
    assert resultado.copia.read_bytes() == original
    # 2) só mudaram as duas linhas; o resto (BOM, CRLF, tab/espaço, chinês) igual
    esperado = BOM + MSG.replace("PTG;7134\t;Comissão", "PTG;7134\t;Enc PHC:").replace(
        "PTG;1220\t;Descrição", "PTG;1220\t;Materiais Usados"
    ).encode("utf-8")
    assert caminho.read_bytes() == esperado
    assert [a.chave for a in resultado.alteradas] == ["PTG;7134", "PTG;1220"]
    assert resultado.alteradas[0].antes == "Comissão"
    assert resultado.certas == ["PTG;32250"]
    assert resultado.nao_encontradas == ["PTG;99999"]
    assert ("muda", "PTG;7134    «Comissão»  →  «Enc PHC:»") in passos
    # 3) sem ficheiros temporários esquecidos
    assert sorted(p.name for p in caminho.parent.iterdir()) == [
        "imos.msg",
        "imos.msg.copia_2026-10-10_101500",
    ]


def test_sem_nada_para_mudar_nao_ha_copia_nem_escrita(tmp_path) -> None:
    caminho = _msg(tmp_path)
    lista = _lista(tmp_path)
    s.aplicar(caminho, lista, agora=datetime(2026, 10, 10, 10, 0), listar_processos=set)
    depois = caminho.read_bytes()
    resultado = s.aplicar(caminho, lista, agora=datetime(2026, 10, 10, 11, 0), listar_processos=set)
    assert resultado.copia is None and resultado.alteradas == []
    assert caminho.read_bytes() == depois
    assert len(s.listar_copias(caminho)) == 1


def test_ficheiro_sem_bom_e_com_lf_fica_assim(tmp_path) -> None:
    caminho = _msg(tmp_path, MSG.replace("\r\n", "\n"), bom=False)
    s.aplicar(caminho, _lista(tmp_path), listar_processos=set)
    dados = caminho.read_bytes()
    assert not dados.startswith(BOM)
    assert b"\r\n" not in dados and b"PTG;7134\t;Enc PHC:\n" in dados


def test_texto_repetido_no_ficheiro_muda_todas_as_vezes(tmp_path) -> None:
    caminho = _msg(tmp_path, MSG + "PTG;7134\t;Comissão\r\n")
    resultado = s.aplicar(caminho, _lista(tmp_path), listar_processos=set)
    assert resultado.alteradas[0].ocorrencias == 2
    assert caminho.read_text(encoding="utf-8-sig").count("PTG;7134\t;Enc PHC:") == 2


def test_com_o_ix_aberto_nao_se_mexe(tmp_path) -> None:
    caminho = _msg(tmp_path)
    original = caminho.read_bytes()
    with pytest.raises(s.ErroTraducoes, match="Feche primeiro: iX CAD e iX Organizer"):
        s.aplicar(caminho, _lista(tmp_path), listar_processos=lambda: {"IMOS.EXE", "Organizer.exe"})
    assert caminho.read_bytes() == original
    assert s.listar_copias(caminho) == []


def test_sem_copia_nao_se_grava(tmp_path, monkeypatch) -> None:
    caminho = _msg(tmp_path)
    original = caminho.read_bytes()

    def sem_permissao(*_a, **_k):
        raise PermissionError("negado")

    monkeypatch.setattr(s.shutil, "copy2", sem_permissao)
    with pytest.raises(s.ErroTraducoes, match="Sem cópia não se mexe"):
        s.aplicar(caminho, _lista(tmp_path), listar_processos=set)
    assert caminho.read_bytes() == original


def test_falha_a_gravar_deixa_o_ficheiro_como_estava(tmp_path, monkeypatch) -> None:
    caminho = _msg(tmp_path)
    original = caminho.read_bytes()

    def falha(*_a, **_k):
        raise PermissionError("em uso")

    monkeypatch.setattr(s.os, "replace", falha)
    with pytest.raises(s.ErroTraducoes, match="ficou como estava"):
        s.aplicar(caminho, _lista(tmp_path), listar_processos=set)
    assert caminho.read_bytes() == original
    assert not list(caminho.parent.glob("*.tmp"))


# ---- cópias -------------------------------------------------------------------------------
def test_listar_e_repor_copia_guarda_o_atual(tmp_path) -> None:
    caminho = _msg(tmp_path)
    original = caminho.read_bytes()
    (caminho.parent / "imos.msg.bak_20250101_120000").write_bytes(original)
    (caminho.parent / "imos.msg.backup_20240101_120000").write_bytes(original)
    (caminho.parent / "outra coisa.txt").write_text("x", encoding="utf-8")
    resultado = s.aplicar(caminho, _lista(tmp_path), agora=datetime(2026, 10, 10, 9, 0), listar_processos=set)
    alterado = caminho.read_bytes()
    nomes = {c.caminho.name for c in s.listar_copias(caminho)}
    assert nomes == {
        "imos.msg.copia_2026-10-10_090000",
        "imos.msg.bak_20250101_120000",
        "imos.msg.backup_20240101_120000",
    }
    guardado = s.repor_copia(
        caminho, resultado.copia, agora=datetime(2026, 10, 10, 9, 30), listar_processos=set
    )
    assert caminho.read_bytes() == original
    assert guardado.name == "imos.msg.copia_2026-10-10_093000_antes_de_repor"
    assert guardado.read_bytes() == alterado
    assert resultado.copia.exists()  # repor não apaga a cópia


def test_repor_com_o_ix_aberto_nao_mexe(tmp_path) -> None:
    caminho = _msg(tmp_path)
    copia = caminho.with_name("imos.msg.copia_x")
    copia.write_bytes(b"PTG;1\t;x\r\n")
    with pytest.raises(s.ErroTraducoes, match="Feche primeiro"):
        s.repor_copia(caminho, copia, listar_processos=lambda: {"imos.exe"})


# ---- onde está o iX e que programas estão abertos ------------------------------------------
def test_localizar_pelo_registo_e_pela_pasta(tmp_path) -> None:
    caminho = _msg(tmp_path)
    instalacao = SimpleNamespace(pasta_trabalho=caminho.parent.parent)
    assert s.localizar_imos_msg(localizar_instalacao=lambda: instalacao) == caminho
    # Sem registo: a pasta «iX CAD <ano>» mais recente que tenha o ficheiro.
    (tmp_path / "iX CAD 2023" / "BIN" / "MSG").mkdir(parents=True)
    (tmp_path / "iX CAD 2023" / "BIN" / "MSG" / "imos.msg").write_bytes(b"x")
    (tmp_path / "iX CAD 2027").mkdir()  # instalação a meio, sem ficheiro
    assert s.localizar_imos_msg(localizar_instalacao=lambda: None, pasta_programas=tmp_path) == caminho
    assert s.localizar_imos_msg(localizar_instalacao=lambda: None, pasta_programas=tmp_path / "x") is None
    assert s.versao_do_caminho(caminho) == "iX CAD 2025"


def test_estado_dos_programas() -> None:
    assert s.estado_programas(lambda: {"Imos.exe", "excel.exe"}) == {
        "iX CAD": True,
        "iX Organizer": False,
    }
    assert s.programas_abertos(lambda: {"organizer.exe"}) == ["iX Organizer"]
    assert isinstance(s._executaveis_a_correr(), set)
