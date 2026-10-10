"""Traduções do iX CAD: pôr no ``imos.msg`` deste PC os textos que a empresa usa.

O iX CAD lê os textos de todos os campos de um ficheiro de traduções,
``<iX CAD>\\BIN\\MSG\\imos.msg`` (UTF-8 com BOM, 34 línguas, ~650 mil linhas
``PTG;7134<tab>;Texto``). Alguns campos têm, em português, nomes que não
dizem nada na Lança Encanto — o «Kommission» do iX aparece como «Enc PHC:»
porque é aí que se escreve o número da encomenda do PHC.

A lista do que mudar vive num Excel partilhado (``I:\\imos_msg.xlsx``: coluna B
= referência, ``PTG;10280``; coluna C = texto da empresa), que vai crescendo.
Era aplicada PC a PC por um executável (``AtualizaIMOSMsgPTG.exe``, com a
lista escrita no código e o caminho do iX 2023) e depois pelo Martelo V2
(``imos_msg_sync.py``). Este módulo junta os dois, para o iX 2025 e seguintes:

* o ficheiro do iX encontra-se sozinho (o iX mais recente instalado no PC);
* só se grava com o iX CAD e o iX Organizer FECHADOS;
* antes de mexer, faz-se SEMPRE uma cópia do ``imos.msg`` na mesma pasta;
* trocam-se só as linhas das referências do Excel; tudo o resto do ficheiro
  (BOM, fins de linha, separadores) fica igual, byte a byte;
* a escrita é atómica (ficheiro temporário + troca) e confirmada a seguir.

Nunca apaga nada: as cópias ficam na pasta do iX e repor uma cópia também faz
primeiro uma cópia do ficheiro que lá está.
"""

from __future__ import annotations

import ctypes
import os
import re
import shutil
import sys
import tempfile
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

CHAVE_EXCEL = "imos_traducoes_excel"
EXCEL_PADRAO = r"I:\imos_msg.xlsx"
NOME_FICHEIRO = "imos.msg"
PASTA_PROGRAMAS_IMOS = Path(r"C:\Program Files\imos AG")

#: As cópias novas chamam-se ``imos.msg.copia_2026-10-10_101500``; as antigas
#: do executável (``.backup_``) e do V2 (``.bak_``) também se podem repor.
PREFIXO_COPIA = NOME_FICHEIRO + ".copia_"
PREFIXOS_COPIAS = (PREFIXO_COPIA, NOME_FICHEIRO + ".bak_", NOME_FICHEIRO + ".backup_")

#: Os programas que têm o ``imos.msg`` aberto: executável → nome que se mostra.
PROGRAMAS_IX = {"imos.exe": "iX CAD", "organizer.exe": "iX Organizer"}

ESTADO_CERTA = "certa"
ESTADO_POR_APLICAR = "por_aplicar"
ESTADO_NAO_EXISTE = "nao_existe"

_BOM = b"\xef\xbb\xbf"
#: ``PTG;10280<tab>;Texto`` — o separador antes do segundo «;» varia (tab ou
#: espaços) e mantém-se como está.
_LINHA_RE = re.compile(
    r"^(?P<lingua>[A-Z0-9]{3});(?P<ident>\d+)(?P<separador>[ \t]*;)(?P<texto>[^\r\n]*)(?P<fim>\r?\n)?\Z"
)
_LINGUA_RE = re.compile(r"^[A-Z0-9]{3}$")


class ErroTraducoes(RuntimeError):
    """Algo impede ler ou gravar; a mensagem já está pronta para o utilizador."""


# --------------------------------------------------------------------------
# O Excel da empresa
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class Traducao:
    chave: str  # «PTG;10280»
    texto: str
    linha_excel: int


@dataclass(frozen=True)
class ListaTraducoes:
    excel: Path
    traducoes: tuple[Traducao, ...]
    #: Referências que aparecem mais do que uma vez com textos diferentes (vale a última).
    repetidas: tuple[str, ...] = ()

    @property
    def por_chave(self) -> dict[str, str]:
        return {t.chave: t.texto for t in self.traducoes}


def normalizar_referencia(valor: object) -> str | None:
    """«PTG;10280», «ptg ; 10280», «10280» ou 10280.0 → «PTG;10280»; ``None`` se não servir."""
    texto = "" if valor is None else str(valor).strip()
    if not texto:
        return None
    lingua, _, ident = texto.partition(";") if ";" in texto else ("PTG", "", texto)
    lingua = re.sub(r"\s+", "", lingua).upper()
    ident = ident.strip()
    if ident.endswith(".0"):
        ident = ident[:-2]
    if not _LINGUA_RE.fullmatch(lingua) or not ident.isdigit():
        return None
    return f"{lingua};{int(ident)}"


def normalizar_texto(valor: object) -> str:
    """O texto da coluna C numa só linha: as mudanças de linha passam a ``\\n``, como no iX."""
    texto = "" if valor is None else str(valor)
    texto = texto.replace("\r\n", "\n").replace("\r", "\n").replace("\n", r"\n")
    return texto.strip()


def ler_excel(caminho: str | Path) -> ListaTraducoes:
    """As traduções do Excel (colunas B e C da primeira folha)."""
    from openpyxl import load_workbook

    excel = Path(caminho)
    try:
        dados = excel.read_bytes()
    except FileNotFoundError as erro:
        raise ErroTraducoes(
            f"Não encontrei o Excel das traduções:\n{excel}\n\n"
            "Confirme que a unidade I: está ligada."
        ) from erro
    except OSError as erro:
        raise ErroTraducoes(f"Não consegui ler o Excel das traduções:\n{excel}\n\n{erro}") from erro
    try:
        livro = load_workbook(BytesIO(dados), read_only=True, data_only=True)
    except (BadZipFile, OSError, ValueError, KeyError) as erro:
        raise ErroTraducoes(
            f"O ficheiro não é um Excel válido (.xlsx):\n{excel}"
        ) from erro
    por_chave: dict[str, Traducao] = {}
    repetidas: list[str] = []
    try:
        folha = livro.worksheets[0]
        for numero, (referencia, texto) in enumerate(
            folha.iter_rows(min_col=2, max_col=3, values_only=True), start=1
        ):
            chave = normalizar_referencia(referencia)
            novo = normalizar_texto(texto)
            if not chave or not novo:
                continue
            anterior = por_chave.get(chave)
            if anterior is not None and anterior.texto != novo and chave not in repetidas:
                repetidas.append(chave)
            por_chave[chave] = Traducao(chave, novo, numero)
    finally:
        livro.close()
    if not por_chave:
        raise ErroTraducoes(
            "O Excel não tem traduções nas colunas B e C.\n\n"
            "Esperado: coluna B = referência (ex.: PTG;10280); coluna C = texto da empresa."
        )
    traducoes = tuple(sorted(por_chave.values(), key=lambda t: t.linha_excel))
    return ListaTraducoes(excel, traducoes, tuple(repetidas))


# --------------------------------------------------------------------------
# O ficheiro do iX
# --------------------------------------------------------------------------


def localizar_imos_msg(
    *,
    localizar_instalacao: Callable[[], object] | None = None,
    pasta_programas: Path = PASTA_PROGRAMAS_IMOS,
) -> Path | None:
    """O ``imos.msg`` do iX CAD mais recente instalado neste PC, ou ``None``.

    Primeiro pelo registo do Windows (o mesmo que o Organizer usa); se falhar,
    a pasta ``iX CAD <ano>`` mais recente em ``C:\\Program Files\\imos AG``.
    """
    if localizar_instalacao is None:
        from app.services.imos_cad_service import localizar_instalacao as _localizar

        localizar_instalacao = _localizar
    try:
        instalacao = localizar_instalacao()
    except Exception:  # noqa: BLE001 - registo ilegível: tenta-se a pasta
        instalacao = None
    if instalacao is not None:
        candidato = Path(instalacao.pasta_trabalho) / "MSG" / NOME_FICHEIRO
        if candidato.is_file():
            return candidato
    try:
        pastas = sorted(
            (p for p in pasta_programas.glob("iX CAD *") if p.is_dir()),
            key=lambda p: p.name,
            reverse=True,
        )
    except OSError:
        return None
    for pasta in pastas:
        candidato = pasta / "BIN" / "MSG" / NOME_FICHEIRO
        if candidato.is_file():
            return candidato
    return None


def versao_do_caminho(caminho: Path | None) -> str:
    """«iX CAD 2025» a partir do caminho do ficheiro (vazio se não se perceber)."""
    if caminho is None:
        return ""
    for parte in Path(caminho).parts:
        if parte.lower().startswith("ix cad"):
            return parte
    return ""


@dataclass
class _FicheiroMsg:
    caminho: Path
    bom: bool
    linhas: list[str]


def _ler_msg(caminho: Path) -> _FicheiroMsg:
    try:
        dados = caminho.read_bytes()
    except FileNotFoundError as erro:
        raise ErroTraducoes(f"Não encontrei o ficheiro do iX:\n{caminho}") from erro
    except PermissionError as erro:
        raise ErroTraducoes(
            f"Sem permissão para ler o ficheiro do iX:\n{caminho}"
        ) from erro
    except OSError as erro:
        raise ErroTraducoes(f"Não consegui ler o ficheiro do iX:\n{caminho}\n\n{erro}") from erro
    try:
        texto = dados.decode("utf-8-sig")
    except UnicodeDecodeError as erro:
        raise ErroTraducoes(
            f"O ficheiro do iX não está em UTF-8 e não lhe vou mexer:\n{caminho}"
        ) from erro
    return _FicheiroMsg(caminho, dados.startswith(_BOM), _partir_linhas(texto))


def _partir_linhas(texto: str) -> list[str]:
    """As linhas COM o seu fim (``\\r\\n``), partidas só no ``\\n``.

    O ``splitlines`` também parte em separadores que algumas línguas usam
    dentro do texto (``\\u2028``, ``\\x85``…); aqui juntar as linhas devolve
    sempre o ficheiro tal e qual.
    """
    partes = texto.split("\n")
    linhas = [parte + "\n" for parte in partes[:-1]]
    if partes[-1]:
        linhas.append(partes[-1])
    return linhas


def _chave_da_linha(linha: str, prefixos: tuple[str, ...]) -> tuple[str, re.Match] | None:
    if not linha.startswith(prefixos):
        return None
    encontrado = _LINHA_RE.match(linha)
    if encontrado is None:
        return None
    return f"{encontrado['lingua']};{int(encontrado['ident'])}", encontrado


def _prefixos(chaves: Iterable[str]) -> tuple[str, ...]:
    return tuple(sorted({chave.split(";", 1)[0] + ";" for chave in chaves}))


@dataclass(frozen=True)
class EstadoTraducao:
    chave: str
    texto_atual: str | None  # ``None`` = a referência não existe no imos.msg
    texto_novo: str

    @property
    def estado(self) -> str:
        if self.texto_atual is None:
            return ESTADO_NAO_EXISTE
        return ESTADO_CERTA if self.texto_atual == self.texto_novo else ESTADO_POR_APLICAR


def comparar(caminho_msg: str | Path, lista: ListaTraducoes) -> list[EstadoTraducao]:
    """Para cada tradução do Excel, o que o iX mostra hoje neste PC."""
    ficheiro = _ler_msg(Path(caminho_msg))
    desejado = lista.por_chave
    prefixos = _prefixos(desejado)
    atual: dict[str, str] = {}
    for linha in ficheiro.linhas:
        par = _chave_da_linha(linha, prefixos)
        if par is None:
            continue
        chave, encontrado = par
        if chave in desejado and chave not in atual:
            atual[chave] = encontrado["texto"]
    return [EstadoTraducao(t.chave, atual.get(t.chave), t.texto) for t in lista.traducoes]


def contar_por_aplicar(estados: Iterable[EstadoTraducao]) -> int:
    return sum(1 for e in estados if e.estado == ESTADO_POR_APLICAR)


# --------------------------------------------------------------------------
# Programas abertos
# --------------------------------------------------------------------------


def _executaveis_a_correr() -> set[str]:
    """Os nomes dos executáveis a correr neste PC (todas as sessões), em minúsculas."""
    if sys.platform != "win32":
        return set()
    from ctypes import wintypes

    class _PROCESSENTRY32W(ctypes.Structure):
        _fields_ = [
            ("dwSize", wintypes.DWORD),
            ("cntUsage", wintypes.DWORD),
            ("th32ProcessID", wintypes.DWORD),
            ("th32DefaultHeapID", ctypes.c_void_p),
            ("th32ModuleID", wintypes.DWORD),
            ("cntThreads", wintypes.DWORD),
            ("th32ParentProcessID", wintypes.DWORD),
            ("pcPriClassBase", ctypes.c_long),
            ("dwFlags", wintypes.DWORD),
            ("szExeFile", ctypes.c_wchar * 260),
        ]

    k = ctypes.WinDLL("kernel32", use_last_error=True)
    k.CreateToolhelp32Snapshot.restype = wintypes.HANDLE
    k.CreateToolhelp32Snapshot.argtypes = [wintypes.DWORD, wintypes.DWORD]
    k.Process32FirstW.argtypes = [wintypes.HANDLE, ctypes.POINTER(_PROCESSENTRY32W)]
    k.Process32NextW.argtypes = [wintypes.HANDLE, ctypes.POINTER(_PROCESSENTRY32W)]
    k.CloseHandle.argtypes = [wintypes.HANDLE]
    snapshot = k.CreateToolhelp32Snapshot(0x00000002, 0)  # TH32CS_SNAPPROCESS
    if not snapshot or snapshot == wintypes.HANDLE(-1).value:
        return set()
    nomes: set[str] = set()
    try:
        entrada = _PROCESSENTRY32W()
        entrada.dwSize = ctypes.sizeof(_PROCESSENTRY32W)
        continua = k.Process32FirstW(snapshot, ctypes.byref(entrada))
        while continua:
            nomes.add(entrada.szExeFile.casefold())
            continua = k.Process32NextW(snapshot, ctypes.byref(entrada))
    finally:
        k.CloseHandle(snapshot)
    return nomes


def estado_programas(
    listar: Callable[[], set[str]] = _executaveis_a_correr,
) -> dict[str, bool]:
    """Nome a mostrar → está aberto? (iX CAD e iX Organizer)."""
    a_correr = {nome.casefold() for nome in listar()}
    return {mostrar: exe in a_correr for exe, mostrar in PROGRAMAS_IX.items()}


def programas_abertos(listar: Callable[[], set[str]] = _executaveis_a_correr) -> list[str]:
    return [nome for nome, aberto in estado_programas(listar).items() if aberto]


# --------------------------------------------------------------------------
# Aplicar e repor
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class Alteracao:
    chave: str
    antes: str
    depois: str
    ocorrencias: int = 1


@dataclass
class ResultadoAplicacao:
    caminho_msg: Path
    copia: Path | None = None
    alteradas: list[Alteracao] = field(default_factory=list)
    certas: list[str] = field(default_factory=list)
    nao_encontradas: list[str] = field(default_factory=list)
    total_linhas: int = 0


def _nome_copia(caminho_msg: Path, agora: datetime, sufixo: str = "") -> Path:
    base = caminho_msg.with_name(f"{PREFIXO_COPIA}{agora:%Y-%m-%d_%H%M%S}{sufixo}")
    candidato, numero = base, 2
    while candidato.exists():
        candidato = base.with_name(f"{base.name}_{numero}")
        numero += 1
    return candidato


def _fazer_copia(caminho_msg: Path, agora: datetime, sufixo: str = "") -> Path:
    destino = _nome_copia(caminho_msg, agora, sufixo)
    try:
        shutil.copy2(caminho_msg, destino)
    except PermissionError as erro:
        raise ErroTraducoes(
            "Sem permissão para fazer a cópia do imos.msg na pasta do iX:\n"
            f"{caminho_msg.parent}\n\n"
            "Sem cópia não se mexe no ficheiro. Peça para darem permissão de "
            "escrita a esta pasta (a do iX CAD) neste PC."
        ) from erro
    except OSError as erro:
        raise ErroTraducoes(
            f"Não consegui fazer a cópia do imos.msg:\n{destino}\n\n{erro}"
        ) from erro
    if destino.stat().st_size != caminho_msg.stat().st_size:
        raise ErroTraducoes(f"A cópia do imos.msg ficou incompleta:\n{destino}")
    return destino


def _gravar_atomico(destino: Path, dados: bytes) -> None:
    temporario: Path | None = None
    try:
        descritor, nome = tempfile.mkstemp(prefix=f"{destino.name}.", suffix=".tmp", dir=str(destino.parent))
        temporario = Path(nome)
        with os.fdopen(descritor, "wb") as saida:
            saida.write(dados)
        os.replace(temporario, destino)
        temporario = None
    except PermissionError as erro:
        raise ErroTraducoes(
            f"Sem permissão para gravar o ficheiro do iX:\n{destino}\n\n"
            "Confirme que o iX CAD e o iX Organizer estão fechados. O ficheiro "
            "ficou como estava."
        ) from erro
    except OSError as erro:
        raise ErroTraducoes(
            f"Não consegui gravar o ficheiro do iX:\n{destino}\n\n{erro}\n\nO ficheiro ficou como estava."
        ) from erro
    finally:
        if temporario is not None and temporario.exists():
            # O temporário é nosso e nunca chegou a ser o imos.msg.
            temporario.unlink(missing_ok=True)


def _confirmar_fechados(listar: Callable[[], set[str]] | None) -> None:
    abertos = programas_abertos(listar) if listar is not None else programas_abertos()
    if abertos:
        raise ErroTraducoes(
            "Feche primeiro: " + " e ".join(abertos) + ".\n\n"
            "Enquanto estão abertos, o iX tem o ficheiro das traduções em uso."
        )


def aplicar(
    caminho_msg: str | Path,
    lista: ListaTraducoes,
    *,
    agora: datetime | None = None,
    listar_processos: Callable[[], set[str]] | None = None,
    ao_passo: Callable[[str, str], None] | None = None,
) -> ResultadoAplicacao:
    """Copia o ``imos.msg`` e põe-lhe os textos do Excel.

    ``ao_passo(texto, tipo)`` conta o que vai acontecendo (``tipo`` = ``ok``,
    ``muda``, ``info``) para a janela que mostra o progresso.
    """
    avisar = ao_passo or (lambda _texto, _tipo: None)
    caminho = Path(caminho_msg)
    agora = agora or datetime.now()
    _confirmar_fechados(listar_processos)
    avisar("iX CAD e iX Organizer fechados.", "ok")

    ficheiro = _ler_msg(caminho)
    desejado = lista.por_chave
    prefixos = _prefixos(desejado)
    avisar(
        f"A procurar {len(desejado)} referências no imos.msg ({len(ficheiro.linhas):,} linhas)…".replace(",", " "),
        "info",
    )

    resultado = ResultadoAplicacao(caminho_msg=caminho, total_linhas=len(ficheiro.linhas))
    encontradas: set[str] = set()
    mudancas: dict[str, Alteracao] = {}
    novas: list[str] = []
    for linha in ficheiro.linhas:
        par = _chave_da_linha(linha, prefixos)
        if par is None or par[0] not in desejado:
            novas.append(linha)
            continue
        chave, encontrado = par
        encontradas.add(chave)
        texto = desejado[chave]
        if encontrado["texto"] == texto:
            novas.append(linha)
            continue
        novas.append(
            f"{encontrado['lingua']};{encontrado['ident']}{encontrado['separador']}{texto}{encontrado['fim'] or ''}"
        )
        anterior = mudancas.get(chave)
        mudancas[chave] = (
            Alteracao(chave, encontrado["texto"], texto)
            if anterior is None
            else Alteracao(chave, anterior.antes, texto, anterior.ocorrencias + 1)
        )

    for traducao in lista.traducoes:
        if traducao.chave not in encontradas:
            resultado.nao_encontradas.append(traducao.chave)
        elif traducao.chave in mudancas:
            resultado.alteradas.append(mudancas[traducao.chave])
        else:
            resultado.certas.append(traducao.chave)

    if not resultado.alteradas:
        avisar("Nada a mudar: o iX deste PC já tem todos os textos do Excel.", "ok")
        return resultado

    # 1.º a cópia; sem ela não se grava nada.
    resultado.copia = _fazer_copia(caminho, agora)
    avisar(f"Cópia de segurança feita: {resultado.copia.name}", "ok")
    for alteracao in resultado.alteradas:
        avisar(f"{alteracao.chave:<11} «{alteracao.antes}»  →  «{alteracao.depois}»", "muda")

    dados = "".join(novas).encode("utf-8")
    if ficheiro.bom:
        dados = _BOM + dados
    _gravar_atomico(caminho, dados)

    # Confirmar no próprio ficheiro, já gravado, que ficou como devia.
    verificacao = comparar(caminho, lista)
    falhas = [e.chave for e in verificacao if e.estado == ESTADO_POR_APLICAR]
    if falhas:
        raise ErroTraducoes(
            "O ficheiro foi gravado mas estas referências não ficaram certas: "
            + ", ".join(falhas)
            + f"\n\nA cópia de antes está em:\n{resultado.copia}"
        )
    avisar(
        f"imos.msg gravado e confirmado: {len(resultado.alteradas)} texto(s) alterado(s), o resto ficou igual.",
        "ok",
    )
    return resultado


@dataclass(frozen=True)
class CopiaMsg:
    caminho: Path
    modificado: datetime
    tamanho: int


def listar_copias(caminho_msg: str | Path) -> list[CopiaMsg]:
    """As cópias do ``imos.msg`` na pasta do iX, a mais recente primeiro."""
    pasta = Path(caminho_msg).parent
    copias: list[CopiaMsg] = []
    try:
        candidatos = list(pasta.iterdir())
    except OSError:
        return copias
    for candidato in candidatos:
        nome = candidato.name.casefold()
        if not nome.startswith(PREFIXOS_COPIAS) or nome.endswith(".tmp") or not candidato.is_file():
            continue
        info = candidato.stat()
        copias.append(CopiaMsg(candidato, datetime.fromtimestamp(info.st_mtime), info.st_size))
    return sorted(copias, key=lambda c: c.modificado, reverse=True)


def repor_copia(
    caminho_msg: str | Path,
    copia: str | Path,
    *,
    agora: datetime | None = None,
    listar_processos: Callable[[], set[str]] | None = None,
) -> Path:
    """Volta a pôr uma cópia no lugar do ``imos.msg``.

    Antes, o ficheiro que lá está também vai para uma cópia — repor nunca
    deita fora nada. Devolve o caminho dessa cópia do ficheiro substituído.
    """
    caminho = Path(caminho_msg)
    origem = Path(copia)
    if not origem.is_file():
        raise ErroTraducoes(f"A cópia já não existe:\n{origem}")
    _confirmar_fechados(listar_processos)
    try:
        origem.read_bytes().decode("utf-8-sig")
    except UnicodeDecodeError as erro:
        raise ErroTraducoes(f"Esta cópia não é um ficheiro de traduções válido:\n{origem}") from erro
    guardado = _fazer_copia(caminho, agora or datetime.now(), "_antes_de_repor")
    _gravar_atomico(caminho, origem.read_bytes())
    return guardado
