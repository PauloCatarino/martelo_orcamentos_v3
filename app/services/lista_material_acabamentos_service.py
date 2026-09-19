"""Listagem de Acabamentos (separador Lacagem) do Centro de Exportação PDF.

O separador ``Lacagem`` do Excel da Lista Material é a folha que vai para o
fornecedor de lacagem/verniz/pintura. Cada linha espelha, por fórmula, uma
linha da ``LISTAGEM_CUT_RITE``; a coluna ``Legenda`` mostra o ``Tipo_Lacagem``
que o utilizador escolheu para a peça. Só as peças com um dos cinco tipos
seguem no PDF.

Duas armadilhas guardadas aqui:

* o separador tem um número FIXO de linhas de fórmulas (27 no modelo). Uma
  obra com mais peças deixava as últimas de fora do PDF sem ninguém dar por
  isso; na exportação acrescentam-se as linhas que faltam (só em memória, o
  Excel nunca é gravado);
* a peça só entra no PDF se o ``Tipo_Lacagem`` estiver preenchido. Quem
  escreve "LACAR 9001" nas Notas e se esquece da coluna perdia a peça; a
  análise lê as Notas (com tolerância a erros de escrita) e avisa.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

FOLHA_LACAGEM = "Lacagem"
FOLHA_LISTAGEM = "LISTAGEM_CUT_RITE"

# Os cinco valores da lista de validação da coluna Tipo_Lacagem do modelo.
TIPOS_ACABAMENTO = (
    "A-1 Face + Topos",
    "B-2 Face + Topos",
    "c-1 Face + Topos + 50mm Contraface",
    "D-1 Face",
    "E- Outro",
)

# Palavras que indicam acabamento, sem acentos e em maiúsculas.
PALAVRAS_ACABAMENTO = (
    "ACABAMENTO",
    "ACABAMENTOS",
    "PINTURA",
    "PINTAR",
    "PINTADO",
    "PINTADA",
    "LACAGEM",
    "LACAR",
    "LACADO",
    "LACADA",
    "LACADOS",
    "LACADAS",
    "LAQUEAR",
    "LAQUEADO",
    "VERNIZ",
    "ENVERNIZAR",
    "ENVERNIZADO",
    "ENVERNIZADA",
    "ENVERNIZAMENTO",
    "VELATURA",
)
# Um início de palavra destes chega (LACADAS, ENVERNIZADOS, PINTURAS...).
_RAIZES_ACABAMENTO = (
    "LACAG",
    "LACAD",
    "LACAR",
    "LAQUE",
    "VERNI",
    "ENVERN",
    "PINTUR",
    "PINTAD",
    "VELATUR",
    "ACABAMENT",
)
# "S/ LACAGEM", "SEM VERNIZ", "NAO LACAR": a palavra está lá, mas a negar.
_NEGACOES = frozenset({"S", "SEM", "NAO", "N"})
_RAL = re.compile(r"\bRAL\s*\d{4}\b")


def _sem_acentos(texto: str) -> str:
    decomposto = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in decomposto if not unicodedata.combining(c))


def _texto(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def chave_tipo(valor: object) -> str:
    """Forma de comparar tipos: sem acentos, maiúsculas e espaços."""
    return re.sub(r"\s+", "", _sem_acentos(_texto(valor))).upper()


_TIPOS_POR_CHAVE = {chave_tipo(tipo): tipo for tipo in TIPOS_ACABAMENTO}


def tipo_acabamento(valor: object) -> str:
    """Devolve o tipo da lista a que o valor corresponde, ou ``""``."""
    return _TIPOS_POR_CHAVE.get(chave_tipo(valor), "")


def _distancia(a: str, b: str) -> int:
    """Distância de edição com troca de letras vizinhas (Damerau/OSA)."""
    anterior2: list[int] | None = None
    anterior = list(range(len(b) + 1))
    for i, ca in enumerate(a, start=1):
        atual = [i] + [0] * len(b)
        for j, cb in enumerate(b, start=1):
            custo = 0 if ca == cb else 1
            atual[j] = min(anterior[j] + 1, atual[j - 1] + 1, anterior[j - 1] + custo)
            if (
                anterior2 is not None
                and i > 1
                and j > 1
                and ca == b[j - 2]
                and a[i - 2] == cb
            ):
                atual[j] = min(atual[j], anterior2[j - 2] + 1)
        anterior2, anterior = anterior, atual
    return anterior[-1]


def _parece_acabamento(palavra: str) -> bool:
    if len(palavra) < 4:
        return False
    if palavra.startswith(_RAIZES_ACABAMENTO):
        return True
    # Erros de escrita: uma letra trocada/esquecida (duas nas palavras
    # compridas). A primeira letra tem de bater, para "PLACA" não passar
    # por "LACAR".
    tolerancia = 2 if len(palavra) >= 9 else 1
    return any(
        palavra[0] == chave[0]
        and abs(len(palavra) - len(chave)) <= tolerancia
        and _distancia(palavra, chave) <= tolerancia
        for chave in PALAVRAS_ACABAMENTO
    )


def palavra_de_acabamento(texto: object) -> str:
    """A primeira palavra do texto que indica acabamento, ou ``""``.

    Tolera erros de escrita ("LAKAGEM", "VERNIS", "ACABMENTO") e ignora as
    negações ("S/ LACAGEM", "SEM VERNIZ").
    """
    normal = _sem_acentos(_texto(texto)).upper()
    if not normal:
        return ""
    palavras = re.findall(r"[A-Z0-9]+", normal)
    for indice, palavra in enumerate(palavras):
        if not palavra.isalpha() or not _parece_acabamento(palavra):
            continue
        if indice > 0 and palavras[indice - 1] in _NEGACOES:
            continue
        return palavra
    encontrado = _RAL.search(normal)
    return encontrado.group(0) if encontrado else ""


@dataclass(frozen=True)
class PecaAcabamento:
    linha: int
    id: str
    descricao: str
    notas: str
    tipo: str
    quantidade: float


@dataclass(frozen=True)
class AvisoAcabamento:
    separador: str
    linha: int
    id: str
    descricao: str
    texto: str
    palavra: str

    def descrever(self) -> str:
        peca = f"ID {self.id}" if self.id else f"linha {self.linha}"
        nome = f" {self.descricao}" if self.descricao else ""
        return f"{self.separador} linha {self.linha} ({peca}{nome}): «{self.texto}»"


@dataclass(frozen=True)
class PlanoLacagem:
    """Onde está a tabela no separador Lacagem e quantas linhas lhe faltam."""

    linha_cabecalho: int
    primeira_linha: int
    ultima_linha: int
    primeira_coluna: int
    ultima_coluna: int
    coluna_legenda: int
    coluna_quant: int
    linhas_a_acrescentar: int = 0

    @property
    def ultima_linha_final(self) -> int:
        return self.ultima_linha + self.linhas_a_acrescentar

    @property
    def linha_total(self) -> int:
        return self.ultima_linha_final + 1


@dataclass
class AnaliseAcabamentos:
    tem_listagem: bool = False
    tem_lacagem: bool = False
    tem_coluna_tipo: bool = False
    pecas: list[PecaAcabamento] = field(default_factory=list)
    sem_tipo: list[AvisoAcabamento] = field(default_factory=list)
    tipo_invalido: list[AvisoAcabamento] = field(default_factory=list)
    cabecalho_em_falta: list[str] = field(default_factory=list)
    plano: PlanoLacagem | None = None
    erro_plano: str = ""

    @property
    def total_pecas(self) -> float:
        return sum(peca.quantidade for peca in self.pecas)

    @property
    def disponivel(self) -> bool:
        return self.tem_lacagem and bool(self.pecas) and self.plano is not None

    @property
    def motivo(self) -> str:
        if not self.tem_lacagem:
            return f"O separador {FOLHA_LACAGEM} não existe neste Excel."
        if not self.tem_listagem:
            return f"O separador {FOLHA_LISTAGEM} não existe neste Excel."
        if not self.tem_coluna_tipo:
            return f"A {FOLHA_LISTAGEM} não tem a coluna Tipo_Lacagem."
        if self.plano is None:
            return self.erro_plano or (
                f"Não encontrei a tabela de peças no separador {FOLHA_LACAGEM}."
            )
        if not self.pecas:
            return (
                "Nenhuma peça tem a coluna Tipo_Lacagem preenchida na "
                f"{FOLHA_LISTAGEM}."
            )
        total = _formatar_quantidade(self.total_pecas)
        return (
            f"{len(self.pecas)} linha(s) com acabamento, {total} peça(s) a "
            f"enviar ao fornecedor (separador {FOLHA_LACAGEM})."
        )

    def avisos(self) -> list[str]:
        linhas: list[str] = []
        if self.sem_tipo:
            linhas.append(
                "Peças com acabamento escrito nas Notas/Observações mas sem "
                "Tipo_Lacagem (não vão sair no PDF):"
            )
            linhas.extend(f"  - {aviso.descrever()}" for aviso in self.sem_tipo)
        if self.tipo_invalido:
            linhas.append(
                "Tipo_Lacagem com um valor que não é nenhum dos 5 da lista "
                "(não vai sair no PDF):"
            )
            linhas.extend(f"  - {aviso.descrever()}" for aviso in self.tipo_invalido)
        if self.cabecalho_em_falta:
            linhas.append(
                f"Cabeçalho do separador {FOLHA_LACAGEM} por preencher: "
                + ", ".join(self.cabecalho_em_falta)
                + "."
            )
        return linhas


def _formatar_quantidade(valor: float) -> str:
    return str(int(valor)) if float(valor).is_integer() else f"{valor:g}"


def _quantidade(valor: object) -> float:
    try:
        return float(valor or 0)
    except (TypeError, ValueError):
        return 0.0


def _chave_cabecalho(valor: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", _sem_acentos(_texto(valor)).upper())


def _colunas_listagem(linhas: list[tuple]) -> tuple[int, dict[str, int]] | None:
    """Linha do cabeçalho da LISTAGEM_CUT_RITE e índice (0-based) das colunas."""
    for numero, linha in enumerate(linhas[:6], start=1):
        chaves = {_chave_cabecalho(valor): indice for indice, valor in enumerate(linha)}
        if "DESCRICAO" in chaves and ("NOTAS" in chaves or "TIPOLACAGEM" in chaves):
            return numero, chaves
    return None


def _ler_listagem(folha, analise: AnaliseAcabamentos) -> None:
    linhas = list(folha.iter_rows(values_only=True))
    cabecalho = _colunas_listagem(linhas)
    if cabecalho is None:
        return
    linha_cabecalho, colunas = cabecalho
    analise.tem_coluna_tipo = "TIPOLACAGEM" in colunas

    def valor(linha: tuple, chave: str) -> object:
        indice = colunas.get(chave)
        if indice is None or indice >= len(linha):
            return None
        return linha[indice]

    for numero, linha in enumerate(linhas[linha_cabecalho:], start=linha_cabecalho + 1):
        descricao = _texto(valor(linha, "DESCRICAO"))
        notas = _texto(valor(linha, "NOTAS"))
        tipo_bruto = _texto(valor(linha, "TIPOLACAGEM"))
        if not descricao and not notas and not tipo_bruto:
            continue
        ident = _texto(valor(linha, "ID"))
        if tipo_bruto and tipo_bruto != "0":
            if tipo_acabamento(tipo_bruto):
                analise.pecas.append(
                    PecaAcabamento(
                        numero,
                        ident,
                        descricao,
                        notas,
                        tipo_bruto,
                        _quantidade(valor(linha, "QT")),
                    )
                )
            else:
                analise.tipo_invalido.append(
                    AvisoAcabamento(
                        FOLHA_LISTAGEM, numero, ident, descricao, tipo_bruto, ""
                    )
                )
            continue
        palavra = palavra_de_acabamento(notas)
        if palavra and _quantidade(valor(linha, "QT")) > 0:
            analise.sem_tipo.append(
                AvisoAcabamento(FOLHA_LISTAGEM, numero, ident, descricao, notas, palavra)
            )


_ORIGEM_TIPO = re.compile(
    r"LISTAGEM_CUT_RITE'?!\$?Z\$?(\d+)", re.IGNORECASE
)
_CAMPOS_CABECALHO = (
    ("Tipo", lambda chave: chave.startswith("TIPO")),
    ("Acabamento", lambda chave: chave.startswith("ACABAMENTO") and "BRILHO" in chave),
    ("Cor", lambda chave: chave.startswith("COR")),
)


def _ler_lacagem(
    formulas, valores, analise: AnaliseAcabamentos, pecas_por_linha: dict[int, PecaAcabamento]
) -> None:
    linhas_f = list(formulas.iter_rows(values_only=True))
    linhas_v = list(valores.iter_rows(values_only=True))

    linha_cabecalho = 0
    colunas: dict[str, int] = {}
    for numero, linha in enumerate(linhas_v[:40], start=1):
        chaves = {_chave_cabecalho(v): indice for indice, v in enumerate(linha) if v}
        if "LEGENDA" in chaves and "QUANT" in chaves:
            linha_cabecalho, colunas = numero, chaves
            break
    if not linha_cabecalho:
        analise.erro_plano = (
            f"O separador {FOLHA_LACAGEM} não tem o cabeçalho com as colunas "
            "Quant e Legenda."
        )
        return

    # Cabeçalho que o utilizador preenche (Tipo / Acabamento / Cor).
    for linha in linhas_v[: linha_cabecalho - 1]:
        for indice, celula in enumerate(linha):
            chave = _chave_cabecalho(celula)
            for nome, corresponde in _CAMPOS_CABECALHO:
                if chave and corresponde(chave):
                    resto = [v for v in linha[indice + 1 :] if _texto(v)]
                    if not resto and nome not in analise.cabecalho_em_falta:
                        analise.cabecalho_em_falta.append(nome)

    col_legenda = colunas["LEGENDA"]
    col_obs = next(
        (indice for chave, indice in colunas.items() if chave.startswith("OBSERVA")),
        None,
    )
    # Linha da Lacagem -> linha da LISTAGEM_CUT_RITE, lida da fórmula da Legenda.
    origem: dict[int, int] = {}
    for numero, linha in enumerate(linhas_f[linha_cabecalho:], start=linha_cabecalho + 1):
        formula = linha[col_legenda] if col_legenda < len(linha) else None
        encontrado = _ORIGEM_TIPO.search(str(formula or ""))
        if encontrado:
            origem[numero] = int(encontrado.group(1))
    if not origem:
        analise.erro_plano = (
            f"O separador {FOLHA_LACAGEM} não tem as fórmulas que ligam a "
            f"Legenda à coluna Tipo_Lacagem da {FOLHA_LISTAGEM}."
        )
        return

    primeira, ultima = min(origem), max(origem)
    desvio = primeira - origem[primeira]
    linear = all(origem.get(n) == n - desvio for n in range(primeira, ultima + 1))
    em_falta = [
        peca for linha_origem, peca in pecas_por_linha.items()
        if linha_origem + desvio not in origem
    ]
    acrescentar = 0
    if em_falta:
        acima = [peca for peca in em_falta if peca.linha < origem[primeira]]
        if not linear or acima:
            analise.erro_plano = (
                f"O separador {FOLHA_LACAGEM} não cobre todas as linhas da "
                f"{FOLHA_LISTAGEM} (fórmulas fora de ordem). Peças que ficariam "
                "de fora: "
                + ", ".join(f"linha {peca.linha}" for peca in em_falta[:10])
            )
            return
        acrescentar = max(peca.linha for peca in em_falta) - origem[ultima]

    for numero in range(primeira, ultima + 1):
        if col_obs is None:
            break
        linha_v = linhas_v[numero - 1] if numero - 1 < len(linhas_v) else ()
        observacao = _texto(linha_v[col_obs]) if col_obs < len(linha_v) else ""
        linha_origem = origem.get(numero)
        if not observacao or linha_origem in pecas_por_linha:
            continue
        palavra = palavra_de_acabamento(observacao)
        if not palavra:
            continue
        ident = _texto(linha_v[colunas["ID"]]) if "ID" in colunas and colunas["ID"] < len(linha_v) else ""
        descricao = (
            _texto(linha_v[colunas["DESCRICAO"]])
            if "DESCRICAO" in colunas and colunas["DESCRICAO"] < len(linha_v)
            else ""
        )
        analise.sem_tipo.append(
            AvisoAcabamento(FOLHA_LACAGEM, numero, ident, descricao, observacao, palavra)
        )

    colunas_tabela = [indice for indice in colunas.values()]
    analise.plano = PlanoLacagem(
        linha_cabecalho=linha_cabecalho,
        primeira_linha=primeira,
        ultima_linha=ultima,
        primeira_coluna=min(colunas_tabela) + 1,
        ultima_coluna=max(colunas_tabela) + 1,
        coluna_legenda=col_legenda + 1,
        coluna_quant=colunas["QUANT"] + 1,
        linhas_a_acrescentar=acrescentar,
    )


def analisar_acabamentos(workbook_path: Path) -> AnaliseAcabamentos:
    """Lê a LISTAGEM_CUT_RITE e o separador Lacagem (só leitura)."""
    analise = AnaliseAcabamentos()
    valores = load_workbook(Path(workbook_path), read_only=True, data_only=True)
    try:
        nomes = set(valores.sheetnames)
        analise.tem_listagem = FOLHA_LISTAGEM in nomes
        analise.tem_lacagem = FOLHA_LACAGEM in nomes
        if analise.tem_listagem:
            _ler_listagem(valores[FOLHA_LISTAGEM], analise)
        if not analise.tem_lacagem:
            return analise
        formulas = load_workbook(Path(workbook_path), read_only=True, data_only=False)
        try:
            _ler_lacagem(
                formulas[FOLHA_LACAGEM],
                valores[FOLHA_LACAGEM],
                analise,
                {peca.linha: peca for peca in analise.pecas},
            )
        finally:
            formulas.close()
    finally:
        valores.close()
    return analise


# --------------------------------------------------------------------------
# Preparação do separador no Excel (COM), antes de exportar o PDF.
# O livro é aberto só de leitura e fechado sem gravar: nada disto fica no
# ficheiro do utilizador.
# --------------------------------------------------------------------------

XL_CALCULO_MANUAL = -4135
XL_LANDSCAPE = 2
XL_PAPER_A4 = 9
XL_FILTER_VALUES = 7
SUBTOTAL_SOMA_VISIVEIS = 109


def abrir_em_calculo_manual(excel) -> None:
    """Põe o Excel em cálculo manual ANTES de abrir a Lista Material.

    O Centro abre o livro com as macros desligadas. A Esp.Mat da
    LISTAGEM_CUT_RITE é uma função VBA (``EspMaterialMM``): se o Excel
    recalcular ao abrir, dá ``#NOME?`` e todas as linhas do separador
    Lacagem ficam em erro. Em cálculo manual ficam os valores que o
    utilizador gravou. O modo só se muda com um livro aberto, daí o livro
    em branco (que morre com o ``Quit``).
    """
    excel.Workbooks.Add()
    excel.Calculation = XL_CALCULO_MANUAL


def texto_rodape(valor: str) -> str:
    """O ``&`` é código nos rodapés do Excel; tem de ir dobrado."""
    return str(valor or "").replace("&", "&&")


def _definir(objeto, atributo: str, valor) -> None:
    try:
        setattr(objeto, atributo, valor)
    except Exception:
        pass


def _acrescentar_linhas(folha, plano: PlanoLacagem) -> None:
    if plano.linhas_a_acrescentar <= 0:
        return
    ultima = plano.ultima_linha
    c1, c2 = plano.primeira_coluna, plano.ultima_coluna
    folha.Rows(f"{ultima + 1}:{ultima + plano.linhas_a_acrescentar}").Insert()
    origem = folha.Range(folha.Cells(ultima, c1), folha.Cells(ultima, c2))
    # Valores escritos à mão (ex.: Observações) não se copiam para as novas.
    constantes = [
        coluna
        for coluna in range(c1, c2 + 1)
        if not folha.Cells(ultima, coluna).HasFormula
    ]
    for numero in range(ultima + 1, ultima + plano.linhas_a_acrescentar + 1):
        origem.Copy(folha.Range(folha.Cells(numero, c1), folha.Cells(numero, c2)))
        for coluna in constantes:
            folha.Cells(numero, coluna).ClearContents()


def _legendas(folha, plano: PlanoLacagem) -> list[tuple[int, str]]:
    coluna = plano.coluna_legenda
    valores = folha.Range(
        folha.Cells(plano.primeira_linha, coluna),
        folha.Cells(plano.ultima_linha_final, coluna),
    ).Value
    if not isinstance(valores, tuple):
        valores = ((valores,),)
    return [
        (plano.primeira_linha + indice, _texto(linha[0] if isinstance(linha, tuple) else linha))
        for indice, linha in enumerate(valores)
    ]


def _filtrar_legenda(folha, plano: PlanoLacagem) -> int:
    """Deixa visíveis só as peças com um dos 5 tipos. Devolve quantas."""
    if folha.AutoFilterMode:
        try:
            folha.ShowAllData()
        except Exception:
            pass
        folha.AutoFilterMode = False
    folha.Rows(f"{plano.primeira_linha}:{plano.ultima_linha_final}").Hidden = False

    legendas = _legendas(folha, plano)
    validas = [(linha, valor) for linha, valor in legendas if tipo_acabamento(valor)]
    if not validas:
        raise ValueError(
            "Nenhuma peça tem Tipo_Lacagem preenchido; não há nada para enviar "
            "ao fornecedor."
        )
    criterios = sorted({valor for _, valor in validas})
    # Um total escrito à mão por baixo da tabela faz o Excel estender o
    # filtro até ele e esconder essa linha; o total é reescrito a seguir.
    folha.Cells(plano.linha_total, plano.coluna_quant).ClearContents()
    tabela = folha.Range(
        folha.Cells(plano.linha_cabecalho, plano.primeira_coluna),
        folha.Cells(plano.ultima_linha_final, plano.ultima_coluna),
    )
    campo = plano.coluna_legenda - plano.primeira_coluna + 1
    try:
        tabela.AutoFilter(Field=campo, Criteria1=criterios, Operator=XL_FILTER_VALUES)
    except Exception:
        # Sem filtro automático, esconder as linhas imprime o mesmo.
        for linha, valor in legendas:
            if not tipo_acabamento(valor):
                folha.Rows(linha).Hidden = True
    folha.Rows(plano.linha_total).Hidden = False
    return len(validas)


def _escrever_total(folha, plano: PlanoLacagem) -> None:
    coluna = get_column_letter(plano.coluna_quant)
    celula = folha.Cells(plano.linha_total, plano.coluna_quant)
    celula.Formula = (
        f"=SUBTOTAL({SUBTOTAL_SOMA_VISIVEIS},{coluna}{plano.primeira_linha}:"
        f"{coluna}{plano.ultima_linha_final})"
    )
    celula.Font.Bold = True
    celula.Calculate()
    if plano.coluna_quant > plano.primeira_coluna:
        rotulo = folha.Cells(plano.linha_total, plano.coluna_quant - 1)
        if rotulo.Value in (None, ""):
            rotulo.Value = "Total"
            rotulo.Font.Bold = True
            _definir(rotulo, "HorizontalAlignment", -4152)  # à direita


def _area_impressao(folha, plano: PlanoLacagem) -> str:
    atual = str(getattr(folha.PageSetup, "PrintArea", "") or "")
    inicio = atual.split(":")[0] if atual else ""
    if not re.fullmatch(r"\$?[A-Z]{1,3}\$?\d+", inicio):
        inicio = f"${get_column_letter(plano.primeira_coluna)}$1"
    fim = f"${get_column_letter(plano.ultima_coluna)}${plano.linha_total}"
    return f"{inicio}:{fim}"


def _configurar_pagina(excel, folha, plano: PlanoLacagem, obra: str, dia: date) -> None:
    _definir(excel, "PrintCommunication", False)
    try:
        pagina = folha.PageSetup
        _definir(pagina, "PrintArea", _area_impressao(folha, plano))
        _definir(
            pagina,
            "PrintTitleRows",
            f"${plano.linha_cabecalho}:${plano.primeira_linha - 1}",
        )
        _definir(pagina, "Orientation", XL_LANDSCAPE)
        _definir(pagina, "PaperSize", XL_PAPER_A4)
        _definir(pagina, "Zoom", False)
        _definir(pagina, "FitToPagesWide", 1)
        _definir(pagina, "FitToPagesTall", False)
        _definir(pagina, "LeftFooter", texto_rodape(dia.strftime("%d-%m-%Y")))
        _definir(pagina, "CenterFooter", texto_rodape(obra))
        _definir(pagina, "RightFooter", "&P/&N")
    finally:
        _definir(excel, "PrintCommunication", True)


def preparar_separador_lacagem(
    excel,
    workbook,
    plano: PlanoLacagem,
    *,
    obra: str,
    dia: date | None = None,
) -> int:
    """Põe o separador Lacagem pronto a imprimir. Devolve as linhas visíveis.

    O livro tem de ter sido aberto em cálculo manual (ver
    ``abrir_em_calculo_manual``).
    """
    folha = workbook.Worksheets.Item(FOLHA_LACAGEM)
    _acrescentar_linhas(folha, plano)
    # Só este separador: as linhas acrescentadas (que vêm com os valores da
    # linha copiada) e um Tipo_Lacagem mudado sem o Excel ter recalculado
    # ficam certos. A LISTAGEM_CUT_RITE não se recalcula (função VBA).
    folha.Calculate()
    visiveis = _filtrar_legenda(folha, plano)
    _escrever_total(folha, plano)
    _configurar_pagina(excel, folha, plano, obra, dia or date.today())
    return visiveis
