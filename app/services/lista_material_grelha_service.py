"""Grelha da LISTAGEM_CUT_RITE: editar como no Excel e levar ao Excel só a diferença.

Nas obras lowcost o trabalho na listagem é de folha de cálculo: filtrar,
copiar linhas, colar antes/depois, eliminar, limpar. A grelha guarda o estado
final pretendido e, no fim, o Excel recebe apenas o que mudou:

1. células editadas (pelo número de linha ORIGINAL, antes de mexer na estrutura);
2. linhas eliminadas (de baixo para cima);
3. linhas novas, inseridas na posição final — pela tabela do Excel
   (`ListRows.Add`), para as colunas com fórmula (Ref_Cliente, Processo,
   Esp.Mat…) se preencherem sozinhas.

As colunas com fórmula nunca são escritas: são do Excel.
"""
from __future__ import annotations

import hashlib
import importlib
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.services import lista_material_excel_com as excel_com
from app.services.lista_material_assistente_service import (
    SHEET_CUTRITE, AssistantDecision, MaterialRow,
)

# Técnicas: nunca se editam à mão.
PROTEGIDAS = frozenset({"SourceID", "Estado_Assistente", "ID"})
# Numa linha colada, só estas não vêm da linha copiada. O ID do IMOS vem:
# liga a peça ao programa CNC.
NAO_COPIAR = frozenset({"SourceID", "Estado_Assistente"})
_NUMERO = re.compile(r"^-?\d+(?:[.,]\d+)?$")


def colunas_com_formula(path: Path) -> set[str]:
    """As colunas cuja 1.ª linha de dados é uma fórmula (ficam para o Excel)."""
    book = load_workbook(Path(path), read_only=True, data_only=False)
    try:
        if SHEET_CUTRITE not in book.sheetnames:
            return set()
        linhas = book[SHEET_CUTRITE].iter_rows(min_row=2, max_row=3, values_only=True)
        cabecalho = next(linhas, ())
        primeira = next(linhas, ())
        return {
            str(nome).strip()
            for nome, valor in zip(cabecalho, primeira)
            if nome and isinstance(valor, str) and valor.startswith("=")
        }
    finally:
        book.close()


@dataclass
class LinhaGrelha:
    original_row: int | None            # linha no Excel; None = linha nova
    original: dict[str, str]
    values: dict[str, str]
    removed: bool = False
    proposals: dict[str, object] = field(default_factory=dict)   # campo -> sugestão
    delete_proposal: object | None = None
    confirm: set[str] = field(default_factory=set)               # campos a confirmar

    @property
    def is_new(self) -> bool:
        return self.original_row is None


class GrelhaListagem:
    """Estado da grelha, sem Qt: é aqui que vivem as regras e os testes."""

    def __init__(self, columns, rows: list[MaterialRow], suggestions=(), formula_columns=()):
        self.columns = tuple(columns)
        self.formula_columns = set(formula_columns)
        self.rows: list[LinhaGrelha] = []
        by_number: dict[int, LinhaGrelha] = {}
        for row in rows:
            values = {col: str(row.values.get(col, "") or "") for col in self.columns}
            line = LinhaGrelha(row.row_number, dict(values), values)
            self.rows.append(line)
            by_number[row.row_number] = line
        self.suggestions = list(suggestions)
        for suggestion in self.suggestions:
            line = by_number.get(suggestion.row_number)
            if line is None:
                continue
            if suggestion.delete_row:
                line.delete_proposal = suggestion
                line.removed = True
            elif suggestion.blocking and not suggestion.suggested:
                line.confirm.add(suggestion.field)
                line.proposals.setdefault(suggestion.field, suggestion)
            elif suggestion.field in line.values and suggestion.field not in line.proposals:
                line.proposals[suggestion.field] = suggestion
                line.values[suggestion.field] = str(suggestion.suggested or "")
        self.clipboard: list[dict[str, str]] = []
        self._undo: list[list[LinhaGrelha]] = []
        self._redo: list[list[LinhaGrelha]] = []

    # ---- anular / refazer (Ctrl+Z / Ctrl+Y) --------------------------------
    # Pedido do Paulo ao testar (21-09-2026): colou uma linha e não conseguiu
    # voltar atrás. Cada operação guarda antes uma fotografia da grelha.

    LIMITE_ANULAR = 100

    def _fotografia(self) -> list[LinhaGrelha]:
        return [
            LinhaGrelha(line.original_row, line.original, dict(line.values), line.removed,
                        line.proposals, line.delete_proposal, set(line.confirm))
            for line in self.rows
        ]

    def checkpoint(self) -> None:
        self._undo.append(self._fotografia())
        del self._undo[:-self.LIMITE_ANULAR]
        self._redo.clear()

    def can_undo(self) -> bool:
        return bool(self._undo)

    def can_redo(self) -> bool:
        return bool(self._redo)

    def undo(self) -> bool:
        if not self._undo:
            return False
        self._redo.append(self._fotografia())
        self.rows = self._undo.pop()
        return True

    def redo(self) -> bool:
        if not self._redo:
            return False
        self._undo.append(self._fotografia())
        self.rows = self._redo.pop()
        return True

    # ---- consultas -----------------------------------------------------------

    def editable(self, column: str) -> bool:
        return column in self.columns and column not in PROTEGIDAS and column not in self.formula_columns

    def cell_state(self, index: int, column: str) -> str:
        line = self.rows[index]
        if line.removed:
            return "removida"
        if line.is_new:
            return "nova"
        if column in line.confirm and line.values.get(column) == line.original.get(column):
            return "confirmar"
        if line.values.get(column) == line.original.get(column):
            return "igual"
        proposal = line.proposals.get(column)
        if proposal is not None and line.values.get(column) == str(proposal.suggested or ""):
            return "proposta"
        return "editada"

    def row_state(self, index: int) -> str:
        line = self.rows[index]
        if line.removed:
            return "A remover" if line.original_row else "Removida"
        if line.is_new:
            return "Nova"
        states = {self.cell_state(index, c) for c in self.columns}
        if "confirmar" in states:
            return "Confirmar"
        if "editada" in states:
            return "Editada"
        if "proposta" in states:
            return "Proposta"
        return ""

    def has_changes(self, index: int) -> bool:
        return bool(self.row_state(index)) or bool(self.rows[index].confirm)

    # ---- operações tipo Excel ----------------------------------------------

    def set_value(self, index: int, column: str, value: str, *, checkpoint: bool = True) -> bool:
        if not self.editable(column) or self.rows[index].removed:
            return False
        value = str(value or "").strip()
        if self.rows[index].values.get(column, "") == value:
            return True
        if checkpoint:
            self.checkpoint()
        self.rows[index].values[column] = value
        return True

    def copy_rows(self, indexes) -> list[dict[str, str]]:
        self.clipboard = [dict(self.rows[i].values) for i in sorted(set(indexes))]
        return self.clipboard

    def paste_rows(self, index: int, *, after: bool, rows=None) -> list[int]:
        """Inserir cópias das linhas copiadas antes/depois de `index`; devolve as posições."""
        source = rows if rows is not None else self.clipboard
        if not source:
            return []
        position = index + 1 if after else index
        position = max(0, min(position, len(self.rows)))
        self.checkpoint()
        new_lines = []
        for values in source:
            fresh = {col: values.get(col, "") for col in self.columns}
            # Os campos técnicos e as fórmulas pertencem à linha nova, não à copiada.
            for col in NAO_COPIAR | self.formula_columns:
                if col in fresh:
                    fresh[col] = ""
            new_lines.append(LinhaGrelha(None, {}, fresh))
        self.rows[position:position] = new_lines
        return list(range(position, position + len(new_lines)))

    def delete_rows(self, indexes) -> None:
        if not indexes:
            return
        self.checkpoint()
        for i in sorted(set(indexes), reverse=True):
            if self.rows[i].is_new:
                del self.rows[i]
            else:
                self.rows[i].removed = True

    def clear_cells(self, cells) -> int:
        cells = [(i, c) for i, c in cells if self.editable(c) and not self.rows[i].removed
                 and self.rows[i].values.get(c, "")]
        if not cells:
            return 0
        self.checkpoint()   # uma só fotografia: Ctrl+Z desfaz a limpeza toda
        for index, column in cells:
            self.set_value(index, column, "", checkpoint=False)
        return len(cells)

    def restore_rows(self, indexes) -> None:
        if not indexes:
            return
        self.checkpoint()
        for i in sorted(set(indexes), reverse=True):
            line = self.rows[i]
            if line.is_new:
                continue
            line.removed = False
            line.values = dict(line.original)

    # ---- resultado ------------------------------------------------------------

    def edits(self) -> list[tuple[int, str, str]]:
        result = []
        for line in self.rows:
            if line.is_new or line.removed:
                continue
            for column in self.columns:
                if self.editable(column) and line.values.get(column, "") != line.original.get(column, ""):
                    result.append((line.original_row, column, line.values.get(column, "")))
        return result

    def deletions(self) -> list[int]:
        return [line.original_row for line in self.rows if line.removed and not line.is_new]

    def final_order(self) -> list[object]:
        """A ordem final: número da linha original, ou os valores de uma linha nova."""
        return [
            line.original_row if not line.is_new else {
                col: v for col, v in line.values.items()
                if col not in NAO_COPIAR and col not in self.formula_columns
            }
            for line in self.rows
            if not line.removed
        ]

    def summary(self) -> tuple[int, int, int]:
        return len(self.edits()), len(self.deletions()), sum(1 for line in self.rows if line.is_new and not line.removed)

    def decisions(self) -> list[AssistantDecision]:
        """Para o histórico/aprendizagem: o que se fez a cada proposta."""
        result = []
        for line in self.rows:
            if line.is_new:
                continue
            if line.delete_proposal is not None:
                result.append(AssistantDecision(line.delete_proposal, "aceitar" if line.removed else "rejeitar"))
            for column, suggestion in line.proposals.items():
                if line.removed:
                    continue
                final = line.values.get(column, "")
                if final == str(suggestion.suggested or "") and final != line.original.get(column, ""):
                    action = "aceitar"
                elif final == line.original.get(column, ""):
                    action = "rejeitar"
                else:
                    action = "editar"
                result.append(AssistantDecision(suggestion, action, final))
        return result


def _fingerprint(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _excel_value(value: str):
    text = str(value or "").strip()
    if _NUMERO.match(text):
        number = float(text.replace(",", "."))
        return int(number) if number.is_integer() else number
    return text


def aplicar_grelha(path: Path, grelha: GrelhaListagem, *, expected_hash: str, user_name: str) -> tuple[int, int, int]:
    """Leva ao Excel só a diferença da grelha. Devolve (células, eliminadas, novas)."""
    path = Path(path)
    if _fingerprint(path) != expected_hash:
        raise ValueError("O Excel mudou desde a análise. Volte a analisar antes de aplicar.")
    edits, deletions, order = grelha.edits(), grelha.deletions(), grelha.final_order()
    if not edits and not deletions and not any(isinstance(item, dict) for item in order):
        return 0, 0, 0
    excel = importlib.import_module("win32com.client").DispatchEx("Excel.Application")
    book = None
    try:
        excel_com.preparar_excel(excel)
        book = excel.Workbooks.Open(str(path.resolve()), UpdateLinks=0, ReadOnly=False)
        if book.ReadOnly:
            raise ValueError("Feche a Lista Material no Excel antes de aplicar.")
        sheet = book.Worksheets.Item(SHEET_CUTRITE)
        last_col = int(sheet.UsedRange.Column + sheet.UsedRange.Columns.Count)
        headers = {}
        for col in range(1, last_col + 1):
            name = str(sheet.Cells.Item(2, col).Value or "").strip()
            if name and name not in headers:
                headers[name] = col
        table = None
        try:
            if sheet.ListObjects.Count:
                table = sheet.ListObjects.Item(1)
        except Exception:
            table = None
        first_data = int(table.DataBodyRange.Row) if table is not None and table.DataBodyRange is not None else 3

        def write(row: int, column: str, value: str, *, new_row: bool = False) -> bool:
            col = headers.get(column)
            if not col:
                return False
            if not grelha.editable(column) and not (new_row and column not in NAO_COPIAR):
                return False
            cell = sheet.Cells.Item(row, col)
            if cell.HasFormula:
                return False
            cell.Value2 = _excel_value(value)
            return True

        written = sum(1 for row, column, value in edits if write(row, column, value))
        for row in sorted(deletions, reverse=True):
            sheet.Rows.Item(row).Delete()
        inserted = 0
        for position, item in enumerate(order):
            if not isinstance(item, dict):
                continue
            target = first_data + position
            if table is not None:
                table.ListRows.Add(position + 1)
            else:
                sheet.Rows.Item(target).Insert()
                for name, col in headers.items():
                    below = sheet.Cells.Item(target + 1, col)
                    if below.HasFormula:
                        sheet.Cells.Item(target, col).FormulaR1C1 = below.FormulaR1C1
            for column, value in item.items():
                write(target, column, value, new_row=True)
            inserted += 1

        try:
            log = book.Worksheets.Item("LOG")
        except Exception:
            log = book.Worksheets.Add(After=book.Worksheets.Item(book.Worksheets.Count))
            log.Name = "LOG"
        log_row = int(log.Cells(log.Rows.Count, 1).End(-4162).Row) + 1
        area = log.Range(log.Cells(log_row, 1), log.Cells(log_row, 8))
        area.NumberFormat = "@"
        area.Value = ((datetime.now().strftime("%Y-%m-%d %H:%M:%S"), user_name, "grelha", "",
                       "", "", "", f"{written} células, {len(deletions)} linhas eliminadas, "
                                   f"{inserted} linhas novas"),)
        excel_com.recalcular(excel)
        book.Save()
        return written, len(deletions), inserted
    finally:
        if book is not None:
            book.Close(False)
        excel.Quit()
