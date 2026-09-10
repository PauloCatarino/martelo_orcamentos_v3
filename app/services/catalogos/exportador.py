"""Escrever um Excel dos catálogos a partir da base — o caminho inverso.

Durante anos o Excel foi a fonte: alguém o mantinha à mão, com extractores, e o
Martelo lia-o. Desde a Fase 3 é a base que manda, e este módulo fecha o ciclo —
o Excel volta a existir, mas agora **é produzido**, e por isso está sempre de
acordo com o que a aplicação responde.

O que sai daqui é para **ler**, não para reimportar: um separador por tabela
ativa, na forma a que o Paulo está habituado (uma linha por referência, uma
coluna por espessura), mais um índice à frente que diz o que cada separador é e
de quando é. Nada disto substitui o ficheiro de origem — esse continua a ser a
porta de entrada de quem envia xlsx.

Duas coisas que o Excel de origem nunca conseguiu dizer e este diz:

* **o substrato canónico** ao lado do nome que o fornecedor usa, que é o que
  permite comparar um aglomerado hidrófugo entre três fornecedores;
* **onde a tabela se contradiz** — as 153 linhas em que a mesma referência, na
  mesma medida, tem dois preços. A coluna «Aviso» diz qual é o que vale, e
  porquê.

O ficheiro é gravado em ``_gerado/``, uma pasta que o indexador não varre: um
Excel gerado a partir da base, se voltasse ao índice, punha lá cada artigo duas
vezes.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.catalogos import FornTabelaPreco
from app.services.catalogos.consulta import (
    CHAVE_PRECO_UNITARIO,
    ArtigoConsulta,
    GrupoReferencia,
    agrupar_por_referencia,
    listar_artigos,
)

#: A pasta onde o ficheiro gerado vive, dentro da pasta dos catálogos.
PASTA_GERADOS = "_gerado"

#: O nome do ficheiro. Fixo de propósito: é sempre a mesma coisa, atualizada, e
#: um nome com data acumulava dezenas de cópias iguais na pasta de toda a gente.
FICHEIRO_GERADO = "12_Placas_Referencias_GERADO.xlsx"

_CABECALHO_PLACAS = (
    "Referência",
    "ST/Acab",
    "Nome Design",
    "Grupo",
    "Tipo Produto",
    "Substrato",
    "Substrato do fornecedor",
    "Fornecedor",
    "Fabricante",
    "Unidade",
)

_CABECALHO_FERRAGENS = (
    "Referência",
    "Descrição",
    "Família",
    "Secção",
    "Grupo",
    "Fornecedor",
    "Fabricante",
    "Unidade",
)

_CABECALHO_INDICE = (
    "Separador",
    "Fornecedor",
    "Fabricante",
    "Tabela",
    "Código",
    "Data da tabela",
    "Referências",
    "Preços",
    "Importada em",
)

_TITULO = Font(bold=True, size=13)
_CABECALHO = Font(bold=True)
_FUNDO_CABECALHO = PatternFill("solid", fgColor="E2E2D8")


@dataclass(frozen=True)
class ResultadoExportacao:
    """O que foi escrito, para se poder imprimir e acreditar."""

    caminho: str
    separadores: int
    referencias: int
    precos: int


def _nome_separador(folha: str, usados: set[str]) -> str:
    """Um nome de separador que o Excel aceite e que não se repita.

    O Excel recusa ``[]:*?/\\`` e para nos 31 caracteres. Os nomes de origem
    cabem todos, mas um fornecedor novo pode não caber — e um nome cortado que
    colida com outro perdia um separador inteiro sem dizer nada.
    """
    limpo = "".join("-" if caractere in r"[]:*?/\\" else caractere for caractere in folha)
    limpo = (limpo or "Sem nome")[:31]
    if limpo not in usados:
        usados.add(limpo)
        return limpo
    for sufixo in range(2, 100):
        tentativa = f"{limpo[: 31 - len(str(sufixo)) - 1]}~{sufixo}"
        if tentativa not in usados:
            usados.add(tentativa)
            return tentativa
    raise RuntimeError(f"Não consegui um nome de separador para {folha!r}")


def _etiquetas_ordenadas(grupos: list[GrupoReferencia]) -> list[str]:
    """As colunas de preço do separador, das mais finas para as mais grossas.

    Ordenadas pelo número e não pelo texto: por texto o ``10mm`` vinha antes do
    ``8mm`` e a tabela ficava impossível de ler de relance.
    """
    etiquetas: dict[str, Decimal] = {}
    for grupo in grupos:
        for etiqueta, artigo in grupo.por_etiqueta.items():
            etiquetas.setdefault(
                etiqueta,
                artigo.espessura_mm if artigo.espessura_mm is not None else Decimal(-1),
            )
    return sorted(etiquetas, key=lambda etiqueta: (etiquetas[etiqueta], etiqueta))


def _aviso(grupo: GrupoReferencia) -> str:
    """A linha de aviso desta referência, quando a tabela se contradiz."""
    partes: list[str] = []
    for etiqueta, artigo in grupo.por_etiqueta.items():
        if len(artigo.precos_da_medida) < 2:
            continue
        valores = " / ".join(f"{valor:.2f}" for valor in artigo.precos_da_medida)
        maximo = f"{artigo.precos_da_medida[-1]:.2f}"
        partes.append(
            f"{etiqueta}: a tabela dá {len(artigo.precos_da_medida)} preços "
            f"({valores}); vale o mais alto, {maximo}"
        )
    return " | ".join(partes)


def _e_de_placas(grupos: list[GrupoReferencia]) -> bool:
    return any(
        artigo.espessura
        for grupo in grupos
        for artigo in grupo.por_etiqueta.values()
    )


def _escrever_cabecalho(folha, cabecalho: tuple[str, ...], linha: int) -> None:
    for coluna, texto in enumerate(cabecalho, start=1):
        celula = folha.cell(row=linha, column=coluna, value=texto)
        celula.font = _CABECALHO
        celula.fill = _FUNDO_CABECALHO
        celula.alignment = Alignment(vertical="center", wrap_text=False)
    folha.freeze_panes = folha.cell(row=linha + 1, column=1)


def _larguras(folha, larguras: list[int]) -> None:
    for indice, largura in enumerate(larguras, start=1):
        folha.column_dimensions[get_column_letter(indice)].width = largura


def _linha_de_placa(grupo: GrupoReferencia) -> list[object]:
    principal = grupo.principal
    return [
        principal.referencia,
        principal.acabamento or "",
        principal.nome_design or principal.descricao,
        principal.grupo or "",
        principal.tipo,
        principal.substrato or "",
        principal.substrato_origem or "",
        principal.fornecedor,
        principal.fabricante or "",
        principal.unidade,
    ]


def _linha_de_ferragem(grupo: GrupoReferencia) -> list[object]:
    principal = grupo.principal
    return [
        principal.referencia,
        principal.descricao,
        principal.familia or "",
        principal.seccao or "",
        principal.grupo or "",
        principal.fornecedor,
        principal.fabricante or "",
        principal.unidade,
    ]


def _escrever_separador(
    livro: Workbook, nome: str, tabela: FornTabelaPreco | None, grupos: list[GrupoReferencia]
) -> int:
    """Um separador com uma linha por referência. Devolve quantos preços escreveu."""
    folha = livro.create_sheet(nome)
    principal = grupos[0].principal

    titulo = f"{principal.fornecedor} · {principal.tabela}"
    folha.cell(row=1, column=1, value=titulo).font = _TITULO
    detalhes = [f"Separador de origem: {principal.folha}"]
    if principal.data_tabela:
        detalhes.append(f"Data da tabela: {principal.data_tabela:%d/%m/%Y}")
    if tabela is not None and tabela.referencia_tabela:
        detalhes.append(f"Código: {tabela.referencia_tabela}")
    folha.cell(row=2, column=1, value=" · ".join(detalhes))

    de_placas = _e_de_placas(grupos)
    cabecalho = _CABECALHO_PLACAS if de_placas else _CABECALHO_FERRAGENS
    etiquetas = _etiquetas_ordenadas(grupos)
    if de_placas:
        colunas_preco = [f"Preço Tabela {etiqueta}" for etiqueta in etiquetas]
    else:
        colunas_preco = [CHAVE_PRECO_UNITARIO]
        etiquetas = [CHAVE_PRECO_UNITARIO]
    _escrever_cabecalho(folha, (*cabecalho, *colunas_preco, "Aviso"), 4)

    escritos = 0
    for indice, grupo in enumerate(grupos, start=5):
        valores = _linha_de_placa(grupo) if de_placas else _linha_de_ferragem(grupo)
        for etiqueta in etiquetas:
            artigo = grupo.por_etiqueta.get(etiqueta)
            preco = artigo.preco if artigo is not None else None
            valores.append(float(preco) if preco is not None else None)
            escritos += preco is not None
        valores.append(_aviso(grupo))
        for coluna, valor in enumerate(valores, start=1):
            celula = folha.cell(row=indice, column=coluna, value=valor)
            if coluna > len(cabecalho) and coluna <= len(cabecalho) + len(etiquetas):
                celula.number_format = "#,##0.00 €"

    _larguras(
        folha,
        [16, 10, 26, 14, 26, 12, 20, 18, 12, 9]
        if de_placas
        else [16, 44, 22, 12, 12, 14, 12, 9],
    )
    folha.auto_filter.ref = (
        f"A4:{get_column_letter(len(cabecalho) + len(etiquetas) + 1)}"
        f"{4 + len(grupos)}"
    )
    return escritos


def _escrever_indice(livro: Workbook, resumo: list[list[object]], quando: datetime) -> None:
    folha = livro.create_sheet("Índice", 0)
    folha.cell(row=1, column=1, value="Catálogos de fornecedores").font = _TITULO
    folha.cell(
        row=2,
        column=1,
        value=(
            f"Gerado a partir da base martelo_catalogos em {quando:%d/%m/%Y %H:%M}. "
            "Só as tabelas ativas. Este ficheiro é para consultar — o original "
            "continua a ser a porta de entrada de quem envia xlsx."
        ),
    )
    _escrever_cabecalho(folha, _CABECALHO_INDICE, 4)
    for indice, linha in enumerate(resumo, start=5):
        for coluna, valor in enumerate(linha, start=1):
            celula = folha.cell(row=indice, column=coluna, value=valor)
            # Sem formato, uma data de tabela aparecia como «2026-04-20
            # 00:00:00», que ninguém escreve assim em lado nenhum.
            if _CABECALHO_INDICE[coluna - 1] == "Data da tabela":
                celula.number_format = "DD/MM/YYYY"
            elif _CABECALHO_INDICE[coluna - 1] == "Importada em":
                celula.number_format = "DD/MM/YYYY HH:MM"
    _larguras(folha, [30, 20, 14, 34, 10, 14, 12, 10, 18])


def _tabelas_por_folha(session: Session) -> dict[str, FornTabelaPreco]:
    """A tabela ativa de cada separador, para o índice ir buscar o código.

    O ``ficheiro_origem`` é ``ficheiro.xlsx#Separador`` — é dali que sai a
    ligação entre o que o artigo diz (``atributos["folha"]``) e a linha da
    tabela.
    """
    fora: dict[str, FornTabelaPreco] = {}
    for tabela in session.scalars(
        select(FornTabelaPreco).where(FornTabelaPreco.ativa.is_(True)).order_by(
            FornTabelaPreco.id
        )
    ):
        origem = tabela.ficheiro_origem or ""
        if "#" in origem:
            fora[origem.split("#", 1)[1].split(" ", 1)[0]] = tabela
    return fora


def caminho_por_omissao(pasta_catalogos: str | Path) -> Path:
    """Onde o ficheiro é gravado quando ninguém diz outra coisa."""
    return Path(pasta_catalogos) / PASTA_GERADOS / FICHEIRO_GERADO


def exportar(session: Session, destino: str | Path) -> ResultadoExportacao:
    """Escreve o Excel dos catálogos. Substitui o que lá estiver do anterior.

    Não toca no ficheiro de origem nem em nada dentro da base: lê e escreve um
    ficheiro seu.
    """
    artigos = listar_artigos(session)
    if not artigos:
        raise RuntimeError(
            "A base de catalogos esta vazia. Corra "
            "'python -m scripts.importar_catalogos --fornecedor todos' primeiro."
        )

    por_folha: dict[str, list[ArtigoConsulta]] = {}
    for artigo in artigos:
        por_folha.setdefault(artigo.folha or "(sem separador)", []).append(artigo)

    tabelas = _tabelas_por_folha(session)
    livro = Workbook()
    livro.remove(livro.active)

    usados: set[str] = set()
    resumo: list[list[object]] = []
    referencias = 0
    precos = 0
    for folha, deste in por_folha.items():
        grupos = agrupar_por_referencia(deste)
        nome = _nome_separador(folha, usados)
        tabela = tabelas.get(folha)
        precos += _escrever_separador(livro, nome, tabela, grupos)
        referencias += len(grupos)
        principal = grupos[0].principal
        resumo.append(
            [
                nome,
                principal.fornecedor,
                principal.fabricante or "",
                principal.tabela,
                (tabela.referencia_tabela if tabela is not None else "") or "",
                principal.data_tabela or "",
                len(grupos),
                sum(1 for artigo in deste if artigo.preco is not None),
                tabela.importada_em if tabela is not None else "",
            ]
        )

    _escrever_indice(livro, resumo, datetime.now())

    caminho = Path(destino)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    livro.save(caminho)
    return ResultadoExportacao(
        caminho=str(caminho),
        separadores=len(por_folha),
        referencias=referencias,
        precos=precos,
    )
