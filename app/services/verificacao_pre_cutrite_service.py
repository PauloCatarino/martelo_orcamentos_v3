"""Verificação da LISTAGEM_CUT_RITE antes de a mandar para o Cut-Rite.

Duas coisas que o Cut-Rite não avisa e que só se descobrem na fábrica:

* **Material que o Woodstore não conhece** — o Cut-Rite não corta essas peças.
  Não se bloqueia o envio: na prática corta-se o resto e as peças que ficaram
  de fora vão numa versão nova do plano, depois de o material ser criado no
  Woodstore (ou com um nome temporário, para materiais de uma ou duas vezes).
* **Células com erro do Excel** (`#NAME?`...) nas colunas que seguem para o
  Cut-Rite — o envio lê os valores guardados no ficheiro, e um erro guardado
  vai tal e qual para o plano (Ref_Cliente e Processo com `#NAME?`).

Só leitura: nada aqui altera o Excel nem o Woodstore.
"""
from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from app.services.cutrite_service import (
    CUTRITE_EXPORT_HEADERS,
    CUTRITE_SOURCE_DATA_START_ROW,
    CUTRITE_SOURCE_HEADER_ROW,
    LISTAGEM_CUT_RITE_SHEET,
)

# Em inglês e em português: depende da língua do Excel que gravou o livro.
ERROS_EXCEL = frozenset({
    "#NAME?", "#NOME?", "#N/A", "#N/D", "#VALUE!", "#VALOR!", "#REF!",
    "#DIV/0!", "#NUM!", "#NÚM!", "#NULL!", "#NULO!", "#CALC!", "#SPILL!",
})


@dataclass(frozen=True)
class MaterialEmFalta:
    material: str
    pecas: int
    linhas: int
    descricoes: tuple[str, ...]


@dataclass(frozen=True)
class ColunaComErro:
    coluna: str
    linhas: int
    exemplo: str


@dataclass
class VerificacaoPreCutRite:
    materiais_em_falta: list[MaterialEmFalta] = field(default_factory=list)
    colunas_com_erro: list[ColunaComErro] = field(default_factory=list)
    woodstore_verificado: bool = True
    woodstore_aviso: str = ""

    @property
    def tudo_certo(self) -> bool:
        return (
            not self.materiais_em_falta
            and not self.colunas_com_erro
            and self.woodstore_verificado
        )


def _inteiro(valor) -> int:
    try:
        return int(float(str(valor).replace(",", ".")))
    except (TypeError, ValueError):
        return 0


def ler_linhas(path: Path) -> list[dict]:
    """As linhas da LISTAGEM_CUT_RITE, só com as colunas que vão para o Cut-Rite."""
    book = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        if LISTAGEM_CUT_RITE_SHEET not in book.sheetnames:
            raise ValueError(f"Folha {LISTAGEM_CUT_RITE_SHEET} em falta.")
        sheet = book[LISTAGEM_CUT_RITE_SHEET]
        colunas = len(CUTRITE_EXPORT_HEADERS)
        cabecalho = next(
            sheet.iter_rows(
                min_row=CUTRITE_SOURCE_HEADER_ROW,
                max_row=CUTRITE_SOURCE_HEADER_ROW,
                max_col=colunas,
                values_only=True,
            ),
            (),
        )
        # Os nomes do próprio livro mandam; o "+comp" do modelo traz uma
        # mudança de linha no fim.
        nomes = [
            str(nome or "").strip() or CUTRITE_EXPORT_HEADERS[i]
            for i, nome in enumerate(list(cabecalho) + [None] * (colunas - len(cabecalho)))
        ]
        linhas = []
        for valores in sheet.iter_rows(
            min_row=CUTRITE_SOURCE_DATA_START_ROW, max_col=colunas, values_only=True
        ):
            if not any(v not in (None, "") for v in valores):
                continue
            linhas.append(dict(zip(nomes, valores)))
        return linhas
    finally:
        book.close()


def verificar(
    linhas: list[dict],
    codigos_woodstore: set[str] | None,
    *,
    woodstore_aviso: str = "",
) -> VerificacaoPreCutRite:
    """`codigos_woodstore=None` quer dizer que não foi possível ler o Woodstore."""
    resultado = VerificacaoPreCutRite()

    erros: dict[str, list[str]] = {}
    for linha in linhas:
        for coluna, valor in linha.items():
            if isinstance(valor, str) and valor.strip().upper() in ERROS_EXCEL:
                erros.setdefault(coluna, []).append(valor.strip())
    resultado.colunas_com_erro = [
        ColunaComErro(coluna, len(valores), valores[0])
        for coluna, valores in erros.items()
    ]

    if codigos_woodstore is None:
        resultado.woodstore_verificado = False
        resultado.woodstore_aviso = woodstore_aviso or (
            "Não foi possível ler o Woodstore; os materiais não foram verificados."
        )
        return resultado

    codigos = {str(c or "").strip() for c in codigos_woodstore} - {""}
    pecas: Counter = Counter()
    contagem: Counter = Counter()
    descricoes: dict[str, set[str]] = {}
    for linha in linhas:
        material = str(linha.get("Material") or "").strip()
        if not material or material in codigos:
            continue
        pecas[material] += _inteiro(linha.get("Qt"))
        contagem[material] += 1
        descricao = str(linha.get("Descricao") or "").strip()
        if descricao:
            descricoes.setdefault(material, set()).add(descricao)
    resultado.materiais_em_falta = [
        MaterialEmFalta(
            material,
            pecas[material],
            contagem[material],
            tuple(sorted(descricoes.get(material, ()))),
        )
        for material in sorted(contagem)
    ]
    return resultado


def texto_do_aviso(resultado: VerificacaoPreCutRite, *, pode_reparar: bool = True) -> str:
    """O texto que o utilizador lê antes de decidir se envia."""
    partes: list[str] = []
    if resultado.materiais_em_falta:
        total = sum(m.pecas for m in resultado.materiais_em_falta)
        partes.append(
            f"MATERIAL QUE NÃO EXISTE NO WOODSTORE — {total} peça(s) NÃO vão ser "
            "cortadas neste plano:"
        )
        for m in resultado.materiais_em_falta:
            que = f" ({', '.join(m.descricoes[:3])})" if m.descricoes else ""
            partes.append(f"   • {m.material} — {m.pecas} peça(s) em {m.linhas} linha(s){que}")
        partes.append(
            "\nO habitual: enviar o resto agora e cortar estas peças numa versão "
            "nova do plano, depois de o material ser criado no Woodstore. Para um "
            "material de uma ou duas vezes, pode usar um nome temporário. Se a "
            "peça é comprada ou cortada à parte, pode ignorar este aviso."
        )
    if resultado.colunas_com_erro:
        if partes:
            partes.append("")
        partes.append(
            "CÉLULAS COM ERRO DO EXCEL — iam para o Cut-Rite tal e qual:"
        )
        for c in resultado.colunas_com_erro:
            partes.append(f"   • {c.coluna}: {c.exemplo} em {c.linhas} linha(s)")
        partes.append(
            "\nAcontece quando o livro é gravado com as macros desligadas."
            + (
                " «Reparar e enviar» recalcula a lista com as macros e grava-a "
                "antes do envio."
                if pode_reparar
                else ""
            )
        )
    if not resultado.woodstore_verificado:
        if partes:
            partes.append("")
        partes.append(resultado.woodstore_aviso)
    return "\n".join(partes)
