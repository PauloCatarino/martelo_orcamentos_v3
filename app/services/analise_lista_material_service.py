"""Análise por obra: Woodstore só de leitura e custos com preços congelados.

O layout MAT1 V12/2.15 foi reconciliado com as páginas 1–2 do PDF 0722:
campo 11 = superfície das placas usadas, incluindo restos que saem da obra.
Não é o campo 8 (área das peças), nem se deduz o campo 21 (restos).
"""
from __future__ import annotations

import csv
import hashlib
import importlib
import io
import json
import re
import shutil
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, ROUND_CEILING
from difflib import SequenceMatcher
from pathlib import Path
from uuid import uuid4
from types import SimpleNamespace

from openpyxl import load_workbook

from app.services import lista_material_excel_com as excel_com


def nominal_thickness(value):
    value = number(value)
    return value.quantize(Decimal('1'), rounding=ROUND_HALF_UP) if value is not None and value > 0 else None


def number(value):
    try:
        result = Decimal(str(value).strip().replace(",", "."))
        return result if result.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def fingerprint(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_material_inputs(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if 'LISTAGEM_CUT_RITE' not in book.sheetnames:
            raise ValueError('Folha LISTAGEM_CUT_RITE em falta.')
        rows = book['LISTAGEM_CUT_RITE'].iter_rows(min_row=2, values_only=True)
        header = next(rows, ())
        if 'Material' not in header or 'Qt' not in header:
            raise ValueError('Cabeçalhos Material/Qt em falta.')
        result = []
        for row in rows:
            values = dict(zip(header, row))
            if not values.get('Material') and not values.get('Descricao'):
                continue
            result.append(SimpleNamespace(material=str(values.get('Material') or ''),
                          quantity=number(values.get('Qt')), values=values))
        return result
    finally:
        book.close()


def writable_workbook(path: Path) -> Path:
    path = Path(path).resolve()
    if any(part.upper() == "ORIGINAL" for part in path.parts):
        raise ValueError("A pasta ORIGINAL é exclusivamente de consulta.")
    if not path.is_file():
        raise ValueError("Lista Material não encontrada.")
    return path


def backup_path(path, label):
    folder = path.parent / 'Analise_Lista_Material' / 'Copias'
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f'{path.stem}_{label}_{uuid4().hex[:12]}{path.suffix}'


@dataclass(frozen=True)
class PlanName:
    order: str
    version: str
    plan: str
    year: str
    client: str

    @classmethod
    def parse(cls, name):
        match = re.fullmatch(r"(\d+)_(\d{2})_(\d{2})_(\d{2})_(.+)", name)
        if not match:
            raise ValueError(f"Nome de plano inválido: {name}")
        return cls(*match.groups())

    @property
    def version_key(self):
        return f"{self.order}_{self.version}_{self.year}_{self.client}"


def read_plan(path: Path) -> dict:
    raw = path.read_bytes()
    rows = list(csv.reader(io.StringIO(raw.decode("cp1252"))))
    identity = PlanName.parse(path.stem)
    if not rows or rows[0][:3] != ["VER", "V12.00.5.1", "2.15"]:
        raise ValueError(f"Formato PTN por validar: {path.name}")
    summaries = [r for r in rows if r and r[0] == "SUM1"]
    names = [r for r in rows if r and r[0] == "FN1"]
    if len(summaries) != 1 or len(names) != 1 or names[0][1] != path.stem:
        raise ValueError(f"Identificação/resumo PTN inválido: {path.name}")
    materials = []
    seen = set()
    for row in rows:
        if not row or row[0] != "MAT1":
            continue
        if len(row) < 12 or not row[1] or row[1] in seen:
            raise ValueError(f"Material PTN inválido/duplicado: {path.name}")
        area = number(row[11])
        if area is None or area < 0:
            raise ValueError(f"Consumo inválido: {path.name} / {row[1]}")
        seen.add(row[1])
        materials.append({"material": row[1], "area": str(area)})
    total = number(summaries[0][2])
    if not materials or total is None or abs(sum(Decimal(r['area']) for r in materials) - total) > Decimal("0.01") * len(materials):
        raise ValueError(f"Consumos não reconciliados: {path.name}")
    return {"name": path.stem, "version": identity.version_key,
            "hash": hashlib.sha256(raw).hexdigest(), "path": str(path),
            "total": str(total), "materials": materials}


def discover_plans(folder: Path, plan_name: str):
    identity = PlanName.parse(plan_name)
    if not folder.is_dir():
        raise ValueError(f"Pasta de dados Cut-Rite indisponível: {folder}")
    names = set()
    for path in folder.glob(f"{identity.order}_{identity.version}_*_{identity.year}_{identity.client}.*"):
        try:
            other = PlanName.parse(path.stem)
        except ValueError:
            continue
        if other.version_key == identity.version_key:
            names.add(path.stem)
    plans, warnings = [], []
    for name in sorted(names):
        path = folder / (name + ".ptn")
        if not path.is_file():
            warnings.append(f"{name}: sem resultado PTN; consumo por apurar.")
            continue
        try:
            plans.append(read_plan(path))
        except (ValueError, IndexError, UnicodeError) as exc:
            warnings.append(str(exc))
    if not names:
        warnings.append("Ainda não existem planos de corte para esta versão.")
    return plans, warnings


def material_traits(code):
    code = str(code).upper().strip()
    tokens = tuple(t for t in re.split(r"[_\s/]+", code) if t)
    thickness = re.search(r"(\d+(?:[.,]\d+)?)\s*MM$", code)
    family = next((t for t in tokens if t in {"AGL", "MDF", "FENOLICO", "CONTRAPLACADO", "OSB"}), "")
    decor = tuple(t for t in tokens if re.search(r"\d", t) and not t.endswith("MM"))
    finish = tuple(t for t in tokens if t in {"MLM", "CRU", "MR", "STD", "TERM", "FOL"})
    return tokens, number(thickness[1]) if thickness else None, family, decor, finish


def material_candidates(original, catalog, thickness=None):
    tokens, parsed, family, decor, finish = material_traits(original)
    thickness = nominal_thickness(thickness if thickness is not None else parsed)
    candidates = []
    by_code = {}
    for row in catalog:
        code = str(row.get("Codigo") or "").strip()
        if code:
            by_code.setdefault(code, []).append(row)
    for code, records in by_code.items():
        ct, cp, cf, cd, cfinish = material_traits(code)
        stock_thicknesses = {nominal_thickness(r.get("Espessura")) for r in records} - {None}
        if cp is not None:
            stock_thicknesses.add(nominal_thickness(cp))
        reasons = []
        compatible = True
        if thickness is not None and stock_thicknesses and thickness not in stock_thicknesses:
            compatible = False
            reasons.append("Espessura diferente")
        if family and cf and family != cf:
            compatible = False
            reasons.append("Tipo de placa diferente")
        if decor != cd:
            reasons.append("Confirmar decoração/acabamento")
        if finish != cfinish:
            reasons.append("Confirmar composição/acabamento")
        ratio = SequenceMatcher(None, str(original).upper(), code.upper()).ratio()
        overlap = len(set(tokens) & set(ct)) / max(1, len(set(tokens) | set(ct)))
        score = .6 * ratio + .4 * overlap
        if compatible and score >= .40:
            reasons.insert(0, "Espessura/tipo compatíveis" if thickness is not None and stock_thicknesses and family and cf else "Características por confirmar")
            candidates.append({"code": code, "score": score, "reason": "; ".join(reasons)})
    return sorted(candidates, key=lambda x: (-x["score"], x["code"]))[:8]


def apply_material_codes(path, expected_hash, replacements, user_name):
    """Só altera Material e acrescenta log; eventos desativados, fórmulas recalculadas."""
    path = writable_workbook(path)
    if fingerprint(path) != expected_hash:
        raise ValueError("O Excel mudou desde a análise. Volte a analisar antes de aplicar.")
    if not replacements:
        return 0
    excel = importlib.import_module("win32com.client").DispatchEx("Excel.Application")
    book = None
    try:
        excel_com.preparar_excel(excel)
        book = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)
        if book.ReadOnly:
            raise ValueError("Guarde e feche o Excel antes de aplicar as correções.")
        if fingerprint(path) != expected_hash:
            raise ValueError("O Excel mudou durante a abertura. Repita a análise.")
        sheet = book.Worksheets.Item("LISTAGEM_CUT_RITE")
        headers = {str(sheet.Cells.Item(2, c).Value or '').strip(): c
                   for c in range(1, int(sheet.UsedRange.Column + sheet.UsedRange.Columns.Count))}
        column = headers.get("Material")
        if not column:
            raise ValueError("Coluna Material em falta.")
        changes = []
        for r in range(3, int(sheet.UsedRange.Row + sheet.UsedRange.Rows.Count)):
            old = str(sheet.Cells.Item(r, column).Value or '')
            new = replacements.get(old)
            if new and new != old:
                changes.append((r, old, new))
        if not changes:
            return 0
        backup = backup_path(path, 'antes_analise')
        book.SaveCopyAs(str(backup))
        try:
            log = book.Worksheets.Item("LOG_MATERIAIS_V3")
        except Exception:
            log = book.Worksheets.Add(After=book.Worksheets.Item(book.Worksheets.Count))
            log.Name = "LOG_MATERIAIS_V3"
            log.Range("A1:E1").Value = (("Data", "Utilizador", "Linha", "Original", "Materialcode"),)
        log_row = max(2, int(log.UsedRange.Row + log.UsedRange.Rows.Count))
        for r, old, new in changes:
            sheet.Cells.Item(r, column).NumberFormat = '@'
            sheet.Cells.Item(r, column).Value2 = new
            log.Range(f'A{log_row}:E{log_row}').NumberFormat = '@'
            log.Range(f"A{log_row}:E{log_row}").Value = ((datetime.now().isoformat(timespec='seconds'), user_name, r, old, new),)
            log_row += 1
        excel_com.recalcular(excel)
        book.Save()
        return len(changes)
    finally:
        if book is not None:
            book.Close(False)
        excel.Quit()


def cost_line(kind, key, name, quantity, unit, **extra):
    return {"kind": kind, "key": key, "name": name, "quantity": str(quantity) if quantity is not None else None,
            "unit": unit, **extra}


def board_cost_lines(plans):
    totals = {}
    seen = set()
    version = None
    for plan in plans:
        if version is not None and version != plan['version']:
            raise ValueError("Não pode misturar versões da obra.")
        version = plan['version']
        if plan['name'] in seen:
            raise ValueError("Plano de corte duplicado.")
        seen.add(plan['name'])
        for row in plan['materials']:
            totals[row['material']] = totals.get(row['material'], Decimal(0)) + Decimal(row['area'])
    return [cost_line("Placas", "placa:" + k, k, v, "m2", thickness=str(nominal_thickness(material_traits(k)[1]) or '')) for k, v in sorted(totals.items())]


def edge_width(value):
    text = str(value or '').upper().replace(',', '.')
    match = re.search(r"H\s*(\d+(?:\.\d+)?)", text)
    if match:
        return number(match[1])
    match = re.match(r"\s*(\d+(?:\.\d+)?)\s*[X×]", text)
    return number(match[1]) if match else None


def hardware_sheet_names(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return [name for name in book.sheetnames if 'custo_obra_ferragens' in name.lower()]
    finally:
        book.close()


def hardware_sources(path, version, output_folder=Path(r'C:\IMOS_Output_Batches')):
    """Only this job/version; never search ORIGINAL or other subfolders."""
    folder = Path(path).parent
    local = sorted({p for p in folder.glob(f'{version}_5_Custo_Obra_Ferragens*.xlsx') if p.is_file()})
    canonical = folder / '5_Custo_Obra_Ferragens.xlsx'
    if canonical.is_file():
        local.insert(0, canonical)
    if local:
        return local
    return sorted(p for p in Path(output_folder).glob(f'{version}_5_Custo_Obra_Ferragens*.xlsx') if p.is_file())


def workbook_cost_lines(path):
    book = load_workbook(path, read_only=True, data_only=True)
    result, warnings = [], []
    try:
        if 'ResumoOrlas' in book.sheetnames:
            sheet = book['ResumoOrlas']
            headers = None
            for row in sheet.values:
                row = tuple(row)
                if 'ML_QT' in row and 'Nome_Orlas' in row:
                    headers = {str(v): i for i, v in enumerate(row) if v}
                    continue
                if not headers:
                    continue
                def val(key):
                    i = headers.get(key, len(row))
                    return row[i] if i < len(row) else None
                if not val('Nome_Orlas'):
                    continue
                name, material, size = str(val('Nome_Orlas')), str(val('Material') or ''), str(val('LARG X ESP') or '')
                result.append(cost_line('Orlas', f'orla:{material}:{name}:{size}', name, number(val('ML_QT')), 'ml', width=str(edge_width(size) or ''), thickness=edge_thickness(name), board=material, size=size, source='ResumoOrlas'))
        else:
            warnings.append('ResumoOrlas em falta; custo de orlas por apurar.')
        # Desde 21-09-2026 as ferragens vêm dos separadores 1_FERRAGENS / 2_PURCH /
        # 3_SPP, que o utilizador corrige à mão; o 5_Custo_Obra_Ferragens (IMOS) só
        # dá o preço IMOS de referência e as cavilhas «fora da lista».
        from app.services.custo_ferragens_service import linhas_dos_separadores
        hardware, hardware_warnings = linhas_dos_separadores(book)
        result.extend(hardware)
        warnings.extend(hardware_warnings)
    finally:
        book.close()
    if not any(r['kind'] == 'Orlas' for r in result):
        edges, edge_warnings = calculated_edges(read_material_inputs(path))
        result = edges + result
        warnings.extend(edge_warnings)
        if edges:
            warnings.append('ResumoOrlas vazio: orlas calculadas das peças com a regra do Excel (8% incluído uma vez).')
    return result, warnings


def edge_thickness(name):
    match = re.search(r'(?:PVC|ABS|FOL|ORLA_LASER)[_:](\d+(?:[.,]\d+)?)', str(name), re.I)
    return str(number(match[1])) if match else ''


def calculated_edges(rows):
    """Regra de modResumoOrlas: Fix(mm), larguras por escalão, ceil(ml*1.08)."""
    totals, warnings = {}, []
    for row in rows:
        vals = row.values
        esp = nominal_thickness(vals.get('Esp.Final')) or nominal_thickness(vals.get('Esp.Mat')) or nominal_thickness(material_traits(row.material)[1]) or nominal_thickness(vals.get('Esp'))
        width = next((w for e, w in ((12,15),(16,19),(19,22),(22,25),(25,28),(30,33),(35,38),(40,43),(45,48),(53,55)) if esp is not None and esp <= e), 60) if esp else None
        material = row.material
        if esp and material_traits(material)[1] and nominal_thickness(material_traits(material)[1]) != esp:
            material = re.sub(r'\d+(?:[.,]\d+)?MM$', f'{esp}MM', material, flags=re.I)
        for field, dimension in (('Orla ESQ','Comp'),('Orla DIR','Comp'),('Orla CIMA','Larg'),('Orla BAIXO','Larg')):
            edge = str(vals.get(field) or '').strip()
            if not edge:
                continue
            if not re.match(r'^(PVC_|ABS[_:]|FOL_|ORLA_LASER_)', edge, re.I):
                warnings.append(f'Orla ignorada por não ser material: {edge}')
                continue
            length, qty = number(vals.get(dimension)), row.quantity
            key = (material, edge, width)
            if length is None or qty is None or length < 0 or qty < 0:
                totals[key] = None
            elif key not in totals or totals[key] is not None:
                totals[key] = totals.get(key, Decimal(0)) + int(length) * qty / 1000
    result = []
    for (material, edge, width), length in sorted(totals.items(), key=lambda x: str(x[0])):
        thickness = edge_thickness(edge)
        size = f'{width} x {thickness}' if width else ''
        qty = (length * Decimal('1.08')).to_integral_value(rounding=ROUND_CEILING) if length is not None else None
        result.append(cost_line('Orlas', f'orla:{material}:{edge}:{size}', edge, qty, 'ml', width=str(width or ''), thickness=thickness, board=material, size=size, source='Peças + regra Excel 8%'))
    return result, sorted(set(warnings))


def hardware_rows(rows):
    headers = None
    result = {}
    category = 'Ferragens'
    for index, row in enumerate(rows, 1):
        if 'Ref PHC' in row and 'Qt' in row:
            headers = {str(v): i for i, v in enumerate(row) if v}
            continue
        if not headers:
            continue
        if row and row[0] in ('FERRAGENS', 'SPP', 'PURCHASED PARTS'):
            category = {'FERRAGENS': 'Ferragens', 'SPP': 'SPP', 'PURCHASED PARTS': 'Comprados'}[row[0]]
            continue
        if row and row[0] in ('TOTAL', 'TOTAL DA OBRA'):
            continue
        def val(key):
            i = headers.get(key, len(row))
            return row[i] if i < len(row) else None
        # Os objetos comprados podem vir sem nome IMOS. Nunca os omitir do custo.
        if val('Qt') in (None, '') and not val('Un'):
            continue
        name = str(val('Nome iMos (Nome Uniao)') or val('Descricao') or 'Artigo sem nome')
        ref, description, item_unit = str(val('Ref PHC') or '').strip(), str(val('Descricao') or ''), str(val('Un') or '')
        key = f'ferragem:{category}:{name}:{ref}:{item_unit}:{description}'
        legacy_key = key
        if val('Jogo de Unioes (iMos)'):
            key += ':jogo:' + str(val('Jogo de Unioes (iMos)'))
        quantity = number(val('Qt'))
        if key in result:
            previous = number(result[key]['quantity'])
            result[key]['quantity'] = str(previous + quantity) if previous is not None and quantity is not None else None
            for field, header in (('length','Comp'),('width','Larg'),('thickness','Esp')):
                value = str(val(header) or '')
                previous_values = result[key].get(field, '').split(' / ')
                if value and value not in previous_values:
                    result[key][field] = ' / '.join(v for v in previous_values + [value] if v)
        else:
            result[key] = cost_line(category, key, name, quantity, item_unit, ref_phc=ref, description=description,
                length=str(val('Comp') or ''), width=str(val('Larg') or ''), thickness=str(val('Esp') or ''),
                union_name=str(val('Nome iMos (Nome Uniao)') or ''), union_set=str(val('Jogo de Unioes (iMos)') or ''),
                supplier_ref=str(val('Ref Fornecedor') or ''), legacy_key=legacy_key,
                # Preço do IMOS: só referência, 3.ª opção e sempre provisório.
                imos_price=str(number(val('€ / un')) or '') if number(val('€ / un')) is not None else '',
                supplier=str(val('Fornecedor') or ''), in_list=str(val('Na lista') or ''))
    if headers is None:
        raise ValueError('Cabeçalhos do custo de ferragens não reconhecidos.')
    return list(result.values())


def unit(value):
    value = str(value or '').lower().strip().replace('²', '2')
    return {'m': 'ml', 'mt': 'ml', 'und': 'un', 'uni': 'un', 'unidade': 'un'}.get(value, value)


def calculate_cost(line, price):
    if price is None:
        return None, 'Máquina / tarifa V3 por associar' if line['kind'] == 'Produção' else 'Matéria-prima V3 por associar'
    qty, net = number(line.get('quantity')), number(price.get('net'))
    if qty is None or qty < 0:
        return None, 'Quantidade em falta/inválida'
    if net is None or net < 0:
        return None, 'Preço líquido em falta/inválido'
    source, target = unit(line['unit']), unit(price.get('unit'))
    if source == 'ml' and target == 'm2' and line['kind'] == 'Orlas':
        width = number(line.get('width'))
        if width is None or width <= 0:
            return None, 'Largura de orla em falta'
        qty = qty * width / 1000
    elif source != target or not source:
        return None, f'Unidades por compatibilizar: {source} / {target}'
    return qty * net, 'Calculado' if net else 'Preço líquido zero — confirmar'


def price_record(mp):
    return {'id': mp.id, 'ref': mp.ref_le or mp.ref_phc or '', 'description': mp.descricao,
            'unit': mp.unidade, 'net': str(mp.preco_liquido) if mp.preco_liquido is not None else None,
            'date': datetime.now().isoformat(timespec='seconds')}


def exact_price(line, catalog):
    matches = []
    for mp in catalog:
        if line.get('ref_phc') and line['ref_phc'] in (mp.ref_phc, mp.ref_le):
            matches.append(mp)
        elif mp.nome_imos and mp.nome_imos == line['name']:
            matches.append(mp)
    if not matches and line.get('kind') == 'Comprados' and line.get('supplier_ref'):
        # Objetos comprados: a ref do fornecedor que o Paulo põe na união do IMOS.
        from app.domain.materia_prima_types import normalizar_ref_fornecedor
        forn = normalizar_ref_fornecedor(line['supplier_ref'])
        matches = [mp for mp in catalog if forn and
                   normalizar_ref_fornecedor(getattr(mp, 'referencia_fornecedor', None)) == forn]
    return price_record(matches[0]) if len(matches) == 1 else None


def _sequencia_do_snapshot(ficheiro):
    """O contador do nome, ou -1 se o nome for do formato antigo."""
    partes = Path(ficheiro).stem.split('_')
    if len(partes) >= 5 and len(partes[3]) == 6 and partes[3].isdigit():
        return int(partes[3])
    return -1


def _ordem_do_snapshot(ficheiro):
    """Instante primeiro; dentro do mesmo instante, o contador.

    O relógio do Windows anda aos saltos: vinte mil leituras seguidas de
    `datetime.now()` dão quarenta valores distintos. Dois snapshots gravados um
    a seguir ao outro ficam com o MESMO carimbo no nome — e enquanto o
    desempate era o `uuid` do fim, era aleatório. Metade das vezes
    `latest_snapshot` devolvia o snapshot ANTIGO, e a análise reabria com os
    preços velhos sem ninguém dar por isso.
    """
    partes = Path(ficheiro).stem.split('_')
    return ('_'.join(partes[:3]), _sequencia_do_snapshot(ficheiro), Path(ficheiro).name)


def save_snapshot(path, data):
    path = writable_workbook(path)
    folder = path.parent / 'Analise_Lista_Material'
    folder.mkdir(exist_ok=True)
    # O contador é por pasta e só sobe. O uuid fica, mas passa a ser o que era
    # para ser: garantia de que dois processos não escolhem o mesmo nome —
    # nunca o critério de qual é o mais recente.
    sequencia = max((_sequencia_do_snapshot(f) for f in folder.glob('*.json')), default=0) + 1
    destination = folder / f"{datetime.now():%Y%m%d_%H%M%S_%f}_{sequencia:06d}_{uuid4().hex[:8]}.json"
    payload = {**data, 'workbook': path.name, 'saved_at': datetime.now().isoformat(timespec='seconds')}
    with destination.open('x', encoding='utf-8') as stream:
        json.dump(payload, stream, ensure_ascii=False, indent=2)
    return destination


def latest_snapshot(path, version):
    pasta = Path(path).parent / 'Analise_Lista_Material'
    for f in sorted(pasta.glob('*.json'), key=_ordem_do_snapshot, reverse=True):
        data = json.loads(f.read_text(encoding='utf-8'))
        if data.get('version') == version and data.get('workbook') == Path(path).name:
            return data
    return {}


# ---- Separador Custo_V3 (pedido do Paulo, 22-09-2026) ------------------------
#
# O separador é para AJUSTAR à mão: alterar quantidades e preços, tirar linhas
# (Incluir = N ou apagar a linha), inserir linhas novas — e os totais acertam.
# As linhas são uma Tabela do Excel para as fórmulas acompanharem as linhas
# inseridas. Margem por categoria e um coeficiente de segurança (em %) sobre o
# total já com as margens dão o TOTAL FINAL.

# Ordem pedida pelo Paulo (24-09-2026): placas | orlas | ferragens | purch | spp.
# A Produção não entra na tabela das linhas: vive no quadro dos tempos por setor.
CATEGORIAS_MATERIAIS = ('Placas', 'Orlas', 'Ferragens', 'Comprados', 'SPP')
CATEGORIAS_CUSTO = CATEGORIAS_MATERIAIS + ('Produção',)
MARGENS_CUSTO = {'Placas': 0, 'Orlas': 0.10, 'Ferragens': 0.15, 'SPP': 0.15, 'Comprados': 0.10, 'Produção': 0.12}
COEFICIENTE_CUSTO = 0.10
COR_CABECALHO = 0xD9EAD3     # o verde que o separador sempre teve (BGR)
COR_TOTAL_FINAL = 0xB5D5A9
COR_ENTRADA = 0xCCF2FF       # amarelo claro: células para alterar à mão
COR_SEPARADOR = 0xC9D3DD     # linha vazia entre categorias, num tom acastanhado
COR_SEM_TARIFA = 0xCED3F6    # rosado: falta o €/h no Martelo


def linhas_por_categoria(lines):
    """As linhas de material pela ordem das categorias, com None entre elas.

    O None é a linha separadora (vazia e de outra cor) no Martelo e no Excel.
    A Produção fica de fora; uma categoria desconhecida vai para o fim.
    """
    ordem = list(CATEGORIAS_MATERIAIS)
    for line in lines:
        if line['kind'] not in ordem and line['kind'] != 'Produção':
            ordem.append(line['kind'])
    result = []
    for kind in ordem:
        grupo = [line for line in lines if line['kind'] == kind]
        if grupo:
            if result:
                result.append(None)
            result.extend(grupo)
    return result
COLUNAS_CUSTO = ('Categoria', 'Artigo / material', 'Comp (mm)', 'Larg (mm)', 'Esp (mm)', 'Quantidade', 'Un.',
                 'Ref_LE — Descrição V3', 'Preço líquido', 'Un. preço', 'Fator', 'Incluir (S/N)', 'Custo €',
                 'Estado', 'Data do preço')
ULTIMA_COLUNA = 'O'


def _texto_formula(texto):
    return '"' + str(texto).replace('"', '""') + '"'


def _escrever_relatorio_custo(sheet, version, lines, prices, warnings, production):
    from app.services import tempos_lista_material_service as tempos

    fim = ULTIMA_COLUNA
    sheet.Range(f'A1:{fim}1').Merge()
    sheet.Cells(1, 1).Value2 = 'Custo de produção (parcial) — ' + version
    sheet.Range(f'A2:{fim}2').Merge()
    sheet.Cells(2, 1).Value2 = (
        ('Preços e tarifas V3 guardados nesta análise. Consulta Streamlit: ' + production['queried_at'] if production
         else 'Preços líquidos V3 guardados nesta análise. Tempos Streamlit ainda não consultados.')
        + ' Células amarelas alteráveis (quantidade, preço, fator, Incluir, horas, €/h, margens, coeficiente): '
        'os totais recalculam; também se podem apagar ou inserir linhas na tabela.')

    # Produção: uma linha por setor/máquina, com o €/h do Martelo sempre à vista.
    producao = tempos.completar_setores([l for l in lines if l['kind'] == 'Produção'],
                                        (production or {}).get('sectors'))
    materiais = linhas_por_categoria(lines)
    estado_setor = {s['sector']: s for s in (production or {}).get('sectors', [])}

    linha_setores = 15
    primeira_hora = linha_setores + 1
    ultima_hora = linha_setores + len(producao)
    linha_total_horas = ultima_hora + 1
    linha_avisos = linha_total_horas + 2
    header_row = linha_avisos + 2
    first, last = header_row + 1, header_row + max(1, len(materiais))
    tabela = 'TabCusto_' + sheet.Name[len('Custo_V3_'):]

    # --- Resumo por categoria, margens e coeficiente -------------------------
    sheet.Range('A4:E4').Value = (('Categoria', 'Custo conhecido €', 'Linhas sem custo', 'Margem %',
                                   'Custo c/ margem €'),)
    for r, category in enumerate(CATEGORIAS_CUSTO, 5):
        sheet.Cells(r, 1).Value2 = category
        sheet.Cells(r, 4).Value2 = MARGENS_CUSTO[category]
        sheet.Cells(r, 5).Formula = f'=B{r}*(1+D{r})'
    sheet.Cells(11, 1).Value2 = 'Total conhecido (parcial)'
    sheet.Cells(11, 2).Formula = '=SUM(B5:B10)'
    sheet.Cells(11, 3).Formula = '=SUM(C5:C10)'
    sheet.Cells(11, 4).Formula = '=IF(B11=0,0,E11/B11-1)'
    sheet.Cells(11, 5).Formula = '=SUM(E5:E10)'
    sheet.Range('A12:C12').Merge()
    sheet.Cells(12, 1).Value2 = 'Coeficiente de segurança («cagaço») — % sobre o total c/ margens'
    sheet.Cells(12, 4).Value2 = COEFICIENTE_CUSTO
    sheet.Cells(12, 5).Formula = '=E11*D12'
    sheet.Range('A13:D13').Merge()
    sheet.Cells(13, 1).Value2 = 'TOTAL FINAL (custo + margens + coeficiente)'
    sheet.Cells(13, 5).Formula = '=E11+E12'
    sheet.Range('B5:B11').NumberFormat = '0.00'
    sheet.Range('E5:E13').NumberFormat = '0.00 €'
    sheet.Range('C5:C11').NumberFormat = '0'
    sheet.Range('D5:D12').NumberFormat = '0%'
    sheet.Range('D11').NumberFormat = '0.0%'
    sheet.Range('A4:E4').WrapText = True
    for area in ('A4:E4', 'A11:E11', 'A12:E12'):
        sheet.Range(area).Font.Bold = True
        sheet.Range(area).Interior.Color = COR_CABECALHO
    for area in ('D5:D10', 'D12'):
        sheet.Range(area).Interior.Color = COR_ENTRADA
    sheet.Range('A13:E13').Font.Bold = True
    sheet.Range('A13:E13').Font.Size = 13
    sheet.Range('A13:E13').Interior.Color = COR_TOTAL_FINAL
    sheet.Range('A4:E13').Borders.LineStyle = 1

    # --- Tempos por setor: horas × €/h do Martelo -------------------------------
    # Pedido do Paulo (24-09-2026): o €/h vem SEMPRE preenchido do V3, haja ou não
    # horas no Streamlit, para se poderem escrever as horas à mão e o custo sair.
    sheet.Range(f'A{linha_setores}:I{linha_setores}').Value = (
        ('Setor', 'Máquina / centro (Streamlit)', 'Horas registadas', 'Horas estimadas', '€/h (V3)',
         'Custo €', 'Máquina V3 — tarifa (Configurações › Operações / Máquinas)', None, 'Estado'),)
    sheet.Range(f'G{linha_setores}:H{linha_setores}').Merge()
    sheet.Range(f'I{linha_setores}:{fim}{linha_setores}').Merge()
    cabecalho_horas = sheet.Range(f'A{linha_setores}:{fim}{linha_setores}')
    cabecalho_horas.WrapText = True
    cabecalho_horas.Font.Bold = True
    cabecalho_horas.Interior.Color = COR_CABECALHO
    sheet.Rows(linha_setores).RowHeight = 32
    setores_vistos = set()
    sem_tarifa = []
    for r, line in enumerate(producao, primeira_hora):
        stage = line.get('sector')
        label = tempos.SECTORS[stage][0] if stage in tempos.SECTORS else line['name']
        entry = estado_setor.get(stage, {})
        price = prices.get(line['key'])
        tarifa, tem_tarifa = tempos.descrever_tarifa(line, price, curto=True)
        estimado = number(entry.get('estimated')) if stage not in setores_vistos else None
        setores_vistos.add(stage)
        horas, euros_hora = number(line.get('quantity')), number((price or {}).get('net'))
        for area in (f'A{r}:B{r}', f'G{r}', f'I{r}'):
            sheet.Range(area).NumberFormat = '@'
        sheet.Range(f'A{r}:I{r}').Value = ((
            label, line.get('machine') or '—',
            float(horas) if horas is not None else None,
            float(estimado) if estimado is not None else None,
            float(euros_hora) if euros_hora is not None else None,
            None, tarifa, None,
            entry.get('state') or ('Tempos não consultados' if not production else 'Estado por confirmar')),)
        sheet.Cells(r, 6).Formula = f'=N(C{r})*N(E{r})'
        sheet.Range(f'G{r}:H{r}').Merge()
        sheet.Range(f'I{r}:{fim}{r}').Merge()
        if not tem_tarifa:
            sem_tarifa.append(r)
    sheet.Range(f'A{linha_total_horas}:B{linha_total_horas}').Merge()
    sheet.Cells(linha_total_horas, 1).Value2 = 'Total produção'
    for coluna in ('C', 'D', 'F'):
        sheet.Range(f'{coluna}{linha_total_horas}').Formula = (
            f'=SUM({coluna}{primeira_hora}:{coluna}{ultima_hora})')
    sheet.Range(f'G{linha_total_horas}:H{linha_total_horas}').Merge()
    sheet.Range(f'I{linha_total_horas}:{fim}{linha_total_horas}').Merge()
    horas_area = sheet.Range(f'A{primeira_hora}:{fim}{linha_total_horas}')
    horas_area.VerticalAlignment = -4108
    sheet.Range(f'C{primeira_hora}:F{linha_total_horas}').NumberFormat = '0.00'
    for coluna in ('C', 'E'):
        sheet.Range(f'{coluna}{primeira_hora}:{coluna}{ultima_hora}').Interior.Color = COR_ENTRADA
    for r in sem_tarifa:
        sheet.Range(f'G{r}:H{r}').Interior.Color = COR_SEM_TARIFA
        sheet.Range(f'G{r}:H{r}').WrapText = True
        sheet.Rows(r).RowHeight = 30   # célula unida não ajusta a altura sozinha
    total_horas = sheet.Range(f'A{linha_total_horas}:{fim}{linha_total_horas}')
    total_horas.Font.Bold = True
    total_horas.Interior.Color = COR_CABECALHO
    sheet.Range(f'A{linha_setores}:{fim}{linha_total_horas}').Borders.LineStyle = 1

    sheet.Range(f'A{linha_avisos}:{fim}{linha_avisos}').Merge()
    sheet.Cells(linha_avisos, 1).Value2 = ' | '.join(
        warnings + (production or {}).get('warnings', [])
        + ([production['last_query_error']] if production and production.get('last_query_error') else []))[:32000]
    sheet.Range(f'A{linha_avisos}:{fim}{linha_avisos}').WrapText = True
    sheet.Rows(linha_avisos).RowHeight = 42

    # --- Linhas: Tabela do Excel com fórmulas -----------------------------------
    sheet.Range(f'A{header_row}:{fim}{header_row}').Value = (COLUNAS_CUSTO,)
    estados = []
    separadores = []
    geral = sheet.Cells(1, 40).NumberFormat   # «Geral» na língua do Excel instalado
    for r, line in enumerate(materiais, first):
        if line is None:
            separadores.append(r)
            continue
        price = prices.get(line['key']) or {}
        cost, state = calculate_cost(line, price or None)
        factor = Decimal(1)
        if line['kind'] == 'Orlas' and unit(price.get('unit')) == 'm2':
            factor = (number(line.get('width')) or Decimal(0)) / 1000
        estados.append((r, state))
        values = [line['kind'], line['name'], line.get('length', ''), line.get('width', ''), line.get('thickness', ''),
                  float(number(line['quantity'])) if number(line['quantity']) is not None else None, line['unit'],
                  ' — '.join(str(price.get(k) or '') for k in ('ref', 'description')) if price else '',
                  float(number(price.get('net'))) if number(price.get('net')) is not None else None,
                  price.get('unit', ''), float(factor), 'S', None, None, price.get('date', '')]
        # Texto externo é sempre texto, nunca uma fórmula do ficheiro de origem.
        for area in (f'A{r}:E{r}', f'G{r}:H{r}', f'J{r}', f'L{r}', f'O{r}'):
            sheet.Range(area).NumberFormat = '@'
        for column, formato in ((6, '0.00'), (9, '0.00###'), (11, '0.000'), (13, '0.00')):
            sheet.Cells(r, column).NumberFormat = formato
        sheet.Range(f'A{r}:{fim}{r}').Value = (tuple(values),)
        for column in (3, 4, 5, 6, 9, 11):
            numeric = number(values[column - 1])
            if numeric is not None:
                if column in (3, 4, 5):
                    sheet.Cells(r, column).NumberFormat = geral
                sheet.Cells(r, column).Value2 = float(numeric)
        # Sem custo mas com preço (unidades por compatibilizar, largura em falta...):
        # o fator fica vazio e a linha dá 0 € até alguém o acertar à mão.
        if cost is None and number(price.get('net')) is not None:
            sheet.Cells(r, 11).Value2 = None
    tab = sheet.ListObjects.Add(1, sheet.Range(f'A{header_row}:{fim}{last}'), None, 1)
    tab.Name = tabela
    tab.TableStyle = ''
    tab.ListColumns('Custo €').DataBodyRange.Formula = (
        '=IF([@[Incluir (S/N)]]="N",0,N([@Quantidade])*N([@[Preço líquido]])*N([@Fator]))')
    # O Estado é diferente em cada linha (guarda o motivo de não haver custo). Escrito
    # linha a linha, o Excel copiava cada fórmula para a coluna inteira da Tabela.
    autocorrect = sheet.Application.AutoCorrect
    preencher = autocorrect.AutoFillFormulasInLists
    autocorrect.AutoFillFormulasInLists = False
    try:
        for r, state in estados:
            sem_custo = state if state != 'Calculado' else 'Sem custo — confirmar preço'
            sheet.Cells(r, 14).Formula = (f'=IF(L{r}="N","Fora da conta (Incluir = N)",IF(M{r}=0,'
                                          f'{_texto_formula(sem_custo)},"Calculado"))')
    finally:
        autocorrect.AutoFillFormulasInLists = preencher
    tab.ShowTotals = True
    tab.ListColumns('Custo €').TotalsCalculation = 1       # soma
    tab.TotalsRowRange.Cells(1, 1).Value2 = 'Total das linhas'
    for r, category in enumerate(CATEGORIAS_CUSTO, 5):
        if category == 'Produção':
            # A produção soma o quadro dos tempos (horas × €/h), não a Tabela.
            sheet.Cells(r, 2).Formula = f'=F{linha_total_horas}'
            sheet.Cells(r, 3).Formula = (f'=COUNTIFS(F{primeira_hora}:F{ultima_hora},0,'
                                         f'I{primeira_hora}:I{ultima_hora},"<>Não aplicável")')
            continue
        sheet.Cells(r, 2).Formula = f'=SUMIFS({tabela}[Custo €],{tabela}[Categoria],A{r})'
        sheet.Cells(r, 3).Formula = (f'=COUNTIFS({tabela}[Categoria],A{r},{tabela}[Incluir (S/N)],"S",'
                                     f'{tabela}[Custo €],0)')
    corpo = tab.DataBodyRange
    incluir = tab.ListColumns('Incluir (S/N)').DataBodyRange
    incluir.Validation.Delete()
    incluir.Validation.Add(3, 1, 1, 'S,N')
    incluir.HorizontalAlignment = -4108
    for coluna in ('Quantidade', 'Preço líquido', 'Fator', 'Incluir (S/N)'):
        tab.ListColumns(coluna).DataBodyRange.Interior.Color = COR_ENTRADA
    for coluna, formato in (('Quantidade', '0.00'), ('Preço líquido', '0.00###'), ('Fator', '0.000'),
                            ('Custo €', '0.00')):
        tab.ListColumns(coluna).DataBodyRange.NumberFormat = formato
    # Linha vazia e de outra cor entre categorias (pedido do Paulo, 24-09-2026). O
    # 0 que a fórmula do custo dá nela fica escondido.
    for r in separadores:
        faixa = sheet.Range(f'A{r}:{fim}{r}')
        faixa.Interior.Color = COR_SEPARADOR
        faixa.NumberFormat = ';;;'
        faixa.Validation.Delete()
        sheet.Rows(r).RowHeight = 8
    tab.TotalsRowRange.Cells(1, 13).NumberFormat = '0.00'
    tab.ListColumns('Data do preço').TotalsCalculation = 0   # sem a contagem que o Excel põe
    tab.HeaderRowRange.Font.Bold = True
    tab.HeaderRowRange.Interior.Color = COR_CABECALHO
    tab.HeaderRowRange.WrapText = True
    tab.TotalsRowRange.Font.Bold = True
    tab.TotalsRowRange.Interior.Color = COR_CABECALHO
    tab.Range.Borders.LineStyle = 1
    # Riscado quando Incluir = N (continua visível para se ver o que foi tirado).
    riscado = corpo.FormatConditions.Add(2, None, f'=$L{first}="N"')
    riscado.Font.Strikethrough = True
    riscado.Font.Color = 0x9A9A9A

    sheet.Range(f'A1:{fim}1').Font.Size = 16
    sheet.Range(f'A1:{fim}1').Font.Bold = True
    for col, width in (('A', 20), ('B', 48), ('C', 13), ('D', 13), ('E', 11), ('F', 13), ('G', 8), ('H', 60),
                       ('I', 14), ('J', 10), ('K', 10), ('L', 9), ('M', 15), ('N', 40), ('O', 23)):
        sheet.Columns(col).ColumnWidth = width
    sheet.Range(f'A2:{fim}2').WrapText = True
    sheet.Rows(2).RowHeight = 30

    last = header_row + tab.ListRows.Count + 1   # linha dos totais da tabela
    if production and production.get('events'):
        event_header = last + 3
        event_rows = [('Data', 'Setor', 'Máquina', 'Pessoa / login', 'Horas', 'Chave Streamlit', 'Plano')]
        # Data ao minuto: com segundos e milésimos não cabia e partia cada linha em duas.
        event_rows += [(str(e.get('data_registo') or '')[:16], e['setor'], e.get('maquina'), e.get('responsavel'),
                        e['horas'], e['bd_key'], e.get('bd_plano_corte')) for e in production['events']]
        for r, values in enumerate(event_rows, event_header):
            for (start, end), value in zip(((1, 1), (2, 2), (3, 4), (5, 7), (8, 8), (9, 11), (12, 15)), values):
                area = sheet.Range(sheet.Cells(r, start), sheet.Cells(r, end))
                area.Merge()
                area.NumberFormat = '@'
                sheet.Cells(r, start).Value2 = value
                if r > event_header and start == 8 and number(value) is not None:
                    area.NumberFormat = '0.00'
                    sheet.Cells(r, start).Value2 = float(number(value))
            last = r
        sheet.Range(f'A{event_header}:{fim}{event_header}').Font.Bold = True
        sheet.Range(f'A{event_header}:{fim}{event_header}').Interior.Color = COR_CABECALHO
    from app.services.lista_material_pdf_service import pagina_a3_ao_baixo
    pagina_a3_ao_baixo(sheet)
    setup = sheet.PageSetup
    setup.PrintTitleRows = f'${header_row}:${header_row}'
    setup.PrintArea = f'A1:{fim}{last}'


def export_cost_report(path, expected_hash, version, lines, prices, warnings, *, production=None):
    """Novo relatório nativo por análise; não substitui folhas editadas pelo utilizador."""
    path = writable_workbook(path)
    if fingerprint(path) != expected_hash:
        raise ValueError('O Excel mudou. Reanalise antes de inserir o relatório.')
    excel = importlib.import_module('win32com.client').DispatchEx('Excel.Application')
    book = None
    try:
        excel_com.preparar_excel(excel)
        book = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)
        if book.ReadOnly or fingerprint(path) != expected_hash:
            raise ValueError('Feche o Excel e reanalise antes de inserir o relatório.')
        book.SaveCopyAs(str(backup_path(path, 'antes_relatorio')))
        sheet = book.Worksheets.Add(After=book.Worksheets.Item(book.Worksheets.Count))
        sheet.Name = 'Custo_V3_' + datetime.now().strftime('%y%m%d_%H%M%S') + '_' + uuid4().hex[:3]
        name = sheet.Name
        _escrever_relatorio_custo(sheet, version, lines, prices, warnings, production)
        sheet.Calculate()
        excel_com.recalcular(excel)
        book.Save()
        return name
    finally:
        if book is not None:
            book.Close(False)
        excel.Quit()


HARDWARE_FILENAME = '5_Custo_Obra_Ferragens.xlsx'


def archive_hardware_file(path, source):
    """Transferência pedida pelo utilizador: só remove a fonte após verificar a cópia."""
    path = writable_workbook(path)
    source = Path(source).resolve()
    if any(part.upper() == 'ORIGINAL' for part in source.parts):
        raise ValueError('Não pode mover ficheiros da pasta ORIGINAL.')
    destination = path.parent / HARDWARE_FILENAME
    version_name = path.stem.removeprefix('Lista_Material_')
    if source != destination and not source.name.startswith(version_name + '_5_Custo_Obra_Ferragens'):
        raise ValueError('Fonte de ferragens não corresponde à obra.')
    if source == destination:
        return destination
    original_hash = fingerprint(source)
    if destination.exists():
        if fingerprint(destination) != original_hash:
            raise ValueError('A obra já tem 5_Custo_Obra_Ferragens.xlsx diferente. Ambos os ficheiros foram preservados.')
    elif source.parent == destination.parent:
        source.rename(destination)
        return destination
    else:
        with source.open('rb') as incoming, destination.open('xb') as outgoing:
            shutil.copyfileobj(incoming, outgoing)
        shutil.copystat(source, destination)
    if fingerprint(destination) != original_hash or fingerprint(source) != original_hash:
        raise ValueError('Ficheiro alterado durante a transferência. A fonte foi preservada.')
    source.unlink()  # Autorizado: mover o export IMOS para a obra, após cópia verificada.
    return destination


def import_hardware_cost(path, source):
    """Importa o separador e só depois transfere a fonte para o nome normalizado."""
    path = writable_workbook(path)
    source = Path(source).resolve()
    version_name = path.stem.removeprefix('Lista_Material_')
    if not source.name.startswith(version_name + '_5_Custo_Obra_Ferragens') and source != path.parent / HARDWARE_FILENAME:
        raise ValueError('O ficheiro de ferragens não corresponde à versão desta Lista Material.')
    destination = path.parent / HARDWARE_FILENAME
    if destination.exists() and fingerprint(destination) != fingerprint(source):
        raise ValueError('A obra já tem um ficheiro de custos diferente. Os ficheiros foram preservados.')
    book_source = load_workbook(source, read_only=True, data_only=True)
    try:
        if len(book_source.sheetnames) != 1:
            raise ValueError('O ficheiro de custos deve ter um único separador.')
        incoming_rows = hardware_rows(book_source.active.values)
    finally:
        book_source.close()
    current = load_workbook(path, read_only=True, data_only=True)
    try:
        existing = [s for s in current if 'custo_obra_ferragens' in s.title.lower()]
        if len(existing) > 1 or (existing and hardware_rows(existing[0].values) != incoming_rows):
            raise ValueError('O separador existente difere da fonte. Foi preservado e a fonte não foi removida.')
    finally:
        current.close()
    excel = importlib.import_module('win32com.client').DispatchEx('Excel.Application')
    book = src = None
    try:
        excel_com.preparar_excel(excel)
        book = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)
        if book.ReadOnly:
            raise ValueError('Feche o Excel antes de importar os custos de ferragens.')
        if not any('custo_obra_ferragens' in str(s.Name).lower() for s in book.Worksheets):
            src = excel.Workbooks.Open(str(source), UpdateLinks=0, ReadOnly=True)
            book.SaveCopyAs(str(backup_path(path, 'antes_ferragens')))
            src.Worksheets.Item(1).Copy(After=book.Worksheets.Item(book.Worksheets.Count))
            book.Worksheets.Item(book.Worksheets.Count).Name = '5_Custo_Obra_Ferragens'
            excel_com.recalcular(excel)
        book.Save()
    finally:
        if src is not None:
            src.Close(False)
        if book is not None:
            book.Close(False)
        excel.Quit()
    return archive_hardware_file(path, source)
